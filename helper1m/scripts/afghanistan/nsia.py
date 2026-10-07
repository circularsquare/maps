"""Readers for the NSIA (now GSIA) "Estimated Population of Afghanistan" tables.

Each province has a table "Settled Population of <Province> Province by
Administrative Unit, Residence and Sex": one row per administrative unit (the
provincial centre, the original districts and the temporary districts), with
rural / urban / total by sex. 1403 and 1404 come as Excel workbooks (sheet
"نفوس ولایات"); 1405 only as a PDF, read from its text layer.

Every reader returns a list of dicts in table order:
    {"prov": <English province name as printed>, "sno": int,
     "name": <English unit name>, "name_fa": <Dari name or "">, "pop": int}
"""
from __future__ import annotations

import re
from pathlib import Path

PROV_RE = re.compile(r"Settled\s+Population\s+of\s+(.+?)\s+Province", re.I)
NUM_RE = re.compile(r"^(-|\d{1,3}(?:,\d{3})*|\d+)(?=\s|$)")
LATIN = re.compile(r"[A-Za-z]")


def clean(s):
    return re.sub(r"\s+", " ", str(s)).strip()


def read_xlsx(path: Path, strict=True):
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    # 1403-1404 keep the district tables on their own sheet; 1390-1402 put
    # everything on the first one
    ws = wb["نفوس ولایات"] if "نفوس ولایات" in wb.sheetnames else wb.worksheets[0]
    out, prov, totals, shift = [], None, {}, None
    for row in ws.iter_rows(values_only=True):
        cells = list(row)
        text = " ".join(str(c) for c in cells if isinstance(c, str))
        m = PROV_RE.search(text)
        if m:
            prov, shift = clean(m.group(1)), None
            continue
        # the table starts in column B in 1404 and column A in 1403: align on
        # the "S.No" heading
        heads = [clean(c) if isinstance(c, str) else "" for c in cells]
        if "S.No" in heads:
            k = heads.index("S.No")
            shift = 1 - k
            continue
        if prov is None or shift is None:
            continue
        cells = [None] * shift + cells if shift > 0 else cells[-shift:]
        if len(cells) < 12:
            continue
        sno, name = cells[1], cells[2]
        nums = [c for c in cells[3:12]]
        if isinstance(name, str) and clean(name).lower() == "total":
            totals[prov] = int(round(nums[8] or 0))
            continue
        if not isinstance(sno, (int, float)) or not isinstance(name, str):
            continue
        vals = [int(round(v)) if isinstance(v, (int, float)) else 0 for v in nums]
        rural, urban, total = vals[2], vals[5], vals[8]
        assert abs(rural + urban - total) <= 2, (prov, name, vals)
        fa = clean(cells[12]) if len(cells) > 12 and isinstance(cells[12], str) else ""
        out.append({"prov": prov, "sno": int(sno), "name": clean(name),
                    "name_fa": fa, "pop": total})
    # every province's units must add to its printed total (the workbook holds
    # unrounded figures, so allow one person of rounding per unit)
    for p, t in totals.items():
        units = [r["pop"] for r in out if r["prov"] == p]
        ok = abs(sum(units) - t) <= len(units)
        if strict:
            assert ok, f"{path.name} {p}: units {sum(units)} != total {t}"
        elif not ok:
            print(f"  ({path.name} {p}: units {sum(units)} != printed total {t})")
    return out


def read_pdf(path: Path):
    """1405: the PDF's text layer prints each unit as S.No, its English name
    (one or more lines), then nine figures (rural F/M/both, urban F/M/both,
    total F/M/both; '-' for none), the last followed by the Dari name. The
    caption with the province name comes after the rows of its table."""
    import fitz
    doc = fitz.open(path)
    out, prev_cap, totals = [], None, {}
    for pno, page in enumerate(doc):
        lines = [ln.strip() for ln in page.get_text().splitlines() if ln.strip()]
        # the running head ("19 / Population of Afghanistan 2026-27") looks
        # like a row start; drop it before parsing
        lines = [ln for ln in lines if "Population of Afghanistan" not in ln]
        rows = _pdf_rows(lines)
        caps = [clean(m.group(1)) for ln in lines
                if (m := PROV_RE.search(ln)) and "(Person)" in ln]
        # a long table runs onto the next page, which has no caption of its own
        prov = caps[0] if caps else prev_cap
        prev_cap = caps[0] if caps else None
        if prov is None:
            continue
        for r in rows:
            r["page"] = pno + 1
            r["prov"] = prov
            # the page number followed by the column heads parses as a row:
            # it carries the province total
            if "S.No" in r["name"]:
                totals[prov] = r["pop"]
            else:
                out.append(r)
    for p, t in totals.items():
        units = [r["pop"] for r in out if r["prov"] == p]
        assert abs(sum(units) - t) <= len(units), f"{path.name} {p}: units {sum(units)} != total {t}"
    return out


def _pdf_rows(lines):
    pending, i = [], 0
    while i < len(lines):
        ln = lines[i]
        if re.fullmatch(r"\d{1,2}", ln) and i + 1 < len(lines) and LATIN.search(lines[i + 1]) \
                and not NUM_RE.match(lines[i + 1]):
            sno = int(ln)
            j, name = i + 1, []
            while j < len(lines) and not NUM_RE.match(lines[j]):
                name.append(lines[j])
                j += 1
            vals, last = [], ""
            while j < len(lines) and len(vals) < 9:
                mm = NUM_RE.match(lines[j])
                if not mm:
                    break
                tok = mm.group(1)
                vals.append(0 if tok == "-" else int(tok.replace(",", "")))
                last = lines[j][mm.end():]
                j += 1
            # rural + urban = total, give or take the table's own rounding
            # (Khulam 1405: 27,984 + 66,610 printed against 94,595)
            if len(vals) == 9 and abs(vals[2] + vals[5] - vals[8]) <= 2:
                fa = re.sub(r"[\d۰-۹]+$", "", clean(last)).strip()
                pending.append({"prov": None, "sno": sno, "name": clean(" ".join(name)),
                                "name_fa": fa, "pop": vals[8]})
                i = j
                continue
        i += 1
    return pending
