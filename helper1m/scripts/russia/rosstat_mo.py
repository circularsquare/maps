"""Parse Rosstat's "Численность населения Российской Федерации по муниципальным
образованиям на 1 января <year>" workbooks (Chisl_MO_01-01-<year>.xlsx and
earlier names), and its subject list.

The data sheet is one long list: a subject header row, then its municipal
districts / okrugs / urban okrugs, each followed by its settlements and
localities. The first column is a ТЕРСОН-МО code. Its first 8 digits are the
OKTMO code; what follows is either the OKTMO locality extension or a pair of
flag digits, and the cell is sometimes text with stray spaces in it
("41754000 0 0", "95701000001 1 0 0 3"). So the parse takes the digits,
keeps the first 8 as the OKTMO, and never trusts the rest.

OKTMO structure used here (8 digits, AA B CC DDD):
  AA     subject
  B      type: 3 intra-city municipality of a federal city, 5 municipal okrug,
         6 municipal district, 7 urban okrug, 9 (Moscow) settlements / urban
         okrugs of New Moscow; 2 = Moscow's administrative okrugs and St
         Petersburg's districts, which are not municipalities
  DDD    000 for the unit itself; anything else is a settlement inside it
Autonomous okrugs inside an oblast (Nenets 118, Khanty-Mansi 718, Yamalo-Nenets
719) move everything one digit right: AA 8 B C DDD.
"""
import re

import openpyxl

# OKTMO subject prefix -> ISO 3166-2. Crimea (35) and Sevastopol (67) are kept
# here so they can be switched on; fetch.py decides whether they ship.
OKTMO_ISO = {
    "01": "RU-ALT", "03": "RU-KDA", "04": "RU-KYA", "05": "RU-PRI", "07": "RU-STA",
    "08": "RU-KHA", "10": "RU-AMU", "11": "RU-ARK", "118": "RU-NEN", "12": "RU-AST",
    "14": "RU-BEL", "15": "RU-BRY", "17": "RU-VLA", "18": "RU-VGG", "19": "RU-VLG",
    "20": "RU-VOR", "22": "RU-NIZ", "24": "RU-IVA", "25": "RU-IRK", "26": "RU-IN",
    "27": "RU-KGD", "28": "RU-TVE", "29": "RU-KLU", "30": "RU-KAM", "32": "RU-KEM",
    "33": "RU-KIR", "34": "RU-KOS", "35": "UA-43", "36": "RU-SAM", "37": "RU-KGN",
    "38": "RU-KRS", "40": "RU-SPE", "41": "RU-LEN", "42": "RU-LIP", "44": "RU-MAG",
    "45": "RU-MOW", "46": "RU-MOS", "47": "RU-MUR", "49": "RU-NGR", "50": "RU-NVS",
    "52": "RU-OMS", "53": "RU-ORE", "54": "RU-ORL", "56": "RU-PNZ", "57": "RU-PER",
    "58": "RU-PSK", "60": "RU-ROS", "61": "RU-RYA", "63": "RU-SAR", "64": "RU-SAK",
    "65": "RU-SVE", "66": "RU-SMO", "67": "UA-40", "68": "RU-TAM", "69": "RU-TOM",
    "70": "RU-TUL", "71": "RU-TYU", "718": "RU-KHM", "719": "RU-YAN", "73": "RU-ULY",
    "75": "RU-CHE", "76": "RU-ZAB", "77": "RU-CHU", "78": "RU-YAR", "79": "RU-AD",
    "80": "RU-BA", "81": "RU-BU", "82": "RU-DA", "83": "RU-KB", "84": "RU-AL",
    "85": "RU-KL", "86": "RU-KR", "87": "RU-KO", "88": "RU-ME", "89": "RU-MO",
    "90": "RU-SE", "91": "RU-KC", "92": "RU-TA", "93": "RU-TY", "94": "RU-UD",
    "95": "RU-KK", "96": "RU-CE", "97": "RU-CU", "98": "RU-SA", "99": "RU-YEV",
}
AO_PREFIXES = ("118", "718", "719")
FEDERAL_CITIES = ("40", "45", "67")


def subject_of(oktmo):
    """ISO code of the subject an 8-digit OKTMO belongs to."""
    if oktmo[:3] in AO_PREFIXES:
        return OKTMO_ISO[oktmo[:3]]
    return OKTMO_ISO[oktmo[:2]]


def kind(oktmo):
    """'subject', 'unit' (a level-2 municipality) or None (anything finer or
    an aggregate that is not a municipality)."""
    if oktmo[:3] in AO_PREFIXES:
        if oktmo[3:] == "00000":
            return "subject"
        if oktmo[5:] == "000" and oktmo[3] != "0":
            return "unit"
        return None
    if oktmo[2:] == "000000":
        return "subject"
    # B CC with CC = 00 is an aggregate (2024 prints the old Koryak and Aga
    # Buryat okrugs, 30800000 and 76800000, as rows of their own).
    if oktmo[5:] != "000" or oktmo[3:5] == "00":
        return None
    if oktmo[:2] in ("40", "45"):
        # Moscow's administrative okrugs and St Petersburg's districts (B = 2)
        # rather than the ~250 intra-city municipalities under them: New
        # Moscow's settlements were regrouped in 2024, so the municipalities
        # do not pair across years, while the okrugs and districts do.
        return "unit" if oktmo[2] == "2" else None
    if oktmo[:2] in FEDERAL_CITIES:
        return "unit" if oktmo[2] in "39" else None
    # 8 = units of the former autonomous okrugs merged into a krai or oblast
    # (Aga Buryat in Zabaykalsky, Koryak in Kamchatka), which keep their codes.
    return "unit" if oktmo[2] in "5678" else None


# A municipality's own name, as opposed to a settlement inside it.
MUNI_NAME = re.compile(r"муниципальн\w* (район|округ)|городск\w* округ", re.I)
SETTLEMENT_NAME = re.compile(r"поселени|сельсовет|межселен", re.I)

# Rows printed without any code at all; the code is filled from OKTMO.
MISSING_CODES = {
    "Городской округ г Иркутск": "25701000",  # 2025 workbook
}


def read_rows(path):
    """[(oktmo8 or None, name, pop_total, raw_code)] for every row of the data
    sheet that carries a population and either a municipal-length code or no
    code at all."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    sheet = next(s for s in wb.sheetnames if re.search(r"М[ОO]|MO", s))
    out = []
    for row in wb[sheet].iter_rows(values_only=True):
        vals = [c for c in row if c is not None and str(c).strip() != ""]
        if len(vals) < 2:
            continue
        digits = re.sub(r"\s", "", str(vals[0]))
        if not digits.isdigit():
            # No code: the name is in the first cell and the population next.
            if isinstance(vals[1], (int, float)) and MUNI_NAME.search(str(vals[0])):
                out.append(("", str(vals[0]).strip(), int(vals[1]), ""))
            continue
        if len(vals) < 3 or not isinstance(vals[2], (int, float)):
            continue
        # Municipalities carry 10 digits (some 7-9, see parse); 11, 14 and
        # 15 are localities inside one. Irkutsk city is printed with 12 in
        # 2024, so a longer code is taken when the name is a municipality's.
        if len(digits) in (11, 12):
            if not MUNI_NAME.search(str(vals[1])) or SETTLEMENT_NAME.search(str(vals[1])):
                continue
        elif not 7 <= len(digits) <= 10:
            continue
        out.append((digits, str(vals[1]).strip(), int(vals[2]), str(vals[0])))
    wb.close()
    return out


def normalise(digits, cur_prefix):
    """8-digit OKTMO from a code cell's digits. A 9-digit code has either lost
    its leading zero to a numeric cell (Khabarovsk's 08...) or one trailing
    digit (Kirov's Nemsky 335260000); the subject block says which. One
    subject header (Nenets, 1180010) has 7."""
    if len(digits) == 9:
        digits = digits + "0" if cur_prefix and cur_prefix.startswith(digits[:2]) else "0" + digits
    if len(digits) == 7:
        digits = digits[:5].ljust(10, "0")
    return digits[:8]


def parse(path):
    """Returns (subjects, units, notes):
    subjects = {iso: total}, from the subject header rows
    units    = {oktmo8: (iso, name, pop)} for the level-2 municipalities
    notes    = the code repairs made, one line each
    """
    subjects, units, own = {}, {}, set()
    iso_prefix = {v: k for k, v in OKTMO_ISO.items()}
    cur = None  # subject of the last header row
    notes = []
    for digits, name, pop, raw in read_rows(path):
        oktmo = normalise(digits, iso_prefix.get(cur)) if digits else None
        if len(digits) == 9:
            notes.append(f"{name}: 9-digit code {raw}; read as {oktmo}")
        if oktmo is None:
            if name not in MISSING_CODES:
                raise ValueError(f"{path}: municipality without a code: {name} {pop}")
            oktmo = MISSING_CODES[name]
            notes.append(f"{name}: no code printed; OKTMO {oktmo} filled in")
        # A unit whose code names another subject than the block it sits in is
        # a typo in the workbook (2025: Soletsky okrug of Novgorod printed as
        # 48538000 for 49538000). The block is right; mend the prefix.
        if cur and not iso_prefix[cur].startswith(oktmo[:2]) and kind(oktmo) == "unit":
            fixed = iso_prefix[cur][:2] + oktmo[2:]
            notes.append(f"{name}: coded {raw} inside {cur}; read as {fixed}")
            oktmo = fixed
        # A municipality's own row whose code points below municipal level
        # (2025: Selemdzhinsky district of Amur printed as 10645151): the name
        # says what it is, so take the code's municipal part.
        if (cur and kind(oktmo) is None and MUNI_NAME.search(name)
                and not SETTLEMENT_NAME.search(name) and oktmo[:2] not in FEDERAL_CITIES):
            fixed = oktmo[:5] + "000"
            notes.append(f"{name}: coded {raw}, a settlement-level code; read as {fixed}")
            oktmo = fixed
        # Arkhangelsk and Tyumen each appear twice: with their autonomous
        # okrugs, and without them ("без" / "кроме", coded 11001000 or the
        # bare subject code again). The "without" row is the subject here.
        if oktmo[:2] in ("11", "71") and re.search(r"\bбез\b|\bкроме\b", name):
            iso = subject_of(oktmo)
            subjects[iso] = pop
            own.add(iso)
            cur = iso
            continue
        k = kind(oktmo)
        if k == "subject":
            iso = subject_of(oktmo)
            cur = iso
            if iso not in own:
                subjects[iso] = pop
        elif k == "unit":
            if oktmo in units:
                raise ValueError(f"{path}: duplicate unit {oktmo} {name}")
            units[oktmo] = (subject_of(oktmo), name, pop)
    return subjects, units, notes
