"""Uruguay — boundaries for the 19 departamentos, and the population that goes on them.

Writes data/geo/uy/uy_departamentos.gpkg, data/geo/uy/uy_lookup.csv and
data/geo/uy/uy_pop_2023.csv.

## THE CODE JOIN IS AVAILABLE, GETS NINE OF NINETEEN RIGHT, AND IS A PERMUTATION

This is `sources/sv_geo.py`'s trap in a country where it is harder to see, because here the
naive join is *mostly* right.

**INE numbers Uruguay's departments with Montevideo first and the other eighteen
alphabetically after it** — 1 Montevideo, 2 Artigas, 3 Canelones, … 19 Treinta y Tres. That
numbering is INE's own `dpto` variable in the ENHA microdata and it is also LAPOP's, whose
`prov` for Uruguay is 1400 plus the same number. **COD-AB's `UY01`..`UY19` are alphabetical
with no Montevideo exception**: UY01 is Artigas and Montevideo is UY10.

So `UY{dpto:02d}` shifts the first ten departments by one place and leaves the last nine
alone:

    INE  1 Montevideo    -> UY01, which COD says is Artigas
    INE  2 Artigas       -> UY02, which COD says is Canelones
    INE  9 Lavalleja     -> UY09, which COD says is Maldonado
    INE 10 Maldonado     -> UY10, which COD says is Montevideo
    INE 11 Paysandú      -> UY11, which is Paysandú           (correct, and so are 12-19)

**Nine of nineteen coincide, and they are the last nine**, which is exactly the
shape that survives a spot check: look at Salto, Soriano and Treinta y Tres and the join
appears fine. The ten that move include Montevideo, 1.29 million people and 37% of the
country, which would have been drawn in Artigas.

A permutation preserves every total, so no reconciliation, no national figure and no row
count would show it (`sources/ni_geo.py`). **The join is therefore on NAME**, and
`check_code_join()` asserts that the code join still mispairs exactly ten, so that an OCHA
re-cut to INE's order stops the build rather than silently changing which department is which.

## THE POPULATION IS INE'S OWN ESTIMATE FOR 2023, THE CENSUS YEAR

Uruguay counted in 2023, so §9bn's rule applies and COD-PS is not the neutral choice here:
its Uruguay file is a projection off the 2011 census. INE's *Estimaciones y proyecciones,
revisión 2025* is the post-census series — `A.1.2 Departamentos.xlsx`, one sheet per
department, both sexes and five-year age bands, 1996 to 2023 — and it is read here for two
years rather than one:

  * **2023**, the magnitude the map is drawn on;
  * **2006**, the ENHA's own year, which is what `sources/uy.py` checks the survey's
    departmental weighting against. Comparing a 2006 survey's unit shares with a 2023
    population would fail on Uruguay's real internal migration rather than on a bad join.

**THE DRAWN UNIVERSE IS AGES 7 AND OVER AND THE BANDS ARE FIVE YEARS WIDE**, so `pop7` is
`total - (0-4) - 0.4 * (5-9)`: ages 5 and 6 are two fifths of the band. Nothing INE
publishes for Uruguay is by single year of age (the 2025 revision's `(100ymas)` files are
five-year bands too, the `100 y más` being the open top band). The 5-9 band is 6.60% of the
country, so an even split within it that is off by a tenth of the band is 0.66% of Uruguay;
the census 2023 persons file puts Montevideo's under-7 share at 6.95% against 7.16% from this
cut. Until 2026-09-15 the factor was 0.6, which removed ages 7 to 9 instead (sources/uy.md
§12.7). Stated here rather than hidden because it is arithmetic on a published band and not a
published number.

## MONTEVIDEO'S 62 BARRIOS, SINCE 2026-09-15

Montevideo is drawn at INE's 62 barrios instead of as one department (`sources/uy.md` §12).
Three things are read for them, and each is joined to the others by a key the next one checks.

  * **Polygons.** The Intendencia de Montevideo's layer `zon_v_sig_barrios`, from its WFS as a
    435 KB shapefile zip, *"BARRIOS DE MONTEVIDEO segun INE"*, free use under municipal
    resolution 640/10, cartography of December 2011. `nrobarrio` is INE's barrio number, 1 to
    62, which is also the ENHA's `barrio`. INE's own layer (`ine_barrios_mvd_nbi85` inside the
    45 MB *mapas vectoriales 2011* zip) carries the same numbers and agrees at IoU 0.990 or
    better on every barrio, so the two are one lineage and only the small one is fetched.
  * **People, all ages.** INE's census 2023 *Cuadro 15, Población por barrios* (weighted,
    1,302,721 people, 1,359 of them *sin dato de barrio*, living on the street). Joined to the
    polygons on the name, since Cuadro 15 carries no number: every pairing must be the other's
    best match, one to one.
  * **The 7+ cut.** Nothing INE tabulates gives age by barrio, and the under-7 share runs from
    4.0% (Tres Cruces) to 10.7% (Casavalle), so INE's one city-wide ratio would put barrios
    3.1% under to 4.2% over their own 7+ population, centre against periphery. The census 2023
    persons file
    (ANDA catalogue 781, the same click-through terms as the ENHA's) has `BARRIO85`, age
    `PERNA01` and a weight `W`. Weighted, it reproduces Cuadro 15 to within one person in every
    barrio and 3,499,451 nationally, the published census count; unweighted it runs 0.78
    (Villa García) to 0.99 (Punta Carretas) of the table, lowest in the poorest barrios, so
    the weight is not optional. **The CSV is latin-1**, and the weight column is a bare `W`.

The barrios then take **INE's own revision-2025 Montevideo total** at 30 June 2023, split by
each barrio's census share: all ages for `pop_2023`, 7 and over for `pop7_2023`. Montevideo's
department total, and every other department, is exactly what it was.

Usage:
    python sources/uy_geo.py --fetch    one ~1 MB zip from HDX, a 0.6 MB xlsx from INE, the
                                        Intendencia's barrio layer and Cuadro 15; the census
                                        persons file is a click-through, see census_extract()
    python sources/uy_geo.py            rebuild from data/raw/uy/
    python sources/uy_geo.py --census   also re-stream the census persons file (about 30 s)
"""

import difflib
import os
import re
import shutil
import ssl
import subprocess
import sys
import unicodedata
import urllib.parse
import urllib.request
import zipfile

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import geopandas as gpd
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
RAW = os.path.join(ROOT, "data", "raw", "uy")
SHP_DIR = os.path.join(RAW, "shp")
OUT_DIR = os.path.join(ROOT, "data", "geo", "uy")
OUT = os.path.join(OUT_DIR, "uy_departamentos.gpkg")
LOOKUP = os.path.join(OUT_DIR, "uy_lookup.csv")
POP_OUT = os.path.join(OUT_DIR, "uy_pop_2023.csv")

POP_XLSX = "A.1.2 Departamentos.xlsx"

UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                    "(KHTML, like Gecko) Chrome/125.0 Safari/537.36"}

DOWNLOADS = {
    # cod-ab-ury on HDX. ADM0..ADM2 plus the boundary lines; 1 MB in total.
    "ury_adm_2020_shp.zip":
        "https://data.humdata.org/dataset/d31c3690-b6d9-461c-8ac6-553dd9c69b81/resource/"
        "75c7c51c-8f05-4b49-8b47-2f94112f986c/download/ury_adm_2020_shp.zip",
    # INE's own post-census estimates, revisión 2025. Linked from
    # gub.uy/instituto-nacional-estadistica/estimaciones2025 and nowhere in the census pages.
    POP_XLSX:
        "https://www5.ine.gub.uy/documents/Demograf%C3%ADayEESS/SERIES%20Y%20OTROS/"
        "Estimaciones%20y%20proyecciones/Revisi%C3%B3n%202025/A.1.2%20Departamentos.xlsx",
    # Montevideo's barrios, INE's numbering, from the Intendencia's geoserver (geonetwork record
    # 1277c8cd-3e7a-4afd-8289-aeae893ce0db).
    "im_barrios.zip":
        "https://montevideo.gub.uy/app/geoserver/ows?service=WFS&version=1.1.0&request=GetFeature"
        "&typeName=mapstore-tematicas%3Azon_v_sig_barrios&outputFormat=SHAPE-ZIP"
        "&format_options=charset%3AUTF-8",
    # Census 2023, Cuadro 15: people by barrio, Montevideo.
    "Cuadro_15_CAR_2023.xlsx":
        "https://www5.ine.gub.uy/documents/CENSO%202023/Tabulados/Personas/Cuadro_15_CAR_2023.xlsx",
}

# www5.ine.gub.uy serves an incomplete certificate chain, the same fault as Ghana's
# StatsBank and Ecuador's censoecuador.gob.ec. Nothing is trusted on the strength of the
# transport: the workbook's own national total is asserted below.
LAX_TLS = {"www5.ine.gub.uy"}

# INE's department numbering, verbatim from the `dpto` value labels in the ENHA 2006 person
# file. Montevideo first, then the other eighteen alphabetically. LAPOP's `prov` for Uruguay
# is 1400 plus these, and its own `prov_es` labels agree name for name, with 1407 (Flores)
# carrying an EMPTY label rather than a wrong one.
INE_DEPARTMENTS = {
    1: "Montevideo",     2: "Artigas",       3: "Canelones",   4: "Cerro Largo",
    5: "Colonia",        6: "Durazno",       7: "Flores",      8: "Florida",
    9: "Lavalleja",     10: "Maldonado",    11: "Paysandú",   12: "Río Negro",
    13: "Rivera",       14: "Rocha",        15: "Salto",      16: "San José",
    17: "Soriano",      18: "Tacuarembó",   19: "Treinta y Tres",
}

# LAPOP's `prov_es` labels for the 1400s, verbatim. Unaccented, and 1407 is blank because
# LAPOP never filled Flores in. Used as an independent witness on the same numbering.
LAPOP_PROVINCES = {
    1401: "Montevideo",  1402: "Artigas",    1403: "Canelones", 1404: "Cerro Largo",
    1405: "Colonia",     1406: "Durazno",    1407: "",          1408: "Florida",
    1409: "Lavalleja",   1410: "Maldonado",  1411: "Paysandu",  1412: "Rio Negro",
    1413: "Rivera",      1414: "Rocha",      1415: "Salto",     1416: "San Jose",
    1417: "Soriano",     1418: "Tacuarembo", 1419: "Treinta y Tres",
}

# How many of the nineteen a naive `UY{dpto:02d}` join happens to get right: the nine from
# Paysandú (11) onwards, plus nothing else. If this ever changes, the pcodes have been re-cut
# and somebody has to decide which numbering is now which.
CODE_JOIN_CORRECT = 9

EXPECTED_DEPARTMENTS = 19

# INE's estimated population at 30 June 2023, from the workbook's own Uruguay sheet. Asserted
# so a re-issued revision fails here instead of quietly moving every dot.
POP_2023 = 3_496_400

# The drawn universe: ENHA 2006 asked the religion question of people OVER 6 (the code-0
# not-applicable group is exactly ages 0 to 6, checked in sources/uy.py). Five-year bands
# cannot cut at 7, so `read_pop` SUBTRACTS two fifths of the 5-9 band, ages 5 and 6, and
# keeps the three fifths that are 7 to 9.
#
# It was 0.6 until 2026-09-15 (sources/uy.md §12.7), which took the 7-to-9-year-olds out and
# left the 5-and-6-year-olds in: every department's 7+ population was 1.3% low. The census 2023
# persons file is the witness: Montevideo's under-7 share is 6.95% there, 8.39% with 0.6 and
# 7.16% with 0.4.
BAND_5_9_OVER_6 = 0.4

# ---- Montevideo's barrios (module docstring, last section)
MONTEVIDEO_PCODE = "UY10"
MONTEVIDEO_INE = 1                  # `DEPARTAMENTO` in the census file, `dpto` in the ENHA
EXPECTED_BARRIOS = 62
BARRIO_ZIP = "im_barrios.zip"
BARRIO_SHP = "zon_v_sig_barrios"
CUADRO15 = "Cuadro_15_CAR_2023.xlsx"
CUADRO15_TOTAL = 1_302_721          # its `Total` row: the 62 barrios plus the street
STREET_LABEL = "Sin dato de barrio"
STREET_CODE = "9898"                # the same people in the persons file's `BARRIO85`
STREET_PEOPLE = 1_359
CENSUS_RAR = "personas_ext_07_2026.rar"
CENSUS_MEMBER = "personas_ext_07_2026.csv"
CENSUS_RAR_BYTES = 135_339_245
CENSUS_PEOPLE = 3_499_451           # `W` summed over the whole file; INE's published count
CENSUS_EXTRACT = os.path.join(RAW, "censo2023_mvd_barrio_edad.csv")
NATIONAL_KEY = "__URUGUAY__"
# A barrio's weighted census total may differ from Cuadro 15 by rounding only. Measured
# 2026-09-15: at most 0.6 of a person.
CUADRO15_TOLERANCE = 1.0
BARRIOS_OUT = os.path.join(OUT_DIR, "uy_barrios.gpkg")
BARRIO_LOOKUP = os.path.join(OUT_DIR, "uy_barrios_lookup.csv")
BARRIO_POP_OUT = os.path.join(OUT_DIR, "uy_barrios_pop.csv")


def fold(s):
    """Accent- and case-insensitive key for a department name."""
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return "".join(ch for ch in s.lower() if ch.isalnum())


def fetch():
    os.makedirs(RAW, exist_ok=True)
    for name, url in DOWNLOADS.items():
        dst = os.path.join(RAW, name)
        if os.path.exists(dst):
            print(f"  have {name} ({os.path.getsize(dst):,} bytes)")
            continue
        host = urllib.parse.urlsplit(url).hostname
        ctx = None
        if host in LAX_TLS:
            ctx = ssl.create_default_context()
            ctx.check_hostname = False
            ctx.verify_mode = ssl.CERT_NONE
        req = urllib.request.Request(url, headers=UA)
        with urllib.request.urlopen(req, timeout=900, context=ctx) as r:
            body = r.read()
        with open(dst + ".part", "wb") as f:
            f.write(body)
        os.replace(dst + ".part", dst)
        # §5a: a 200 is not a download. Both of these are zip containers (.xlsx is one).
        with open(dst, "rb") as f:
            magic = f.read(4)
        if magic[:2] != b"PK":
            raise SystemExit(f"{dst} is not a zip container — starts {magic!r}")
        print(f"  got  {name} ({os.path.getsize(dst):,} bytes)")


def read_pop(years=(2006, 2023)):
    """INE A.1.2: one sheet per department, `Ambos sexos` plus five-year bands, 1996-2023.

    Returns a frame indexed by department name with, for each requested year, `pop_<year>`
    and `pop7_<year>`, plus `T_<lo>_<hi>` / `T_90Plus` columns for the LAST year requested,
    shaped so `lapop._age_bands` could read them.

    The three sexes are stacked in one column, so only the FIRST block is read: `Ambos sexos`
    and the bands under it, stopping at `Hombres`. Reading past that would double the country.
    """
    import openpyxl

    path = os.path.join(RAW, POP_XLSX)
    if not os.path.exists(path):
        raise SystemExit(f"{path} missing — run with --fetch")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)

    sheets = [s for s in wb.sheetnames if s != "Uruguay"]
    if len(sheets) != EXPECTED_DEPARTMENTS:
        raise SystemExit(f"{len(sheets)} department sheets, expected {EXPECTED_DEPARTMENTS}: "
                         f"{sheets}")

    def read_sheet(name):
        ws = wb[name]
        year_col, block, out = None, False, {}
        for row in ws.iter_rows(values_only=True):
            label = "" if row[0] is None else str(row[0]).strip()
            if year_col is None:
                # THE HEADER IS NOT ALWAYS A NUMBER. Eighteen sheets write 2023 as an
                # integer and Paysandú writes it as the string `2023*`, a footnote marker
                # that nothing else in the workbook explains. Parsed rather than typed, so
                # that sheet does not silently lose the year the whole map is drawn on.
                yrs = {}
                for i, v in enumerate(row):
                    m = re.fullmatch(r"\s*(\d{4})\D*", str(v)) if v is not None else None
                    if m and 1900 < int(m.group(1)) < 2100:
                        yrs[int(m.group(1))] = i
                if len(yrs) > 20:
                    year_col = yrs
                continue
            if label in ("Hombres", "Mujeres"):
                break
            if label == "Ambos sexos":
                block = True
                out["total"] = row
                continue
            if not block or not label:
                continue
            m = re.fullmatch(r"(\d+)\s*-\s*(\d+)", label)
            if m:
                out[f"T_{int(m.group(1))}_{int(m.group(2))}"] = row
            elif label.lower().startswith("90"):
                out["T_90Plus"] = row
        if year_col is None or "total" not in out:
            raise SystemExit(f"sheet {name!r} has no year header row or no `Ambos sexos`")
        for y in years:
            if y not in year_col:
                raise SystemExit(f"sheet {name!r} has no column for {y}")
        return year_col, out

    rows = {}
    bands_last = {}
    for name in sheets + ["Uruguay"]:
        year_col, blk = read_sheet(name)
        rec = {}
        for y in years:
            i = year_col[y]
            total = float(blk["total"][i])
            b04 = float(blk["T_0_4"][i])
            b59 = float(blk["T_5_9"][i])
            band_sum = sum(float(v[i]) for k, v in blk.items() if k != "total")
            if abs(band_sum - total) > 1.0:
                raise SystemExit(f"{name} {y}: bands sum to {band_sum:,.0f} against a total "
                                 f"of {total:,.0f} — the sheet is not one block")
            rec[f"pop_{y}"] = total
            rec[f"pop7_{y}"] = total - b04 - BAND_5_9_OVER_6 * b59
        rows[name] = rec
        if name != "Uruguay":
            i = year_col[years[-1]]
            bands_last[name] = {k: float(v[i]) for k, v in blk.items() if k != "total"}
    wb.close()

    pop = pd.DataFrame(rows).T
    nat = pop.loc["Uruguay"]
    pop = pop.drop(index="Uruguay")
    pop = pop.join(pd.DataFrame(bands_last).T)

    for y in years:
        off = pop[f"pop_{y}"].sum() - nat[f"pop_{y}"]
        if abs(off) > 1.0:
            raise SystemExit(f"the {len(pop)} departments sum to {pop[f'pop_{y}'].sum():,.0f} "
                             f"in {y}, against the Uruguay sheet's {nat[f'pop_{y}']:,.0f}")
    if round(nat[f"pop_{years[-1]}"]) != POP_2023:
        raise SystemExit(f"INE's {years[-1]} national total is "
                         f"{nat[f'pop_{years[-1]}']:,.0f}, not the {POP_2023:,} this file "
                         "was written against — the revision has been re-issued")
    print(f"  INE A.1.2: {len(pop)} departments summing EXACTLY to {POP_2023:,} in "
          f"{years[-1]}, and to {nat[f'pop_{years[0]}']:,.0f} in {years[0]}")
    over6 = pop[f"pop7_{years[-1]}"].sum()
    print(f"    aged 7 and over in {years[-1]}: {over6:,.0f}, "
          f"{over6 / POP_2023:.2%} of the country")
    return pop


def check_code_join(by_name):
    """Report what a `UY{dpto:02d}` join would have done, and refuse to let it pass.

    Kept as an assertion rather than a comment because the failure it guards is invisible:
    a permutation preserves every total, and here it preserves nine of the pairings too.
    """
    right, wrong = [], []
    for num, name in sorted(INE_DEPARTMENTS.items()):
        naive = f"UY{num:02d}"
        actual = by_name[fold(name)]
        (right if naive == actual else wrong).append((num, name, naive, actual))
    print(f"\n  witness 2 — the code join is NOT used. `UY{{dpto:02d}}` would pair "
          f"{len(right)} of {EXPECTED_DEPARTMENTS} correctly and MISPAIR {len(wrong)}:")
    for num, name, naive, actual in wrong:
        other = next(n for n, p in by_name.items() if p == naive)
        print(f"      INE {num:>2} {name:<15} -> {naive}, which COD says is {other!r} "
              f"(the name says {actual})")
    if len(right) != CODE_JOIN_CORRECT:
        raise SystemExit(
            f"the code join now gets {len(right)} of {EXPECTED_DEPARTMENTS} right, not "
            f"{CODE_JOIN_CORRECT}. Either OCHA has re-cut the pcodes or INE has renumbered. "
            "STOP and decide which numbering is which before anything is drawn.")


# =======================================================================================
# Montevideo's barrios
# =======================================================================================

def barrio_geo_id(nro):
    """`UY10-B01` to `UY10-B62`: Montevideo's pcode, then INE's own barrio number."""
    return f"{MONTEVIDEO_PCODE}-B{int(nro):02d}"


def read_barrio_layer():
    """The Intendencia's 62 barrios, in EPSG:32721 as served, with INE's number asserted."""
    zpath = os.path.join(RAW, BARRIO_ZIP)
    if not os.path.exists(zpath):
        raise SystemExit(f"{zpath} missing — run with --fetch")
    out_dir = os.path.join(SHP_DIR, "barrios")
    os.makedirs(out_dir, exist_ok=True)
    with zipfile.ZipFile(zpath) as z:
        members = [n for n in z.namelist() if n.startswith(BARRIO_SHP + ".")]
        if BARRIO_SHP + ".shp" not in members:
            raise SystemExit(f"{zpath} holds no {BARRIO_SHP}.shp: {z.namelist()}")
        for n in members:
            z.extract(n, out_dir)
    b = gpd.read_file(os.path.join(out_dir, BARRIO_SHP + ".shp"), engine="fiona",
                      encoding="utf-8")
    if len(b) != EXPECTED_BARRIOS:
        raise SystemExit(f"{len(b)} barrio polygons, expected {EXPECTED_BARRIOS}")
    if b.crs is None or b.crs.to_epsg() != 32721:
        raise SystemExit(f"the barrio layer is in {b.crs}, not EPSG:32721 as served in 2026")
    b["nro"] = b["nrobarrio"].astype(int)
    if sorted(b["nro"]) != list(range(1, EXPECTED_BARRIOS + 1)):
        raise SystemExit(f"`nrobarrio` is not 1..{EXPECTED_BARRIOS} once each")
    b["im_name"] = b["barrio"].astype(str).str.strip()
    print(f"\n  barrios: {len(b)} polygons from the Intendencia, INE numbers 1-{len(b)}, "
          f"{b.area.sum() / 1e6:,.1f} km2")
    return b


def read_cuadro15():
    """{barrio name: people} from census 2023 Cuadro 15, with its own arithmetic asserted."""
    import openpyxl

    path = os.path.join(RAW, CUADRO15)
    if not os.path.exists(path):
        raise SystemExit(f"{path} missing — run with --fetch")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    rows, total, street = {}, None, None
    for row in wb.worksheets[0].iter_rows(values_only=True):
        if not row or row[0] is None or len(row) < 2 or not isinstance(row[1], (int, float)):
            continue
        name = str(row[0]).strip()
        if name == "Total":
            total = row[1]
        elif name == STREET_LABEL:
            street = row[1]
        else:
            rows[name] = float(row[1])
    wb.close()
    if len(rows) != EXPECTED_BARRIOS:
        raise SystemExit(f"Cuadro 15 has {len(rows)} barrio rows, expected {EXPECTED_BARRIOS}")
    if total != CUADRO15_TOTAL or street != STREET_PEOPLE:
        raise SystemExit(f"Cuadro 15's Total is {total} and its street row {street}, not "
                         f"{CUADRO15_TOTAL:,} and {STREET_PEOPLE:,} — the table was re-issued")
    if round(sum(rows.values()) + street) != total:
        raise SystemExit("Cuadro 15's barrios plus its street row do not make its Total")
    print(f"  Cuadro 15: {len(rows)} barrios and {street:,} people without one, "
          f"{total:,} in all")
    return rows


def census_extract(refresh=False):
    """Montevideo's census 2023 people by `BARRIO85`, all ages and under 7, weighted by `W`.

    The persons file is a 135 MB RAR holding one 1.9 GB CSV, so it is streamed through bsdtar
    rather than unpacked, and the result is kept as a small CSV beside it. To get the file:
    `www4.ine.gub.uy/Anda5/index.php/catalog/781/get-microdata`, POST `accept=Aceptar` with
    the session cookie (the ENHA's terms, word for word, `sources/uy.md` §2), then
    `.../catalog/781/download/1503` is `personas_ext_07_2026.rar`.

    Under 7 is `PERNA01 <= 6`, the ENHA's own cut (`sources/uy.py` `NOT_ASKED_MAX_AGE`).
    """
    if os.path.exists(CENSUS_EXTRACT) and not refresh:
        ex = pd.read_csv(CENSUS_EXTRACT, dtype={"barrio85": str}, encoding="utf-8")
        print(f"  census 2023 persons: read the extract {CENSUS_EXTRACT} "
              "(--census re-streams it)")
    else:
        rar = os.path.join(RAW, CENSUS_RAR)
        if not os.path.exists(rar):
            raise SystemExit(f"{rar} missing. www4.ine.gub.uy/Anda5/index.php/catalog/781 -> "
                             "Obtener Microdatos -> Aceptar -> download/1503; see "
                             "census_extract()'s docstring")
        if os.path.getsize(rar) != CENSUS_RAR_BYTES:
            raise SystemExit(f"{rar} is {os.path.getsize(rar):,} bytes, not "
                             f"{CENSUS_RAR_BYTES:,}: a partial download or a new release")
        bsdtar = shutil.which("bsdtar")
        if not bsdtar:
            raise SystemExit("bsdtar is not on PATH; anaconda ships it (Library/bin)")
        proc = subprocess.Popen([bsdtar, "-xOf", rar, CENSUS_MEMBER], stdout=subprocess.PIPE)
        # the header's names are quoted ("ID_CENSO"), the values are not
        header = [c.strip().strip('"')
                  for c in proc.stdout.readline().decode("latin-1").rstrip("\r\n").split(",")]
        need = ["DEPARTAMENTO", "BARRIO85", "PERNA01", "W"]
        missing = [c for c in need if c not in header]
        if missing:
            proc.kill()
            raise SystemExit(f"the persons file has no {missing}; its header is {header}")
        parts, n_rows, w_all = [], 0, 0.0
        for ch in pd.read_csv(proc.stdout, header=None, names=header, usecols=need, dtype=str,
                              chunksize=500_000, encoding="latin-1"):
            w = pd.to_numeric(ch["W"], errors="raise")
            age = pd.to_numeric(ch["PERNA01"], errors="raise")
            dep = pd.to_numeric(ch["DEPARTAMENTO"], errors="raise")
            n_rows += len(ch)
            w_all += float(w.sum())
            sel = dep == MONTEVIDEO_INE
            t = pd.DataFrame({"barrio85": ch.loc[sel, "BARRIO85"].str.strip(),
                              "w": w[sel], "u7": w[sel].where(age[sel] <= 6, 0.0)})
            parts.append(t.groupby("barrio85").agg(n=("w", "size"), w_total=("w", "sum"),
                                                   w_under7=("u7", "sum")))
        if proc.wait() != 0:
            raise SystemExit(f"bsdtar exited {proc.returncode} while streaming {rar}")
        ex = pd.concat(parts).groupby(level=0).sum().reset_index()
        ex = pd.concat([ex, pd.DataFrame([{"barrio85": NATIONAL_KEY, "n": n_rows,
                                           "w_total": w_all, "w_under7": float("nan")}])],
                       ignore_index=True)
        ex.to_csv(CENSUS_EXTRACT + ".part", index=False, encoding="utf-8")
        os.replace(CENSUS_EXTRACT + ".part", CENSUS_EXTRACT)
        print(f"  census 2023 persons: streamed {n_rows:,} rows out of {CENSUS_RAR}, wrote "
              f"{CENSUS_EXTRACT}")
    nat = ex.loc[ex["barrio85"] == NATIONAL_KEY, "w_total"]
    if len(nat) != 1 or round(float(nat.iloc[0])) != CENSUS_PEOPLE:
        raise SystemExit(f"the persons file weights to {nat.tolist()}, not the census's "
                         f"{CENSUS_PEOPLE:,}")
    ex = ex[ex["barrio85"] != NATIONAL_KEY].copy()
    street = ex.loc[ex["barrio85"] == STREET_CODE, "w_total"]
    if len(street) != 1 or round(float(street.iloc[0])) != STREET_PEOPLE:
        raise SystemExit(f"`BARRIO85` {STREET_CODE} weights to {street.tolist()}, not "
                         f"Cuadro 15's {STREET_PEOPLE:,} people without a barrio")
    print(f"    weighted to {CENSUS_PEOPLE:,} nationally; Montevideo {ex['w_total'].sum():,.0f} "
          f"in {len(ex) - 1} barrios and {STREET_PEOPLE:,} ({STREET_CODE}) on the street")
    return ex[ex["barrio85"] != STREET_CODE].reset_index(drop=True)


def match_names(left, right, what):
    """Pair each name in `left` with its most similar in `right`; refuse anything but 1:1.

    For two spellings of one list that share no code (Cuadro 15 abbreviates `PQUE. BATLLE, V.
    DOLORES`, the Intendencia writes `PQUE BATLLE VILLA DOLORES`). The narrowest margins
    between the best and second-best candidate are printed, since those are where a wrong
    pairing would hide.
    """
    out, margins = {}, []
    for a in left:
        sims = sorted(((difflib.SequenceMatcher(None, fold(a), fold(b)).ratio(), b)
                       for b in right), reverse=True)
        out[a] = sims[0][1]
        margins.append((sims[0][0] - sims[1][0], sims[0][0], a, sims[0][1], sims[1][1]))
    if len(set(out.values())) != len(left) or set(out.values()) != set(right):
        dup = sorted(v for v in out.values() if list(out.values()).count(v) > 1)
        raise SystemExit(f"{what}: the best-match pairing is not one to one (repeated: {dup})")
    margins.sort()
    print(f"  {what}: all {len(left)} names pair one to one by best match; narrowest margins:")
    for m, s, a, b, c in margins[:3]:
        print(f"      {a!r} -> {b!r} at {s:.2f}, next {c!r} {m:.2f} behind")
    return out


def build_barrios(dept_pop, refresh_census=False):
    """Polygons, lookup and 2023 populations for Montevideo's barrios. See the docstring."""
    b = read_barrio_layer()
    c15 = read_cuadro15()
    ex = census_extract(refresh_census)

    # the persons file writes Cuadro 15's own names, so that join is exact
    by_fold = {fold(n): n for n in c15}
    ex["c15_name"] = ex["barrio85"].map(lambda s: by_fold.get(fold(s)))
    if ex["c15_name"].isna().any() or ex["c15_name"].nunique() != EXPECTED_BARRIOS:
        raise SystemExit(f"census `BARRIO85` names with no Cuadro 15 row: "
                         f"{sorted(ex.loc[ex['c15_name'].isna(), 'barrio85'])}")
    ex["c15"] = ex["c15_name"].map(c15)
    worst = (ex["w_total"] - ex["c15"]).abs().max()
    if worst > CUADRO15_TOLERANCE:
        raise SystemExit(f"a barrio's weighted census total is {worst:.1f} people off Cuadro 15")
    print(f"  census persons file, weighted, against Cuadro 15: every barrio within "
          f"{worst:.2f} of a person")

    to_layer = match_names(list(c15), list(b["im_name"]), "Cuadro 15 -> Intendencia layer")
    nro_by_im = dict(zip(b["im_name"], b["nro"]))
    ex["nro"] = ex["c15_name"].map(lambda n: nro_by_im[to_layer[n]])

    mvd = dept_pop.loc[INE_DEPARTMENTS[MONTEVIDEO_INE]]
    ex["w7"] = ex["w_total"] - ex["w_under7"]
    ex["pop_2023"] = mvd["pop_2023"] * ex["w_total"] / ex["w_total"].sum()
    ex["pop7_2023"] = mvd["pop7_2023"] * ex["w7"] / ex["w7"].sum()
    ex["u7_share"] = ex["w_under7"] / ex["w_total"]
    lo, hi = ex.loc[ex["u7_share"].idxmin()], ex.loc[ex["u7_share"].idxmax()]
    print(f"  under 7 by barrio (census 2023, weighted): {lo['u7_share']:.2%} {lo['c15_name']} "
          f"to {hi['u7_share']:.2%} {hi['c15_name']}; Montevideo "
          f"{ex['w_under7'].sum() / ex['w_total'].sum():.2%} in the census against INE's "
          f"band interpolation {1 - mvd['pop7_2023'] / mvd['pop_2023']:.2%}")
    flat = mvd["pop7_2023"] / mvd["pop_2023"]
    move = ex["pop_2023"] * flat / ex["pop7_2023"] - 1
    print(f"    a flat city-wide 7+ ratio would have drawn barrios {move.min():+.1%} to "
          f"{move.max():+.1%} of their own 7+ population")

    ex["geo_id"] = ex["nro"].map(barrio_geo_id)
    b["geo_id"] = b["nro"].map(barrio_geo_id)
    g = b.merge(ex[["geo_id", "c15_name", "pop_2023"]], on="geo_id", how="left")
    if g["pop_2023"].isna().any():
        raise SystemExit("a barrio polygon came out of the population join with nobody")
    g["unit"] = g["geo_id"]
    g["name"] = g["c15_name"]
    g["pop"] = g["pop_2023"].round().astype("int64")
    g = g.to_crs(4326)
    g[["unit", "geo_id", "nro", "name", "im_name", "pop", "geometry"]].to_file(
        BARRIOS_OUT + ".part.gpkg", layer="barrios", driver="GPKG")
    os.replace(BARRIOS_OUT + ".part.gpkg", BARRIOS_OUT)
    print(f"\nwrote {BARRIOS_OUT} ({len(g)} polygons)")

    lut = g[["geo_id", "unit", "nro", "name", "im_name"]].sort_values("geo_id")
    lut.to_csv(BARRIO_LOOKUP, index=False, encoding="utf-8")
    print(f"wrote {BARRIO_LOOKUP} ({len(lut)} rows)")
    pop = ex.rename(columns={"c15_name": "name", "w_total": "census_2023",
                             "w7": "census_2023_7plus"})
    pop = pop[["geo_id", "nro", "name", "census_2023", "census_2023_7plus", "u7_share",
               "pop_2023", "pop7_2023"]].sort_values("geo_id")
    pop.to_csv(BARRIO_POP_OUT, index=False, encoding="utf-8")
    print(f"wrote {BARRIO_POP_OUT} ({len(pop)} barrios, {pop['pop_2023'].sum():,.0f} people "
          f"of whom {pop['pop7_2023'].sum():,.0f} are 7 or over)")


def main():
    if "--fetch" in sys.argv:
        fetch()

    zpath = os.path.join(RAW, "ury_adm_2020_shp.zip")
    if not os.path.exists(zpath):
        raise SystemExit(f"{zpath} missing — run with --fetch")
    os.makedirs(SHP_DIR, exist_ok=True)
    with zipfile.ZipFile(zpath) as z:
        members = [n for n in z.namelist() if "adm1" in n and not n.endswith(".xml")]
        if not members:
            raise SystemExit("no adm1 members in the COD-AB zip — it has been re-cut")
        for n in members:
            z.extract(n, SHP_DIR)
    print(f"extracted {len(members)} adm1 members")

    shp = os.path.join(SHP_DIR, "ury_admbnda_adm1_2020.shp")
    # `engine="fiona"` on §12's Chile rule: pyogrio is geopandas' default when installed and
    # is the engine that has silently returned zero features. The count is asserted either way.
    g = gpd.read_file(shp, engine="fiona")
    if len(g) != EXPECTED_DEPARTMENTS:
        raise SystemExit(f"{len(g)} ADM1 features, expected {EXPECTED_DEPARTMENTS} — COD has "
                         "re-cut Uruguay")
    if g.crs is None or g.crs.to_epsg() != 4326:
        raise SystemExit(f"unexpected CRS {g.crs}")
    print(f"read {shp}: {len(g)} departments, {g.crs}")

    g["pcode"] = g["ADM1_PCODE"].astype(str).str.strip()
    g["name"] = g["ADM1_ES"].astype(str).str.strip()
    by_name = dict(zip(g["name"].map(fold), g["pcode"]))
    if len(by_name) != EXPECTED_DEPARTMENTS:
        raise SystemExit("COD's department names are not unique — the name join is unsafe")

    # ---- witness 1: every INE department name is a COD name, with no aliases ----
    missing = [n for n in INE_DEPARTMENTS.values() if fold(n) not in by_name]
    if missing:
        raise SystemExit(f"INE names with no polygon: {missing} — the name join FAILED")
    spare = sorted(n for n in g["name"]
                   if fold(n) not in {fold(v) for v in INE_DEPARTMENTS.values()})
    if spare:
        raise SystemExit(f"polygons with no INE department: {spare}")
    print(f"  witness 1 — all {EXPECTED_DEPARTMENTS} INE names match a COD name exactly, "
          "no aliases needed")

    check_code_join(by_name)

    # ---- witness 3: LAPOP's own labels, on the same numbering ----
    bad = []
    for code, label in sorted(LAPOP_PROVINCES.items()):
        if not label:
            continue
        if fold(label) != fold(INE_DEPARTMENTS[code - 1400]):
            bad.append((code, label, INE_DEPARTMENTS[code - 1400]))
    blank = sorted(c for c, v in LAPOP_PROVINCES.items() if not v)
    if bad:
        for code, label, ine in bad:
            print(f"      LAPOP {code} {label!r} against INE {ine!r}")
        raise SystemExit("LAPOP's province labels no longer agree with INE's numbering")
    print(f"  witness 3 — LAPOP's `prov_es` labels agree with INE's numbering on "
          f"{len(LAPOP_PROVINCES) - len(blank)} of {EXPECTED_DEPARTMENTS}; "
          f"{blank} carries an EMPTY label, which is Flores")

    # ---- witness 4: INE's own population workbook, joined on the name ----
    pop = read_pop()
    pop_missing = [n for n in pop.index if fold(n) not in by_name]
    if pop_missing:
        raise SystemExit(f"population sheets with no polygon: {pop_missing}")
    pop["pcode"] = [by_name[fold(n)] for n in pop.index]
    if set(pop["pcode"]) != set(g["pcode"]):
        raise SystemExit("the population workbook and COD-AB do not cover the same pcodes")
    g = g.merge(pop.reset_index(names="ine_name")[["pcode", "pop_2023"]], on="pcode",
                how="left")
    if g["pop_2023"].isna().any():
        raise SystemExit("a department came out of the population join with no total")
    g["pop"] = g["pop_2023"].round().astype("int64")
    print(f"  witness 4 — INE A.1.2 joins on the name, all {len(g)} departments: "
          f"{g['pop'].sum():,} people")

    # Montevideo is the smallest department by area and by far the largest by population, and
    # Durazno the sparsest; a permuted population join is what this catches.
    g["area_km2"] = g.to_crs(3857).geometry.area / 1e6
    g["density"] = g["pop"] / g["area_km2"]
    hi = g.loc[g["density"].idxmax(), "name"]
    biggest = g.loc[g["pop"].idxmax(), "name"]
    print(f"    densest {hi!r}, sparsest {g.loc[g['density'].idxmin(), 'name']!r}; "
          f"largest population {biggest!r}, smallest {g.loc[g['pop'].idxmin(), 'name']!r}")
    if fold(hi) != fold("Montevideo") or fold(biggest) != fold("Montevideo"):
        raise SystemExit("Montevideo is not both the densest and the largest department — "
                         "the population join is permuted")

    g["unit"] = g["pcode"]
    # geo_id IS the pcode. INE's own department numbering never leaves sources/uy.py, so
    # there is one department key in data/ and nobody downstream can pick the wrong one.
    g["geo_id"] = g["pcode"]

    os.makedirs(OUT_DIR, exist_ok=True)
    out = g[["unit", "name", "pcode", "geo_id", "pop", "geometry"]]
    out.to_file(OUT, layer="departamentos", driver="GPKG")
    print(f"\nwrote {OUT} ({len(out)} polygons)")

    num_by_pcode = {by_name[fold(n)]: k for k, n in INE_DEPARTMENTS.items()}
    order = sorted(g["pcode"])
    lut = pd.DataFrame({
        "geo_id": order,
        "unit": order,
        "name": [g.loc[g["pcode"] == p, "name"].iloc[0] for p in order],
        "ine_dpto": [num_by_pcode[p] for p in order],
        "lapop_prov": [1400 + num_by_pcode[p] for p in order],
    })
    lut.to_csv(LOOKUP, index=False, encoding="utf-8")
    print(f"wrote {LOOKUP} ({len(lut)} rows)")

    pop = pop.reset_index(names="ine_name")
    pop["geo_id"] = pop["pcode"]
    pop["ine_dpto"] = pop["geo_id"].map(num_by_pcode)
    band_cols = sorted((c for c in pop.columns if c.startswith("T_")),
                       key=lambda c: int(re.findall(r"\d+", c)[0]))
    cols = ["geo_id", "ine_name", "ine_dpto", "pop_2023", "pop7_2023",
            "pop_2006", "pop7_2006"] + band_cols
    pop.sort_values("geo_id")[cols].to_csv(POP_OUT, index=False, encoding="utf-8")
    print(f"wrote {POP_OUT} ({len(pop)} departments, {len(band_cols)} age bands, "
          f"{pop['pop_2023'].sum():,.0f} people in 2023 of whom "
          f"{pop['pop7_2023'].sum():,.0f} are 7 or over)")

    build_barrios(pop.set_index("ine_name"), refresh_census="--census" in sys.argv)


if __name__ == "__main__":
    main()
