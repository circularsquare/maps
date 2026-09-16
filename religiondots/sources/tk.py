"""Tokelau - 2016 Tokelau Census of Population and Dwellings, religious affiliation by atoll.

Reads (or fetches) the census workbooks and the profile report into data/normalized/tk.csv.
`sources/tk.md` is the write-up; `sources/tk_geo.py` builds the three atolls and the placement grid;
`taxonomy/tk2016.py` is the mapping.

## THE TABLE

Tokelau National Statistics Office and Statistics New Zealand, *2016 Tokelau Census of Population
and Dwellings: Tables about social profile* (published 1 February 2017), sheet **Table 5.8,
"Religious affiliation, by atoll of usual residence"**, 2011 and 2016 side by side, in persons,
"for usually resident population present in Tokelau on census night". Seven rows, one answer per
person: every atoll column sums to its total.

## THE UNIVERSE IS THE RESIDENTS PRESENT, NOT THE OFFICIAL COUNT

Tokelau's headline figure is the de jure usually resident population, 1,499 in 2016, which sets
each atoll's seats in the General Fono. It is the 1,197 usual residents present on census night
plus 254 usual residents overseas and 48 Tokelau Public Service employees and their families in
Apia (profile report, printed p.15). Absentees were described by the head of their household and
the Apia staff answered a short form, so neither has a religion. The map draws the 1,197 and puts
the other 302 in `gap` beside the 9 whose religion was not stated.

## THE FORM

The profile report (printed p.28) says the questionnaire named three churches, Congregational
Christian, Roman Catholic and Presbyterian, with "other, please specify". No religion is a
write-in only (code 8 of the classification on printed pp.76-77). The workbook's contents sheet
says most `not stated` answers come from people whose age was imputed, so the questions after it
went unanswered.

## THE CHECKS

    all three files pinned (size and digest)
    Table 5.8 parsed by label equals the transcription below, both years; captions and universe
    each atoll's rows sum to its printed total, and the atolls to each row's total, both years
    the demography workbook: de jure residents (1.3.1) minus absentees (1.3.2) equals Table 5.8's
        atoll totals, both years; 1,499 de jure in 2016, 48 of them in Samoa
    the profile report: Table 4.1 prints the same present and absent counts; printed p.15 states
        1,499 = 1,197 + 302; printed p.28 names the three churches on the form; pp.76-77 carry
        the no-religion and not-stated codes
    UNSD table 28: its 2011 row equals Table 5.8's 2011 column; its 2016 row differs in three
        cells by four people in all (pinned), so the office's table is the source

Usage:
    python sources/tk.py --fetch    three files from tokelau.org.nz (2.6 MB in all)
    python sources/tk.py            normalise from data/raw/tk/
"""

import csv
import os
import re
import sys
import unicodedata

os.environ.setdefault("OMP_NUM_THREADS", "6")

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
RAW = os.path.join(ROOT, "data", "raw", "tk")
OUT = os.path.join(ROOT, "data", "normalized", "tk.csv")
sys.path.insert(0, HERE)
sys.path.insert(0, os.path.join(ROOT, "tools"))

from fetch_checks import FetchCheckError, check_body, digest   # noqa: E402

SOURCE_ID = "tk_census2016_t5_8"
YEAR = 2016
BASIS = "self_id"
COLUMNS = ["geo_id", "geo_level", "geo_name", "source_category", "count",
           "basis", "year", "source_id", "note"]

BASE = "https://www.tokelau.org.nz/site/tokelau/files/TokelauNSO/2016Census/"
SOCIAL = os.path.join(RAW, "tk2016_tables_social_profile.xlsx")
DEMOG = os.path.join(RAW, "tk2016_tables_demography.xlsx")
PROFILE = os.path.join(RAW, "tk2016_profile_report.pdf")
DOWNLOADS = {
    SOCIAL: (BASE + "2016%20Tokelau%20Census%20of%20Population%20and%20Dwellings%20-%20"
                    "Tables%20about%20social%20profile.xlsx",
             "zip", 177_914, "IAXB35RC7XF25GUELMAR5QHJAHOIDRPH"),
    DEMOG: (BASE + "2016%20Tokelau%20Census%20of%20Population%20and%20Dwellings%20-%20"
                   "Tables%20about%20demography.xlsx",
            "zip", 75_125, "WJK4TS75ILLVUNRZAFBJIGNHTPIXFVPL"),
    PROFILE: (BASE + "profile-tokelau-2016-census-final-to-print28jun17jj.pdf",
              "pdf", 2_395_295, "2IZBMBTJFWAYOYOC4NWCX5THXAU4GTRM"),
}
PROFILE_PAGES = 95
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) religiondots/1.0"}

ATOLLS = ["Atafu", "Fakaofo", "Nukunonu"]
CATS = ["Congregational Christian", "Presbyterian", "Roman Catholic", "Other Christian",
        "Spiritualism and New Age religions", "No religion", "Not stated"]
# The sheet writes `Congregational Christian ` with a trailing space; labels are stripped, and the
# check asserts that this is the only label that needed it.
PADDED = {"Congregational Christian"}

# Table 5.8, transcribed from the workbook: (Atafu, Fakaofo, Nukunonu) per row.
T58 = {
    2011: {"Congregational Christian": (345, 306, 14), "Presbyterian": (5, 11, 5),
           "Roman Catholic": (13, 115, 290), "Other Christian": (21, 11, 0),
           "Spiritualism and New Age religions": (0, 1, 0), "No religion": (0, 0, 0),
           "Not stated": (1, 5, 0)},
    2016: {"Congregational Christian": (318, 250, 35), "Presbyterian": (54, 8, 9),
           "Roman Catholic": (18, 130, 315), "Other Christian": (15, 11, 24),
           "Spiritualism and New Age religions": (0, 0, 0), "No religion": (1, 0, 0),
           "Not stated": (7, 0, 2)},
}
T58_TOTAL = {2011: (385, 449, 309), 2016: (413, 399, 385)}
PRESENT = {2011: 1_143, 2016: 1_197}

# Demography workbook, Total columns: 1.3.1 de jure (Atafu, Fakaofo, Nukunonu, Samoa) and
# 1.3.2 absentees (Atafu, Fakaofo, Nukunonu).
DEJURE = {2011: (482, 490, 397, 42), 2016: (519, 484, 448, 48)}
ABSENT = {2011: (97, 41, 88), 2016: (106, 85, 63)}
DEJURE_2016 = 1_499

# UNSD's names for the office's rows, and the three cells its 2016 row does not share.
UNSD_NAME = {"Congregational Christian": "Christian Congregational", "Presbyterian": "Presbyterian",
             "Roman Catholic": "Roman Catholic", "Other Christian": "Other Christians",
             "Spiritualism and New Age religions": "New Age", "No religion": "No Religion",
             "Not stated": "Not Stated"}
UNSD_2016_DIFF = {"Roman Catholic": (463, 460), "Other Christian": (50, 54), "No religion": (1, 0)}


def fetch():
    import urllib.request

    os.makedirs(RAW, exist_ok=True)
    for dst, (url, kind, size, dig) in DOWNLOADS.items():
        if os.path.exists(dst) and os.path.getsize(dst) == size:
            print(f"  have {os.path.basename(dst)} ({size:,} bytes)")
            continue
        with urllib.request.urlopen(urllib.request.Request(url, headers=UA), timeout=600) as r:
            body = r.read()
        try:
            check_body(body, kind, where=url, pin_size=size, pin_digest=dig)
        except FetchCheckError as e:
            raise SystemExit(f"{os.path.basename(dst)}: {e}")
        with open(dst + ".part", "wb") as fh:
            fh.write(body)
        os.replace(dst + ".part", dst)
        print(f"  got  {os.path.basename(dst)} ({len(body):,} bytes)")


def count(cell, where):
    """A table cell that must be a whole number of people; anything else stops."""
    if isinstance(cell, bool) or cell is None:
        raise SystemExit(f"{where}: {cell!r} is not a count")
    if isinstance(cell, int):
        return cell
    if isinstance(cell, float) and cell.is_integer():
        return int(cell)
    if isinstance(cell, str) and re.fullmatch(r"\d{1,3}(,\d{3})*", cell.strip()):
        return int(cell.strip().replace(",", ""))
    raise SystemExit(f"{where}: {cell!r} is not a count")


def sheet(path, name):
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if name not in wb.sheetnames:
        raise SystemExit(f"{os.path.basename(path)}: no sheet {name!r} in {wb.sheetnames}")
    rows = [list(r) for r in wb[name].iter_rows(values_only=True)]
    wb.close()
    return rows


def _text(rows):
    return " ".join(str(c) for r in rows for c in r if c is not None)


def read_t58():
    """({year: {label: (a, f, n)}}, {year: totals}, padded labels, caption text) off the sheet."""
    rows = sheet(SOCIAL, "Table 5.8")
    years = next((r for r in rows if 2011 in r and 2016 in r), None)
    if years is None:
        raise SystemExit("Table 5.8: no row carrying both 2011 and 2016")
    col = {y: years.index(y) for y in (2011, 2016)}
    head = next((i for i, r in enumerate(rows) if r.count("Atafu") == 2), None)
    if head is None:
        raise SystemExit("Table 5.8: no header row naming the atolls twice")
    for y, c in col.items():
        got = [str(x).strip() for x in rows[head][c:c + 4]]
        if got != ATOLLS + ["Total"]:
            raise SystemExit(f"Table 5.8 {y}: header {got}, expected {ATOLLS + ['Total']}")
    table = {2011: {}, 2016: {}}
    totals, padded = {}, set()
    for r in rows[head + 1:]:
        if not r or not isinstance(r[0], str):
            continue
        raw = r[0]
        label = unicodedata.normalize("NFC", raw.strip())
        if label.startswith("Source"):
            break
        if raw != raw.strip():
            padded.add(label)
        for y, c in col.items():
            vals = tuple(count(r[c + i], f"Table 5.8 {y} {label!r} col {i}") for i in range(4))
            if label == "Total":
                totals[y] = vals
            else:
                if label in table[y]:
                    raise SystemExit(f"Table 5.8: {label!r} twice")
                table[y][label] = vals
    return table, totals, padded, _text(rows[:head])


def read_demography(name):
    """{year: {row: total}} from a 1.3.x sheet (Male, Female, Total per year), and its captions."""
    rows = sheet(DEMOG, name)
    years = next((r for r in rows if 2011 in r and 2016 in r), None)
    if years is None:
        raise SystemExit(f"demography {name}: no year row")
    col = {y: years.index(y) + 2 for y in (2011, 2016)}      # Male, Female, Total
    head = next((i for i, r in enumerate(rows) if "Male" in r), None)
    for y, c in col.items():
        if rows[head][c] != "Total":
            raise SystemExit(f"demography {name} {y}: column {c} is {rows[head][c]!r}, not Total")
    out = {2011: {}, 2016: {}}
    for r in rows[head + 1:]:
        if not r or not isinstance(r[0], str) or r[0].startswith("Source"):
            continue
        for y, c in col.items():
            out[y][r[0].strip()] = count(r[c], f"demography {name} {y} {r[0]!r}")
    return out, _text(rows[:head])


def _flat(s):
    s = unicodedata.normalize("NFKC", s)
    s = s.replace("–", "-").replace("—", "-").replace("‘", "'").replace("’", "'")
    return " ".join(s.split())


def check():
    import fitz
    import oracle

    ok = True

    def say(good, msg):
        nonlocal ok
        ok &= bool(good)
        print(f"  {'OK ' if good else 'BAD'} {msg}")

    print("Tokelau - 2016 census, Table 5.8\n")

    # 0. the files
    for path, (_url, _kind, size, dig) in DOWNLOADS.items():
        with open(path, "rb") as fh:
            body = fh.read()
        say(len(body) == size and digest(body) == dig,
            f"{os.path.basename(path)} is the pinned {size:,} bytes, digest {dig}")

    # 1. the page and the transcription are one table
    table, totals, padded, caption = read_t58()        # each row (Atafu, Fakaofo, Nukunonu, Total)
    atolls_only = {y: {k: v[:3] for k, v in table[y].items()} for y in table}
    say(atolls_only == T58, f"Table 5.8 parsed by label equals the transcription: {len(T58[2016])} "
                            "rows x 3 atolls, 2011 and 2016")
    if atolls_only != T58:
        for y in T58:
            for k in set(T58[y]) | set(atolls_only[y]):
                if T58[y].get(k) != atolls_only[y].get(k):
                    print(f"        {y} {k!r}: sheet {atolls_only[y].get(k)} transcription {T58[y].get(k)}")
    say(padded == PADDED, f"labels with stray whitespace in the sheet: {sorted(padded)}")
    say("For usually resident population present in Tokelau on census night" in caption
        and "By atoll of usual residence" in caption,
        "the caption names the atoll of usual residence and the residents present on census night")

    # 2. closure
    for y in (2011, 2016):
        say(all(totals[y][i] == sum(table[y][k][i] for k in CATS) for i in range(3))
            and totals[y][:3] == T58_TOTAL[y],
            f"{y}: each atoll's rows sum to its printed total {T58_TOTAL[y]}")
        say(all(table[y][k][3] == sum(table[y][k][:3]) for k in CATS)
            and totals[y][3] == PRESENT[y] == sum(T58_TOTAL[y]),
            f"{y}: the atolls sum to every row's Total column, {PRESENT[y]:,} in all")

    # 3. the demography workbook: present = de jure - absent, per atoll
    dj, dj_cap = read_demography("1.3.1")
    ab, ab_cap = read_demography("1.3.2")
    say("For de jure usually resident population" in dj_cap
        and "For usual residents absent from Tokelau on census night" in ab_cap,
        "Tables 1.3.1 and 1.3.2 are the de jure population and the absentees")
    for y in (2011, 2016):
        got_dj = tuple(dj[y][a] for a in ATOLLS + ["Samoa"])
        got_ab = tuple(ab[y][a] for a in ATOLLS)
        say(got_dj == DEJURE[y] and got_ab == ABSENT[y],
            f"{y}: de jure {got_dj} and absentees {got_ab} as transcribed")
        say(tuple(d - a for d, a in zip(got_dj, got_ab)) == T58_TOTAL[y],
            f"{y}: de jure minus absentees equals Table 5.8's atoll totals {T58_TOTAL[y]}")
    say(dj[2016]["Total"] == DEJURE_2016 and DEJURE_2016 - PRESENT[2016] == sum(ABSENT[2016]) + DEJURE[2016][3],
        f"2016: {DEJURE_2016:,} de jure = {PRESENT[2016]:,} present + {sum(ABSENT[2016])} absent + "
        f"{DEJURE[2016][3]} in Samoa")

    # 4. the profile report
    doc = fitz.open(PROFILE)
    say(doc.page_count == PROFILE_PAGES, f"the profile report is {doc.page_count} pages "
                                         f"(expected {PROFILE_PAGES})")
    page = {i: _flat(doc.load_page(i).get_text()) for i in (14, 15, 27, 75, 76)}
    say("The de jure usually resident population in 2016 was 1,499" in page[14]
        and "1,197 usual residents who were present in Tokelau on census night, and 302 usual "
            "residents who were overseas" in page[14]
        and "48 Tokelauan TPS employees and their immediate families based in Apia, and 254" in page[14],
        "printed p.15: 1,499 de jure = 1,197 present + 302 overseas (48 in Apia, 254 elsewhere)")
    say("Atafu 198 53 215 53 413 106 Fakaofo 185 41 214 44 399 85 Nukunonu 218 33 167 30 385 63 "
        "Samoa 0 28 0 20 0 48 Total 601 155 596 147 1,197 302" in page[15],
        "printed p.16, Table 4.1: present 413 / 399 / 385 and absent 106 / 85 / 63 + 48 in Samoa")
    say("three major denominations as options - Congregational Christian, Roman Catholic, and "
        "Presbyterian - along with an 'other, please specify', option" in page[27],
        "printed p.28: the form named Congregational Christian, Roman Catholic and Presbyterian, "
        "with other, please specify")
    codes = page[75] + " " + page[76]
    say("RELIGAFF" in codes and "7 Spiritualism and New Age Religions 8 No Religion 99 Other "
        "999 Not Stated" in codes,
        "printed pp.76-77: the religion classification's codes, no religion and not stated among them")

    # 5. UNSD table 28
    for y in (2011, 2016):
        got = oracle.oracle("Tokelau", y)
        if not got:
            raise SystemExit(f"Tokelau {y} is not in the oracle; run `python tools/oracle.py --fetch`")
        cats, stated, _exact = oracle.partition(got[oracle.TOTAL])
        mine = {UNSD_NAME[k]: sum(T58[y][k]) for k in CATS}
        theirs = {k: cats.get(k, 0) for k in mine}
        extra = sorted(set(cats) - set(mine))
        diff = {k: (mine[UNSD_NAME[k]], theirs[UNSD_NAME[k]]) for k in CATS
                if mine[UNSD_NAME[k]] != theirs[UNSD_NAME[k]]}
        want = {} if y == 2011 else UNSD_2016_DIFF
        say(stated == PRESENT[y] and not extra and diff == want,
            f"UNSD {y}: total {stated:,}; cells differing from the office (office, UNSD): "
            f"{diff or 'none'}{'; extra UNSD rows ' + str(extra) if extra else ''}")

    if not ok:
        raise SystemExit("reconciliation FAILED")


def emit():
    out = []
    for i, atoll in enumerate(ATOLLS):
        for c in CATS:
            n = T58[YEAR][c][i]
            if n <= 0:
                continue
            out.append({
                "geo_id": atoll, "geo_level": "atoll", "geo_name": atoll,
                "source_category": c, "count": n, "basis": BASIS, "year": YEAR,
                "source_id": SOURCE_ID,
                "note": (f"Table 5.8, usual residents present on census night; atoll total "
                         f"{T58_TOTAL[YEAR][i]}"),
            })
    return out


def main():
    if "--fetch" in sys.argv:
        fetch()
    missing = [p for p in DOWNLOADS if not os.path.exists(p)]
    if missing:
        raise SystemExit(f"missing {missing}; run: python sources/tk.py --fetch")
    check()
    out = emit()

    print(f"\n  3 atolls, {sum(r['count'] for r in out):,} people")
    for c in CATS:
        n = sum(T58[YEAR][c])
        print(f"    {n:>6,}  {100.0 * n / PRESENT[YEAR]:6.2f}%  {c}")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT + ".part", "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(out)
    os.replace(OUT + ".part", OUT)
    print("\nwrote", OUT, f"({len(out)} rows)")


if __name__ == "__main__":
    main()
