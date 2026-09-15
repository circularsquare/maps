"""Argentina — CEIL-CONICET's second national religion survey (2019), at its six regions.

Writes data/normalized/ar.csv: 6 answers x 6 survey regions, as PEOPLE.

Usage:
    python sources/ar.py --fetch     # two survey reports, one journal article, two census tables
    python sources/ar.py             # rebuild ar.csv from data/raw/ar/

THE SOURCE. *Segunda Encuesta Nacional sobre Creencias y Actitudes Religiosas en Argentina*,
Programa Sociedad, Cultura y Religión, CEIL-CONICET, funded by MINCyT:

    fieldwork      August - September 2019
    respondents    2,421 adults 18+, resident in localities of 5,000+ (2010 census frame)
    design         89 localities as PSUs, stratified by REGION and city size, PPS; census radios
                   by PPS; dwellings systematically; the respondent by sex and age quota
    margin         +/- 2% at 95%, the report's own figure
    geography      six regions, "posibilitando así una representatividad regional" -- the same
                   six as the first survey (2008, 2,403 cases), deliberately, for comparison

Argentina's census has not asked religion since 1960, it is absent from the UNSD oracle, and
LAPOP's Argentine rounds predate its religion question (sources.md §11ad). This is the only
source with a regional religion table.

WHAT IS PARSED, and from which of three documents:

    Tabla 5, page 8 of Mallimaci, Esquivel & Giménez Béliveau (2020), *Religiones y creencias
        en Argentina (2008-2019)*, Sociedad y Religión 30(55). Six answers by TOTAL and the six
        regions, INCLUDING the cells under 2%. This is what is drawn.
    Page 18 of the CEIL report (Informe de Investigación 25, 2019). The same regional chart,
        but it prints only values over 2%, so it cannot be the source; it is the CHECK, and
        the three large answers agree with Tabla 5 to the decimal in all six regions.
    Page 7 of the first survey's 2008 report. The 2008 regional split, used for the
        cross-wave stability test below and for nothing that is drawn.

**TABLA 7 OF THE SAME ARTICLE IS NOT USED, AND IT IS A TRAP.** It prints 2008 and 2019 side by
side by region and its Cuyo rows have `Sin filiación religiosa` and `Evangélica` SWAPPED in
both years -- against Tabla 5 for 2019, and against the 2008 report for 2008, whose stacked
chart is labelled in full on its Patagonia bar so the order is not in doubt. Its Patagonia
2019 column also differs from Tabla 5 (51.3 against 51.0 Catholic). The CEIL report's own
page 20 repeats the Patagonia figures. `check_table7()` prints the disagreements so a reader
can see them without trusting this paragraph.

WHICH PROVINCES ARE IN WHICH REGION IS PUBLISHED NOWHERE IN WORDS. Neither report, the
article, CEIL's survey page nor the dataset record lists them. The one statement of it is the
map on CONICET's 2019 infographic, whose regions are FILLED SHAPES in the PDF, so they were
read as vector fills rather than by eye (sources/ar.md): NOA is Jujuy, Salta, Tucumán,
Catamarca, Santiago del Estero and La Rioja; NEA is Formosa, Chaco, Misiones, Corrientes AND
ENTRE RÍOS, which INDEC's own regionalisation puts elsewhere; Cuyo is San Juan, San Luis and
Mendoza; Centro is Córdoba, Santa Fe, La Pampa and Buenos Aires province outside the capital's
conurbation; Patagonia is the five southern provinces. **AMBA is the Ciudad Autónoma plus the
24 partidos of Gran Buenos Aires**: the 2008 report labels the same region `Capital y GBA`, and
INDEC publishes that exact aggregate. The infographic's inset is ambiguous about Gran La Plata
(about 0.9M people); the 2008 label decides it.

WHICH ANSWERS CARRY THEIR OWN REGIONAL SHARES. The house test is the split-half (sources.md
§9bi) and it needs microdata, which CEIL deposited at ri.conicet.gov.ar/handle/11336/249205 and
EMBARGOED UNTIL 2026-12-31. So the nearest available thing stands in for it: the 2008 wave is
an independent sample of the same six regions by the same team with the same design, and an
answer whose regional ORDER survives eleven years is not an ordering 400 interviews invented.
The bar is sources/spearman_null.py's exact one-sided 95% bar on six units, +0.8286, and NOT
1.96/sqrt(n-1) = +0.877, which is a standard deviation and not a critical value (§9ct, Anita's
ruling on ask/007-cr). The exact permutation p over all 720 orderings is printed beside it. **This stand-in is conservative, not neutral**: real change
between the waves lowers the correlation exactly as noise does, which is why `OVERRIDE` exists.

Answers that fail are still drawn, at their national proportions inside each region's own
residual, which is Guatemala's construction (§9bi): the partition stays closed and only the
claim to know where those people are is withdrawn.

THE MAGNITUDE IS INDEC'S 2022 CENSUS, definitive results. `c2022_tp_c_resumen.xlsx` gives the
24 jurisdictions' total population (private and collective dwellings and people on the street,
45,892,285); Cuadro 1.2 of `c2022_bsas_est_c1_2.xlsx` gives the printed row `24 Partidos del
Gran Buenos Aires`, 10,849,398, which splits Buenos Aires province between AMBA and Centro. The
24 partido rows beneath it are asserted to sum to it.

2019 SHARES ON 2022 PEOPLE, spec §3.4. AND THE SURVEY IS ADULTS IN TOWNS OF 5,000 OR MORE while
the dots are everybody: children and small-town and rural Argentines are drawn at their region's
urban adult rates, because nothing measures them and a blank would read as an absence (§6.12).
"""

import argparse
import itertools
import math
import os
import re
import sys
import unicodedata

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
RAW = os.path.join(ROOT, "data", "raw", "ar")
OUT = os.path.join(ROOT, "data", "normalized", "ar.csv")

UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
      "Chrome/124.0 Safari/537.36")

REPORT_URL = "http://www.ceil-conicet.gov.ar/wp-content/uploads/2019/11/ii25-2encuestacreencias.pdf"
REPORT_PDF = os.path.join(RAW, "ceil_ii25_2019.pdf")
ARTICLE_URL = ("https://ri.conicet.gov.ar/bitstream/handle/11336/144739/"
               "CONICET_Digital_Nro.2f081de2-1c65-4a90-a4f8-b403287deddd_A.pdf"
               "?sequence=2&isAllowed=y")
ARTICLE_PDF = os.path.join(RAW, "mallimaci_2020_sociedad_y_religion_55.pdf")
REPORT2008_URL = "http://www.ceil-conicet.gov.ar/wp-content/uploads/2013/02/encuesta1.pdf"
REPORT2008_PDF = os.path.join(RAW, "ceil_encuesta1_2008.pdf")
RESUMEN_URL = "https://censo.gob.ar/wp-content/uploads/2024/01/c2022_tp_c_resumen.xlsx"
RESUMEN_XLSX = os.path.join(RAW, "c2022_tp_c_resumen.xlsx")
BSAS_URL = "https://censo.gob.ar/wp-content/uploads/2023/11/c2022_bsas_est_c1_2.xlsx"
BSAS_XLSX = os.path.join(RAW, "c2022_bsas_est_c1_2.xlsx")

# (url, destination, expected page count for a PDF)
DOWNLOADS = [
    (REPORT_URL, REPORT_PDF, 72),
    (ARTICLE_URL, ARTICLE_PDF, 31),
    (REPORT2008_URL, REPORT2008_PDF, 29),
    (RESUMEN_URL, RESUMEN_XLSX, None),
    (BSAS_URL, BSAS_XLSX, None),
]

SOURCE_ID = "ar_ceil_2019"
YEAR = 2019
POP_YEAR = 2022
N_RESPONDENTS = 2421

TABLE5_PAGE = 8          # article, 1-based
REPORT_REGION_PAGE = 18  # CEIL report, 1-based
REPORT_EVANG_PAGE = 14   # CEIL report, national chart with the evangelical split
TABLE1_PAGE = 4          # article, national Tabla 1 with the no-affiliation split
TABLE7_PAGE = 10         # article, the 2008 vs 2019 table that is NOT used
REGION2008_PAGE = 7      # 2008 report, 1-based

REGIONS = ["AMBA", "CENTRO", "NEA", "NOA", "CUYO", "PATAGONIA"]
REGION_NAMES = {"AMBA": "AMBA", "CENTRO": "Centro", "NEA": "NEA", "NOA": "NOA",
                "CUYO": "Cuyo", "PATAGONIA": "Patagonia"}

# Tabla 5's rows, in the order it prints them. These strings are `source_category`.
CATS = ["Católica", "Sin filiación religiosa", "Evangélica", "Testigos de Jehová/Mormones",
        "Otras", "No sabe"]

# The census's own jurisdiction names -> CEIL region. Buenos Aires province is split below.
CENSUS_TO_REGION = {
    "Ciudad Autónoma de Buenos Aires": "AMBA",
    "Buenos Aires": None,
    "Córdoba": "CENTRO", "Santa Fe": "CENTRO", "La Pampa": "CENTRO",
    "Formosa": "NEA", "Chaco": "NEA", "Misiones": "NEA", "Corrientes": "NEA",
    "Entre Ríos": "NEA",
    "Jujuy": "NOA", "Salta": "NOA", "Tucumán": "NOA", "Catamarca": "NOA",
    "Santiago del Estero": "NOA", "La Rioja": "NOA",
    "San Juan": "CUYO", "San Luis": "CUYO", "Mendoza": "CUYO",
    "Neuquén": "PATAGONIA", "Rio Negro": "PATAGONIA", "Chubut": "PATAGONIA",
    "Santa Cruz": "PATAGONIA",
    "Tierra del Fuego, Antártida e Islas del Atlántico Sur": "PATAGONIA",
}

# ---- the stability gate ------------------------------------------------------------------
OVERRIDE_CHI2_P = 1e-3


def _spearman_null():
    """The project's shared exact null (§9ct). Imported lazily; it pulls in numpy."""
    sys.path.insert(0, HERE)
    import spearman_null
    return spearman_null

# One named answer, decided by a person, with the reason printed on every run (§9bi). The bar
# is not moved: moving it would silently change every answer's treatment at once.
OVERRIDE = {
    "Evangélica": (
        "fails the cross-wave rank test because it MOVED, not because it is noise. NOA went "
        "from 3.7% evangelical in 2008 to 16.7% in 2019, the largest change in either table "
        "and the one the authors lead with, and a single region jumping four places is "
        "enough to sink a six-unit Spearman. The ends hold: Patagonia is the most evangelical "
        "region in both waves and Centro is in the bottom two in both. The regions differ at "
        f"chi-square p < {OVERRIDE_CHI2_P:g} under both sample allocations (asserted). Drawn "
        "flat, it would say Centro at 11.3% and Patagonia at 24.4% are the same."),
}

NUM = re.compile(r"^\d{1,3}(?:[.,]\d)?$")


def _num(s):
    return float(s.replace(",", "."))


def _norm(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    s = s.replace("\xa0", " ").replace(",", " ").replace(".", " ").replace("-", " ")
    return " ".join(s.lower().split())


def fetch():
    import requests
    import urllib3

    os.makedirs(RAW, exist_ok=True)
    for url, dest, _pages in DOWNLOADS:
        if os.path.exists(dest) and os.path.getsize(dest) > 10_000:
            print("already have", os.path.basename(dest))
            continue
        # censo.gob.ar serves an incomplete certificate chain ("unable to verify the first
        # certificate"), so verification is off for THAT HOST ONLY. The parser asserts the
        # workbook's content, which is the check that matters for a table.
        verify = "censo.gob.ar" not in url
        if not verify:
            urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)
        print("GET", url)
        r = requests.get(url, timeout=600, verify=verify, headers={"User-Agent": UA})
        r.raise_for_status()
        with open(dest, "wb") as fh:
            fh.write(r.content)
        print(f"  {os.path.basename(dest)}  {len(r.content):,} bytes")

    # [[reference_pdf_truncated_at_source]]: a Content-Length can match a damaged file.
    for _url, dest, _pages in DOWNLOADS:
        if not dest.endswith(".pdf"):
            continue
        with open(dest, "rb") as fh:
            head = fh.read(5)
            fh.seek(-2048, os.SEEK_END)
            tail = fh.read()
        if head != b"%PDF-" or b"%%EOF" not in tail:
            raise SystemExit(f"{dest} is not a complete PDF (no %%EOF in the last 2 KB)")


def _lines(pdf, page_no, pages_expected):
    import fitz

    if not os.path.exists(pdf):
        raise SystemExit(f"missing {pdf} -- run sources/ar.py --fetch first")
    doc = fitz.open(pdf)
    if doc.page_count != pages_expected:
        raise SystemExit(f"{os.path.basename(pdf)} has {doc.page_count} pages, expected "
                         f"{pages_expected} -- a different edition")
    return [ln.strip() for ln in doc[page_no - 1].get_text().split("\n") if ln.strip()]


def _label(parts):
    """Rejoin a row label the PDF broke over lines, including hyphenated breaks."""
    s = ""
    for p in parts:
        s = s[:-1] + p if s.endswith("-") else (s + " " + p).strip()
    s = re.sub(r"\s*/\s*", "/", s)
    return " ".join(s.split())


def _rows(lines, width, stop="Base"):
    """label lines followed by `width` numbers, repeated, until a line starting with `stop`.

    `Base` and not `Base:` -- Tabla 7's footnote prints `Base  2008:` with the year between.
    """
    rows, parts, nums = {}, [], []
    for ln in lines:
        if ln.startswith(stop):
            break
        if NUM.match(ln):
            if not parts:
                raise SystemExit(f"a number with no row label before it: {ln!r}")
            nums.append(_num(ln))
            if len(nums) == width:
                rows[_label(parts)] = nums
                parts, nums = [], []
        else:
            if nums:
                raise SystemExit(f"a label interrupted a row after {len(nums)} numbers: {ln!r}")
            parts.append(ln)
    if nums or parts:
        raise SystemExit(f"a row was left incomplete: {parts} {nums}")
    return rows


def read_table5():
    """Tabla 5 -> {answer: {TOTAL|region: percent}}."""
    lines = _lines(ARTICLE_PDF, TABLE5_PAGE, 31)
    t = next((i for i, ln in enumerate(lines) if ln.startswith("Tabla 5.")), None)
    if t is None or "región de residencia" not in lines[t]:
        raise SystemExit(f"article page {TABLE5_PAGE} is not Tabla 5 -- the PDF changed")
    i = lines.index("TOTAL", t)
    if lines[i:i + 7] != ["TOTAL"] + REGIONS:
        raise SystemExit(f"Tabla 5's column header is {lines[i:i + 7]}")
    rows = _rows(lines[i + 7:], 7)
    if list(rows) != CATS:
        raise SystemExit(f"Tabla 5's rows are {list(rows)}, expected {CATS}")

    table = {c: dict(zip(["TOTAL"] + REGIONS, v)) for c, v in rows.items()}
    for col in ["TOTAL"] + REGIONS:
        s = sum(table[c][col] for c in CATS)
        if abs(s - 100.0) > 0.3:
            raise SystemExit(f"Tabla 5's {col} column sums to {s:.1f}, not 100")
    # Two values read off the page by eye, so a silent column shift cannot pass.
    if table["Católica"]["TOTAL"] != 62.9 or table["Evangélica"]["PATAGONIA"] != 24.4:
        raise SystemExit("Tabla 5's anchors moved: national Católica should be 62.9 and "
                         "Patagonia Evangélica 24.4")
    return table


def check_report(table):
    """The CEIL report's own regional chart must agree with Tabla 5 on the three answers it
    prints (it omits values under 2%, which is why it is the check and not the source)."""
    lines = _lines(REPORT_PDF, REPORT_REGION_PAGE, 72)
    if "Adscripción religiosa según región" not in lines:
        raise SystemExit(f"report page {REPORT_REGION_PAGE} is not the regional chart")
    start = lines.index("Total")
    names = {"Total": "TOTAL", "AMBA": "AMBA", "NEA": "NEA", "NOA": "NOA",
             "Patagonia": "PATAGONIA", "Centro": "CENTRO", "Cuyo": "CUYO"}
    for name, col in names.items():
        j = lines.index(name, start)
        vals = []
        for ln in lines[j + 1:j + 4]:
            if not NUM.match(ln):
                raise SystemExit(f"report p{REPORT_REGION_PAGE}: {name} is followed by {ln!r}")
            vals.append(_num(ln))
        want = [table[c][col] for c in CATS[:3]]
        if vals != want:
            raise SystemExit(f"report p{REPORT_REGION_PAGE} gives {name} = {vals}, Tabla 5 "
                             f"gives {want}")
    print(f"  report p{REPORT_REGION_PAGE} agrees with Tabla 5 on Católica, Sin religión and "
          "Evangélica in all six regions and the total")


def read_2008():
    """The 2008 report's stacked regional chart -> {region: {answer: percent}}.

    Each bar prints its five values bottom to top and then its region name, except the
    Patagonia bar (`SUR`), whose name comes first and whose values are interleaved with the
    legend. That legend is what proves the stacking order, so it is asserted.
    """
    lines = _lines(REPORT2008_PDF, REGION2008_PAGE, 29)
    if not any("según región" in ln for ln in lines):
        raise SystemExit(f"2008 report page {REGION2008_PAGE} is not the regional chart")
    order = ["Otras", "Testigos de Jehová/Mormones", "Evangélica", "Sin filiación religiosa",
             "Católica"]
    labels = {"CUYO": "CUYO", "CAPITAL": "AMBA", "CENTRO": "CENTRO", "NEA": "NEA",
              "NOA": "NOA"}
    out, buf, sur, sur_words = {}, [], None, []
    for ln in lines:
        if ln.startswith("Base:"):
            break
        if NUM.match(ln):
            (sur if sur is not None else buf).append(_num(ln))
        elif sur is not None:
            sur_words.append(ln)
        elif ln in labels:
            if len(buf) != 5:
                raise SystemExit(f"2008 chart: {ln} has {len(buf)} values, expected 5")
            out[labels[ln]] = dict(zip(order, buf))
            buf = []
        elif ln == "SUR":
            if buf:
                raise SystemExit("2008 chart: values left over before SUR")
            sur = []
    if sur is None or len(sur) != 5:
        raise SystemExit("2008 chart: the SUR bar did not parse")
    out["PATAGONIA"] = dict(zip(order, sur))

    legend = ["Otras", "Testigos de", "Evangélica", "Indiferentes", "Católica"]
    pos = [next((k for k, w in enumerate(sur_words) if w.startswith(x)), -1) for x in legend]
    if -1 in pos or pos != sorted(pos):
        raise SystemExit(f"2008 chart: the SUR legend is not in stacking order: {sur_words}")
    if set(out) != set(REGIONS):
        raise SystemExit(f"2008 chart regions are {sorted(out)}")
    for reg, v in out.items():
        if abs(sum(v.values()) - 100.0) > 0.25:
            raise SystemExit(f"2008 chart: {reg} sums to {sum(v.values()):.1f}")
    if out["AMBA"]["Sin filiación religiosa"] != 18.0 or out["PATAGONIA"]["Evangélica"] != 21.6:
        raise SystemExit("2008 chart anchors moved: AMBA indiferentes 18.0, SUR evangélica 21.6")
    return out


def check_table7(table, y2008):
    """Print where the article's own 2008-vs-2019 table disagrees with the two sources used."""
    lines = _lines(ARTICLE_PDF, TABLE7_PAGE, 31)
    t = next((i for i, ln in enumerate(lines) if ln.startswith("Tabla 7.")), None)
    if t is None:
        raise SystemExit(f"article page {TABLE7_PAGE} has no Tabla 7")
    i = lines.index("Católica", t)
    rows = _rows(lines[i:], 12)
    canon = {_norm(c): c for c in CATS}
    bad = []
    for label, vals in rows.items():
        cat = canon.get(_norm(label))
        if cat is None:
            raise SystemExit(f"Tabla 7 row {label!r} is not a Tabla 5 answer")
        for k, reg in enumerate(REGIONS):
            v08, v19 = vals[2 * k], vals[2 * k + 1]
            w19 = table[cat][reg]
            w08 = y2008[reg].get(cat, 0.0)
            if abs(v19 - w19) > 0.05:
                bad.append(f"2019 {reg:9s} {cat:28s} Tabla 7 {v19:5.1f}  Tabla 5 {w19:5.1f}")
            if abs(v08 - w08) > 0.05:
                bad.append(f"2008 {reg:9s} {cat:28s} Tabla 7 {v08:5.1f}  2008 report {w08:5.1f}")
    print(f"\n  Tabla 7 (NOT used) disagrees with the sources in {len(bad)} cells:")
    for b in bad:
        print("    " + b)
    if len(bad) > 20:
        raise SystemExit("more than 20 disagreements -- that is a mis-parse, not typos")


def check_national_splits():
    """The finer national splits the regional table cannot carry. Asserted, not drawn."""
    lines = _lines(ARTICLE_PDF, TABLE1_PAGE, 31)
    for word, val in (("Atea", "6,0"), ("Agnóstica", "3,2"), ("Ninguna", "9,7")):
        if word not in lines or lines[lines.index(word) + 1] != val:
            raise SystemExit(f"Tabla 1: {word} should be {val}")
    rep = _lines(REPORT_PDF, REPORT_EVANG_PAGE, 72)
    for word, val in (("Pentecostales", "13"), ("evangélicos", "2.3")):
        if word not in rep or rep[rep.index(word) + 1] != val:
            raise SystemExit(f"report p{REPORT_EVANG_PAGE}: {word} should be {val}")
    print("  national-only splits present: Atea 6.0 / Agnóstica 3.2 / Ninguna 9.7 inside Sin "
          "filiación 18.9; Pentecostales 13 / Otros evangélicos 2.3 inside Evangélica 15.3")


def _int(v):
    return int(str(v).replace(".", "").strip())


def _cells(row):
    """A worksheet row without its empty cells. INDEC's sheets do not start their tables in
    column A, so reading by position silently finds nothing; read by content instead."""
    return [c for c in (row or ()) if c is not None and str(c).strip() != ""]


def _is_count(c):
    return isinstance(c, int) or bool(re.fullmatch(r"\d+", str(c).strip()))


def read_gba24():
    """Cuadro 1.2 -> ({code: (name, people 2022)} for the 24 partidos, their printed total,
    the province total). Shared with sources/ar_geo.py so the two cannot disagree.

    Each row is a code, a name (absent on the Total row), then 2010, 2022, the absolute and the
    relative change -- so 2022 is the SECOND whole number after the code.
    """
    import openpyxl

    if not os.path.exists(BSAS_XLSX):
        raise SystemExit(f"missing {BSAS_XLSX} -- run sources/ar.py --fetch first")
    ws = openpyxl.load_workbook(BSAS_XLSX, read_only=True, data_only=True)["Cuadro1.2"]
    total = gba = None
    partidos, in_gba = {}, False
    for row in ws.iter_rows(values_only=True):
        cells = _cells(row)
        if len(cells) < 3:
            continue
        first = cells[0]
        if isinstance(first, int):
            code = f"{first:05d}" if first > 99 else f"{first:02d}"
        else:
            code = " ".join(str(first).replace("\xa0", " ").split())
        name = next((" ".join(str(c).replace("\xa0", " ").split())
                     for c in cells[1:] if not _is_count(c)), "")
        counts = [_int(c) for c in cells[1:] if _is_count(c)]
        if len(counts) < 2:
            continue
        if code == "Total":
            total = counts[1]
        elif code == "06" and name.startswith("24 Partidos del Gran Buenos Aires"):
            gba, in_gba = counts[1], True
        elif in_gba and re.fullmatch(r"06\d{3}", code):
            partidos[code] = (name, counts[1])
        elif in_gba:
            in_gba = False
    if gba is None or total is None:
        raise SystemExit("Cuadro 1.2 has no `24 Partidos del Gran Buenos Aires` or no Total row")
    if len(partidos) != 24:
        raise SystemExit(f"Cuadro 1.2 lists {len(partidos)} partidos under the GBA row, not 24")
    if sum(p for _n, p in partidos.values()) != gba:
        raise SystemExit("the 24 partido rows do not sum to the printed GBA row")
    return partidos, gba, total


def read_population():
    """INDEC 2022 -> {region: people}."""
    import openpyxl

    if not os.path.exists(RESUMEN_XLSX):
        raise SystemExit(f"missing {RESUMEN_XLSX} -- run sources/ar.py --fetch first")
    ws = openpyxl.load_workbook(RESUMEN_XLSX, read_only=True, data_only=True)["cuadro_resumen"]
    prov, country = {}, None
    for row in ws.iter_rows(values_only=True):
        cells = _cells(row)
        if len(cells) < 2 or not isinstance(cells[0], str) or not _is_count(cells[1]):
            continue
        name = " ".join(cells[0].replace("\xa0", " ").split())
        if name == "Total del país":
            country = _int(cells[1])
        elif name in CENSUS_TO_REGION:
            prov[name] = _int(cells[1])
    if set(prov) != set(CENSUS_TO_REGION):
        raise SystemExit("the census jurisdictions and the region crosswalk disagree:\n"
                         f"  missing: {sorted(set(CENSUS_TO_REGION) - set(prov))}")
    if sum(prov.values()) != country:
        raise SystemExit(f"the 24 jurisdictions sum to {sum(prov.values()):,}, the table's "
                         f"total is {country:,}")

    _partidos, gba, ba_total = read_gba24()
    if ba_total != prov["Buenos Aires"]:
        raise SystemExit(f"Cuadro 1.2's province total {ba_total:,} is not the summary "
                         f"table's {prov['Buenos Aires']:,}")

    pop = dict.fromkeys(REGIONS, 0)
    for name, reg in CENSUS_TO_REGION.items():
        if reg is not None:
            pop[reg] += prov[name]
    pop["AMBA"] += gba
    pop["CENTRO"] += prov["Buenos Aires"] - gba
    if sum(pop.values()) != country:
        raise SystemExit("the regions lost people")
    print(f"  population: {country:,} (INDEC census {POP_YEAR}, definitive); AMBA takes the "
          f"printed 24-partido row, {gba:,}")
    return pop


# ---- stability -----------------------------------------------------------------------------

def _ranks(xs):
    order = sorted(range(len(xs)), key=lambda k: xs[k])
    ranks, i = [0.0] * len(xs), 0
    while i < len(order):
        j = i
        while j + 1 < len(order) and xs[order[j + 1]] == xs[order[i]]:
            j += 1
        for k in range(i, j + 1):
            ranks[order[k]] = (i + j) / 2 + 1
        i = j + 1
    return ranks


def spearman(a, b):
    ra, rb = _ranks(a), _ranks(b)
    ma, mb = sum(ra) / len(ra), sum(rb) / len(rb)
    va = sum((x - ma) ** 2 for x in ra)
    vb = sum((y - mb) ** 2 for y in rb)
    if va == 0 or vb == 0:
        return float("nan")
    return sum((x - ma) * (y - mb) for x, y in zip(ra, rb)) / math.sqrt(va * vb)


def permutation_p(a, b):
    """Exact: the share of all orderings of `b` whose correlation with `a` is at least as high."""
    obs = spearman(a, b)
    if math.isnan(obs):
        return obs, float("nan")
    perms = list(itertools.permutations(b))
    hits = sum(1 for p in perms if spearman(a, list(p)) >= obs - 1e-12)
    return obs, hits / len(perms)


def chi2_sf_df5(x):
    """Upper tail of chi-square on 5 degrees of freedom, closed form for odd df."""
    if x <= 0:
        return 1.0
    h = x / 2.0
    return math.erfc(math.sqrt(h)) + math.exp(-h) * (
        math.sqrt(h) / math.gamma(1.5) + h ** 1.5 / math.gamma(2.5))


def chi2_regions(shares, n):
    """2 x 6 homogeneity test for one answer, from shares (percent) and an assumed n per region."""
    N = sum(n[r] for r in REGIONS)
    pbar = sum(n[r] * shares[r] / 100.0 for r in REGIONS) / N
    if pbar <= 0 or pbar >= 1:
        return float("nan"), float("nan")
    x = sum(n[r] * (shares[r] / 100.0 - pbar) ** 2 for r in REGIONS) / (pbar * (1 - pbar))
    return x, chi2_sf_df5(x)


def stability(table, y2008, pop):
    """Decide which answers keep their own regional shares. Prints the evidence."""
    total = sum(pop.values())
    alloc = {
        "proportional": {r: N_RESPONDENTS * pop[r] / total for r in REGIONS},
        "equal": {r: N_RESPONDENTS / len(REGIONS) for r in REGIONS},
    }
    print("\n  per-region sample size is NOT published, so the chi-square is run under two "
          "allocations:")
    for k, v in alloc.items():
        print(f"    {k:12s} " + "  ".join(f"{r} {v[r]:.0f}" for r in REGIONS))

    sn = _spearman_null()
    bar, how = sn.critical_rho(len(REGIONS))
    print(f"\n  cross-wave rank test, 2008 vs 2019, six regions; bar {bar:+.4f} ({how})")
    print(f"    {'':28s} {'rho':>6s} {'perm p':>7s} {'chi2 prop':>10s} {'p':>9s} "
          f"{'chi2 eq':>8s} {'p':>9s}  decision")
    own = []
    for cat in CATS:
        a = [y2008[r].get(cat, 0.0) for r in REGIONS]
        b = [table[cat][r] for r in REGIONS]
        rho, pp = permutation_p(a, b)
        s = {r: table[cat][r] for r in REGIONS}
        x1, p1 = chi2_regions(s, alloc["proportional"])
        x2, p2 = chi2_regions(s, alloc["equal"])
        ties = sn.ties_note(a, b)
        if ties:
            print(f"    ({cat}: {ties})")
        if not math.isnan(rho) and rho >= bar:
            decision = "own regional shares"
            own.append(cat)
        elif cat in OVERRIDE:
            if not (p1 < OVERRIDE_CHI2_P and p2 < OVERRIDE_CHI2_P):
                raise SystemExit(f"OVERRIDE on {cat} is not available: the regions do not "
                                 f"differ at p < {OVERRIDE_CHI2_P:g} under both allocations")
            decision = "own regional shares (OVERRIDE)"
            own.append(cat)
        else:
            decision = "national rate inside the residual"
        print(f"    {cat:28s} {rho:+6.3f} {pp:7.4f} {x1:10.1f} {p1:9.2e} {x2:8.1f} {p2:9.2e}"
              f"  {decision}")
    for cat, why in OVERRIDE.items():
        print(f"\n  OVERRIDE {cat}: {why}")
    return own


def build_shares(table, own):
    """Region -> {answer: percent}, closed to the region's own total."""
    rest = [c for c in CATS if c not in own]
    nat_rest = sum(table[c]["TOTAL"] for c in rest)
    out = {}
    for reg in REGIONS:
        s = {c: table[c][reg] for c in own}
        resid = 100.0 - sum(s.values())
        if resid < -0.3:
            raise SystemExit(f"{reg}: the measured answers exceed 100 by {-resid:.2f}")
        resid = max(resid, 0.0)
        for c in rest:
            s[c] = resid * table[c]["TOTAL"] / nat_rest
        out[reg] = s
    return out


def to_people(shares, population):
    """shares x population -> whole people summing EXACTLY to population (largest remainder)."""
    scale = sum(shares.values())
    exact = {k: v / scale * population for k, v in shares.items()}
    floors = {k: int(v) for k, v in exact.items()}
    short = population - sum(floors.values())
    for k in sorted(exact, key=lambda k: exact[k] - floors[k], reverse=True)[:short]:
        floors[k] += 1
    return floors


def main():
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--fetch", action="store_true", help="download the reports and census tables")
    args = ap.parse_args()

    os.makedirs(RAW, exist_ok=True)
    if args.fetch:
        fetch()

    import pandas as pd

    table = read_table5()
    print(f"  Tabla 5: {len(CATS)} answers x {len(REGIONS)} regions, n={N_RESPONDENTS:,}")
    check_report(table)
    check_national_splits()
    y2008 = read_2008()
    check_table7(table, y2008)
    pop = read_population()

    own = stability(table, y2008, pop)
    shares = build_shares(table, own)

    rows = []
    for reg in REGIONS:
        people = pop[reg]
        counts = to_people(shares[reg], people)
        if sum(counts.values()) != people:
            raise SystemExit(f"{reg}: apportionment lost people")
        for cat in CATS:
            rule = ("own regional share" if cat in own
                    else "national proportions inside the region's residual")
            rows.append({
                "geo_id": reg, "geo_level": "region", "geo_name": REGION_NAMES[reg],
                "source_category": cat, "count": counts[cat],
                "basis": "self_id", "year": YEAR, "source_id": SOURCE_ID,
                "note": (f"share={shares[reg][cat]:.4f}% of the region; rule={rule}; "
                         f"published={table[cat][reg]:.1f}%; population={people} (INDEC "
                         f"census {POP_YEAR}); CEIL-CONICET 2019, n={N_RESPONDENTS} nationally"),
            })

    total = sum(pop.values())
    for cat in CATS:
        rows.append({
            "geo_id": "AR", "geo_level": "country", "geo_name": "Argentina",
            "source_category": cat, "count": round(table[cat]["TOTAL"] / 100.0 * total),
            "basis": "self_id", "year": YEAR, "source_id": SOURCE_ID,
            "note": f"share={table[cat]['TOTAL']:.1f}%; Tabla 5's own TOTAL column",
        })

    out = pd.DataFrame(rows)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    out.to_csv(OUT, index=False, encoding="utf-8")
    drawn = out.loc[out["geo_level"] == "region", "count"].sum()
    if drawn != total:
        raise SystemExit(f"the region rows total {drawn:,}, not {total:,}")
    print(f"\n  wrote {OUT}  {len(out)} rows; drawn {drawn:,} people on {len(REGIONS)} regions")

    print("\n  population-weighted regions vs Tabla 5's own TOTAL (the survey weights urban "
          "adults, so this is a relationship, not an identity):")
    for cat in CATS:
        w = sum(shares[r][cat] * pop[r] for r in REGIONS) / total
        print(f"    {cat:28s} {w:6.2f}%  vs {table[cat]['TOTAL']:5.1f}%")

    print("\n  what the map draws (% of each region):")
    print("    " + f"{'':11s}" + "".join(f"{c[:11]:>12s}" for c in CATS) + f"{'people':>13s}")
    for reg in REGIONS:
        print(f"    {REGION_NAMES[reg]:11s}" + "".join(f"{shares[reg][c]:12.2f}" for c in CATS)
              + f"{pop[reg]:13,}")


if __name__ == "__main__":
    main()
