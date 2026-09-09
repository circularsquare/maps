"""Cabo Verde — INE, V Recenseamento Geral da População e Habitação (RGPH-2021).

Reads (or fetches) data/raw/cv/ and writes data/normalized/cv.csv.

**FIFTEEN CATEGORIES ON ALL TWENTY-TWO CONCELHOS, AND THE QUEUE PRICED THIS COUNTRY AS
NATIONAL-ONLY.** §11w had Cabo Verde at "open office, not chased past the catch-all", from
`tools/oracle.py`, whose rows are national and urban/rural. INE publishes one workbook per
municipality — *CABO VERDE EM NÚMEROS*, twenty-two of them plus a national one, filed under
the census area's `Quadros por Concelho 2021` and not under its publications — and every one
of them carries the religion table. 491,233 people over 22 units is 22,000 a unit, which is
finer than the two thirds of this map that are drawn at a first-level province.

**THE TABLE'S UNIVERSE IS THE POPULATION AGED 15 AND OVER**, 352,494 of 491,233. INE asked
nobody younger, so 138,739 people, 28.24% of the country, are not in any religion row here.
That figure is not inferred: Tabela 1 of the same workbook prints 0-4, 5-9 and 10-14 as
45,540 + 46,619 + 46,580 = 138,739, and 491,233 - 352,494 is the same number to the person.
It is the `gap` in countries.py and the reason `gap_share` is 28.5% rather than 0.3%.

**THE ORACLE'S SIXTEENTH CATEGORY IS THOSE CHILDREN.** `tools/oracle.py "Cabo Verde"` prints
sixteen rows summing exactly to 491,233, of which `Unknown` is 138,739 — the under-15s, not
a non-response. The queue's "16 categories" was therefore one too many, and the actual
published list is fifteen, of which one is a catch-all (`Outra`, 4,090), one a refusal
(`Não sabe / Não respondeu`, 1,311) and one `Sem religião` (54,814). The other twelve are
named religions, and nine of those are individual Christian churches.

**THE INTERESTING ONE IS `Racionalismo Cristão`, 6,129 people, 1.74%** — the fourth largest
named religion in the country, behind Catholic (255,511), Adventist (6,626) and the Church
of the Nazarene (6,175), and within 500 people of both of those. Christian Rationalism is a
spiritualist doctrine founded in Santos, Brazil, in 1910 by Luís de Mattos, and it reached
Cabo Verde in 1911 with a Cape Verdean medium returning from Rio. **Its geography is the
historical account exactly**: 3,988 of the 6,129 are in São Vicente, where it is 6.86% of
the adult population, and the next four concelhos are Tarrafal de São Nicolau (4.41%), Boa
Vista (3.54%), Paul on Santo Antão (3.17%) and Sal (2.38%) — the same islands the movement's
own history names, in roughly that order. `check()` prints it. See `sources/cv.md`.

**THE TWENTY-TWO FILES RECONCILE TO THE NATIONAL WORKBOOK EXACTLY, ON EVERY CATEGORY.**
That is the whole reconciliation: `data/raw/cv/137_...xlsx` is a twenty-third file published
by the same office, and summing the twenty-two municipal ones reproduces all fifteen of its
figures and its 352,494 total to the person. The UNSD Demographic Yearbook is then a third,
independent transcription of the same table and is asserted too.

**TWO LABELS ARE SPELT TWO WAYS AND THE ROW ORDER NEVER CHANGES.** Fourteen of the workbooks
(the 2022-2023 batch, plus the national one) print `Racionalismo Critão` and `... dos
últimos dias`; the nine published in 2025 print `Racionalismo Cristão` and `... dos Últimos
Dias`. Nothing else differs, and `read()` asserts that all twenty-three files list the
fifteen categories in the same order, which is the check a permuted or re-ordered table
would fail.

**THE JOIN TO BOUNDARIES IS BY HAND AND IS ASSERTED ON POPULATION, BECAUSE THREE PAIRS OF
CONCELHOS SHARE A NAME.** Ribeira Grande is a concelho of Santo Antão and Ribeira Grande de
Santiago is a different one; Santa Catarina is on Santiago and Santa Catarina do Fogo is
not; Tarrafal is on Santiago and Tarrafal de São Nicolau is not. The workbooks' own table
titles use the SHORT form for all three — `concelho de Ribeira Grande`, `concelho de Santa
Catarina`, `concelho de Tarrafal` — so a name join on the sheet would pair each of them with
a coin flip, and [[reference_name_join_wrong_neighbour]] is exactly this. `PCODE` below maps
INE's own content id to the COD-AB pcode, and `sources/cv_geo.py` proves the pairing against
COD-PS's per-concelho population rather than against a row count.

Usage:
    python sources/cv.py --fetch    23 workbooks (~4.8 MB) from INE's site API, seconds
    python sources/cv.py            normalise from data/raw/cv/
"""

import csv
import json
import os
import ssl
import sys
import urllib.request

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
RAW = os.path.join(ROOT, "data", "raw", "cv")
OUT = os.path.join(ROOT, "data", "normalized", "cv.csv")
INDEX = os.path.join(RAW, "_index.json")

SOURCE_ID = "cv_rgph_2021"
YEAR = 2021
BASIS = "self_id"
GEO_LEVEL = "concelho"

COLUMNS = ["geo_id", "geo_level", "geo_name", "source_category", "count",
           "basis", "year", "source_id", "note"]

# ine.cv is an Angular app and its pages carry no links; the API the bundle names is open.
# `Census` -> the two censuses; `Census/2` -> Censo 2021's four content categories, of which
# 3 is `Quadros por Concelho 2021`; `Census/content/paginated/3` -> the twenty-three
# workbooks; `GenericFile/GetFile?codePublication=<id>&publicationType=censo` -> the file.
# sources/cv.md §1.
API = "https://bdmi.ine.cv/site_deploy_api/api"
FILE_BASE = "https://bdmi.ine.cv/site_deploy_api/"
CATEGORY = 3               # "Quadros por Concelho 2021"
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) religiondots/1.0"}

NATIONAL_FILE_ID = 137     # "Cabo Verde - CORRIGIDO", the national workbook

# INE's content id -> the COD-AB adm1 pcode. Written out rather than joined on a name
# because three of the twenty-two names are ambiguous; see the module docstring, and
# sources/cv_geo.py, which asserts every one of these pairings against an independent
# population. The names are INE's own, from the API listing, which disambiguates all three
# where the sheet titles do not.
PCODE = {
    124: ("CV01", "Boa Vista"),
    130: ("CV02", "Brava"),
    133: ("CV03", "Maio"),
    129: ("CV04", "Mosteiros"),
    120: ("CV05", "Paul"),
    121: ("CV06", "Porto Novo"),
    128: ("CV07", "Praia"),
    122: ("CV08", "Ribeira Brava"),
    119: ("CV09", "Ribeira Grande de Santo Antão"),
    125: ("CV10", "Ribeira Grande de Santiago"),
    123: ("CV11", "Sal"),
    127: ("CV12", "Santa Catarina de Santiago"),
    132: ("CV13", "Santa Catarina do Fogo"),
    143: ("CV14", "Santa Cruz"),
    144: ("CV15", "São Domingos"),
    148: ("CV16", "São Filipe"),
    147: ("CV17", "São Lourenço dos Órgãos"),
    145: ("CV18", "São Miguel"),
    146: ("CV19", "São Salvador do Mundo"),
    140: ("CV20", "São Vicente"),
    142: ("CV21", "Tarrafal de Santiago"),
    141: ("CV22", "Tarrafal de São Nicolau"),
}

# The fifteen categories, in INE's own print order, which is identical in all 23 workbooks.
CATEGORIES = [
    "Adventista",
    "Assembleia de Deus",
    "Católica",
    "Deus é Amor",
    "Igreja do Nazareno / Protestante",
    "Islâmica / Muçulmano",
    "Judaica",
    "Nova Apastólica",
    "Racionalismo Cristão",
    "Testemunha de Jeová",
    "Universal do Reino de Deus",
    "Jesus Cristo dos Santos dos Últimos Dias / Mórmons",
    "Outra",
    "Sem religião",
    "Não sabe / Não respondeu",
]

# Two labels are spelt two ways across the batches and nothing else is. `Racionalismo
# Critão` is a typo for `Cristão` that INE corrected between the 2022-2023 workbooks and the
# 2025 ones; `últimos dias` / `Últimos Dias` is a capitalisation. Folded to the corrected
# spelling, and both variants are kept here so that a THIRD spelling stops the run instead
# of arriving as a sixteenth category. `Nova Apastólica` is a typo for `Apostólica` that
# INE has NOT corrected, so it stands as printed (spec §2.4, transcribe).
ALIASES = {
    "Racionalismo Critão": "Racionalismo Cristão",
    "Jesus Cristo dos Santos dos últimos dias / Mórmons":
        "Jesus Cristo dos Santos dos Últimos Dias / Mórmons",
}

# The universe of the religion table, and the country it is a subset of. Tabela 1 of the
# national workbook, asserted in check().
POP_15_PLUS = 352_494
POPULATION = 491_233
UNDER_15 = 138_739

# UNSD Demographic Yearbook table 28, Cabo Verde 2021 — INE's own English transcription of
# the same table, forwarded to New York, plus the under-15 residual as `Unknown`.
# `tools/oracle.py "Cabo Verde"`.
UNSD = {
    "Catholic": 255_511, "Unknown": 138_739, "No Religion": 54_814, "Adventist": 6_626,
    "Church of Nazarene": 6_175, "Christian Rationalism": 6_129, "Islam ": 4_616,
    "Other": 4_090, "Jehovah Witness": 4_083,
    "Church of Jesus Christ of Latter-day Saints": 3_565,
    "Universal of the Kingdom of God": 2_802, "New Apostolic": 1_719,
    "Didn't answer": 1_311, "Assembly of God": 730, "God is Love": 300, "Jewish": 23,
}
UNSD_TO_INE = {
    "Catholic": "Católica",
    "No Religion": "Sem religião",
    "Adventist": "Adventista",
    "Church of Nazarene": "Igreja do Nazareno / Protestante",
    "Christian Rationalism": "Racionalismo Cristão",
    "Islam ": "Islâmica / Muçulmano",
    "Other": "Outra",
    "Jehovah Witness": "Testemunha de Jeová",
    "Church of Jesus Christ of Latter-day Saints":
        "Jesus Cristo dos Santos dos Últimos Dias / Mórmons",
    "Universal of the Kingdom of God": "Universal do Reino de Deus",
    "New Apostolic": "Nova Apastólica",
    "Didn't answer": "Não sabe / Não respondeu",
    "Assembly of God": "Assembleia de Deus",
    "God is Love": "Deus é Amor",
    "Jewish": "Judaica",
}

UNDER_15_BANDS = ("0-4", "5-9", "10-14")


def _key(s):
    return " ".join(str(s).split())


def _get(url):
    ctx = ssl.create_default_context()
    ctx.check_hostname = False
    ctx.verify_mode = ssl.CERT_NONE
    return urllib.request.urlopen(urllib.request.Request(url, headers=UA),
                                  timeout=180, context=ctx).read()


def fetch():
    """The twenty-three `Quadros por Concelho 2021` workbooks, straight off the site API."""
    os.makedirs(RAW, exist_ok=True)
    listing = json.loads(_get(f"{API}/Census/content/paginated/{CATEGORY}"
                              "?codigo_lingua=PT&page_number=1&page_size=200"))
    items = listing["items"]
    if len(items) != len(PCODE) + 1:
        raise SystemExit(f"INE lists {len(items)} workbooks under category {CATEGORY}, "
                         f"expected {len(PCODE) + 1} (22 concelhos + Cabo Verde). The "
                         "category was re-cut; check Census/2 for its current shape.")
    index = []
    for it in items:
        code = it["codigo"]
        files = json.loads(_get(f"{API}/GenericFile/GetFile?codePublication={code}"
                                f"&publicationType=censo"))
        if len(files) != 1:
            raise SystemExit(f"content {code} ({it['titulo']}) has {len(files)} attachments")
        f = files[0]
        name = f"{code}{f['extensao']}"
        path = os.path.join(RAW, name)
        if not os.path.exists(path) or os.path.getsize(path) != int(f["tamanho"]):
            with open(path + ".part", "wb") as fh:
                fh.write(_get(FILE_BASE + f["url"]))
            os.replace(path + ".part", path)
        index.append({"content": code, "titulo": it["titulo"], "file": name,
                      "nome_ficheiro": f["nome_ficheiro"], "url": FILE_BASE + f["url"]})
        print(f"  {code:>4}  {it['titulo'][:34]:<36}{name}  {f['tamanho']:>8} bytes")
    with open(INDEX, "w", encoding="utf-8") as fh:
        json.dump(index, fh, ensure_ascii=False, indent=1)
    print(f"\n  {len(index)} workbooks -> {RAW}")


def _path(code):
    for ext in (".xlsx", ".xls"):
        p = os.path.join(RAW, f"{code}{ext}")
        if os.path.exists(p):
            return p
    raise SystemExit(f"missing data/raw/cv/{code}.xlsx — run `python sources/cv.py --fetch`")


def _labelled_rows(ws):
    """-> [(label, first numeric cell)], the shape both sheet layouts share.

    The national workbook indents its tables by one column and the municipal ones do not,
    and one file merges the title across two rows. Dropping empty cells and taking the first
    label-then-number pair on a row reads both without knowing which is which.
    """
    out = []
    for row in ws.iter_rows(values_only=True):
        cells = [c for c in row if c is not None and c != ""]
        if (len(cells) >= 2 and isinstance(cells[0], str)
                and isinstance(cells[1], (int, float)) and not isinstance(cells[1], bool)):
            out.append((_key(cells[0]), int(cells[1])))
    return out


def _religion_sheet(wb):
    """The one sheet whose name starts RELIG. Six spellings across the 23 workbooks:
    RELIGIÃO_1, RELIGIÃO_ESPIRITUALIDADE_1, RELIGIÃO ESPIRITUALIDADE_1 and so on."""
    got = [s for s in wb.sheetnames if _key(s).upper().startswith("RELIG")]
    if not got:
        raise SystemExit(f"no RELIG* sheet in {wb.sheetnames}")
    return wb[got[0]]


def _population_sheet(wb):
    """Tabela 1, population by age band, FOUND BY ITS CONTENT AND NOT BY ITS NAME.

    The 23 workbooks call it POPULAÇÃO_RESIDENTE, POPULAÇÃO, POPULAÇÃO_RESIDENTE_1 and,
    in Brava's, POP_BRAVA. A name prefix that covers all four also covers nothing else here,
    but the tab is renamed once per edition and the content is not, so this looks for the
    table: a `Total` row and the three under-15 bands, on one sheet.
    """
    for name in wb.sheetnames:
        rows = _labelled_rows(wb[name])
        labs = {lab for lab, _ in rows}
        if "Total" in labs and all(b in labs for b in UNDER_15_BANDS):
            return rows
    raise SystemExit("no sheet carries a Total row and the 0-4 / 5-9 / 10-14 bands: "
                     f"{wb.sheetnames}")


def read():
    """-> {content_id: {"religion": {cat: n}, "total15": n, "population": n, "under15": n}}"""
    out = {}
    orders = {}
    for code in list(PCODE) + [NATIONAL_FILE_ID]:
        wb = openpyxl.load_workbook(_path(code), read_only=True, data_only=True)

        rel = _labelled_rows(_religion_sheet(wb))
        if not rel or rel[0][0] != "Total":
            raise SystemExit(f"{code}: religion sheet does not start with a Total row: "
                             f"{rel[:2]}")
        total15 = rel[0][1]
        cats = {}
        order = []
        for lab, n in rel[1:]:
            lab = ALIASES.get(lab, lab)
            if lab not in CATEGORIES:
                raise SystemExit(f"{code}: unknown religion label {lab!r} — a sixteenth "
                                 "category or a third spelling; see ALIASES")
            if lab in cats:
                raise SystemExit(f"{code}: {lab!r} twice")
            cats[lab] = n
            order.append(lab)
        if sorted(cats) != sorted(CATEGORIES):
            raise SystemExit(f"{code}: {len(cats)} categories, expected {len(CATEGORIES)}: "
                             f"{sorted(set(CATEGORIES) ^ set(cats))}")
        orders.setdefault(tuple(order), []).append(code)

        pop = dict(_population_sheet(wb))
        out[code] = {"religion": cats, "total15": total15, "population": pop["Total"],
                     "under15": sum(pop[b] for b in UNDER_15_BANDS)}

    if len(orders) != 1:
        raise SystemExit(f"the workbooks do not agree on the print order of the fifteen "
                         f"categories: {[(v, list(k)[:3]) for k, v in orders.items()]}")
    return out


def build(t):
    rows = []
    for code, (pcode, name) in sorted(PCODE.items(), key=lambda kv: kv[1][0]):
        for cat in CATEGORIES:
            rows.append({
                "geo_id": pcode, "geo_level": GEO_LEVEL, "geo_name": name,
                "source_category": cat, "count": t[code]["religion"][cat],
                "basis": BASIS, "year": YEAR, "source_id": SOURCE_ID,
                "note": "tier=measured; universe=15+",
            })
    return rows


def check(t, rows):
    ok = True

    def say(good, msg):
        nonlocal ok
        ok = ok and good
        print(f"  {'ok  ' if good else 'FAIL'}  {msg}")

    nat = t[NATIONAL_FILE_ID]
    print("\nreconciliation")
    say(nat["total15"] == POP_15_PLUS,
        f"table  the national religion table counts {nat['total15']:,} == {POP_15_PLUS:,} "
        "people aged 15 and over")
    say(nat["population"] == POPULATION,
        f"table  and Tabela 1 of the same workbook counts {nat['population']:,} residents "
        f"of every age == {POPULATION:,}")
    say(nat["under15"] == UNDER_15 == POPULATION - POP_15_PLUS,
        f"table  its 0-4, 5-9 and 10-14 bands are {nat['under15']:,} == "
        f"{POPULATION:,} - {POP_15_PLUS:,}, so the missing universe is the under-15s "
        "and nothing else")

    # The twenty-two against the twenty-third, category by category.
    bad = []
    for cat in CATEGORIES:
        got = sum(t[c]["religion"][cat] for c in PCODE)
        if got != nat["religion"][cat]:
            bad.append((cat, got, nat["religion"][cat]))
    say(not bad, f"sum    all 15 categories sum over the 22 concelho workbooks to the "
                 f"national workbook's own figures {bad[:3]}  <- THE RECONCILIATION")
    tot = sum(t[c]["total15"] for c in PCODE)
    say(tot == POP_15_PLUS,
        f"sum    and their totals to {tot:,} == {POP_15_PLUS:,}")
    tot_pop = sum(t[c]["population"] for c in PCODE)
    say(tot_pop == POPULATION,
        f"sum    the same workbooks' Tabela 1 totals to {tot_pop:,} == {POPULATION:,}")

    bad = [(c, t[c]["total15"], sum(t[c]["religion"].values())) for c in PCODE
           if sum(t[c]["religion"].values()) != t[c]["total15"]]
    say(not bad, f"unit   every concelho's 15 categories partition its own printed total "
                 f"{bad[:3]}")
    bad = [(c, t[c]["population"] - t[c]["under15"], t[c]["total15"]) for c in PCODE
           if t[c]["population"] - t[c]["under15"] != t[c]["total15"]]
    say(not bad, f"unit   and its religion universe is its own population minus its own "
                 f"under-15s {bad[:3]}")

    say(sum(UNSD.values()) == POPULATION,
        f"UNSD   the oracle's 16 rows partition {sum(UNSD.values()):,} == {POPULATION:,}")
    say(UNSD["Unknown"] == UNDER_15,
        f"UNSD   and its `Unknown` is {UNSD['Unknown']:,} == the under-15 population, "
        "so it is an age cut and not a non-response")
    bad = [(k, UNSD[k], nat["religion"][v]) for k, v in UNSD_TO_INE.items()
           if UNSD[k] != nat["religion"][v]]
    say(not bad, f"UNSD   all 15 named figures reproduce the workbook exactly {bad[:3]}  "
                 "<- INDEPENDENT TRANSCRIPTION")

    # ---- the built rows ----
    say(len(rows) == len(PCODE) * len(CATEGORIES),
        f"built  {len(rows):,} rows == 22 concelhos x 15 categories")
    say(len({r['geo_id'] for r in rows}) == 22,
        f"built  {len({r['geo_id'] for r in rows})} distinct pcodes")
    drawn = sum(r["count"] for r in rows if r["source_category"] != "Não sabe / Não respondeu")
    say(sum(r["count"] for r in rows) == POP_15_PLUS,
        f"built  and they carry {sum(r['count'] for r in rows):,} people == the table's "
        "own universe")
    print(f"\n  {drawn:,} people are placed and {POPULATION - drawn:,} "
          f"({100.0 * (POPULATION - drawn) / POPULATION:.2f}% of Cabo Verde) are not: "
          f"{UNDER_15:,} under-15s never asked, plus "
          f"{nat['religion']['Não sabe / Não respondeu']:,} who did not answer.")

    print("\n  national, as drawn:")
    for cat in sorted(CATEGORIES, key=lambda c: -nat["religion"][c]):
        n = nat["religion"][cat]
        print(f"    {n:>10,}  {100.0 * n / POP_15_PLUS:6.2f}%  {cat}")

    # Christian Rationalism is the country's own religion and its geography is the finding.
    print("\n  Racionalismo Cristão, by concelho (% of the concelho's 15+ population):")
    share = sorted(((100.0 * t[c]["religion"]["Racionalismo Cristão"] / t[c]["total15"],
                     PCODE[c][1], t[c]["religion"]["Racionalismo Cristão"]) for c in PCODE),
                   reverse=True)
    for pct, name, n in share[:6]:
        print(f"    {pct:6.2f}%  {n:>6,}  {name}")
    print(f"    ...   lowest {share[-1][0]:.2f}%  {share[-1][1]}")

    if not ok:
        raise SystemExit("reconciliation FAILED")


def main():
    if "--fetch" in sys.argv:
        fetch()
    t = read()
    rows = build(t)
    check(t, rows)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(rows)
    print("\nwrote", OUT, f"({len(rows):,} rows)")


if __name__ == "__main__":
    main()
