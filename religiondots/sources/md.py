"""Moldova — BNS, Recensământul Populaţiei şi al Locuinţelor 2024, religion, down to the UAT.

Reads (or fetches) data/raw/md/ and writes data/normalized/md.csv.

**FOURTEEN RELIGIOUS AFFILIATIONS AT THE LEVEL OF THE UAT** — the first tier of local
government, an oraş/municipiu or a sat/comună — 901 of them for 2.41M people, so about
2,700 people a unit. That is the second-finest count geography on this map after Ireland's
Small Areas, and it is unusual for a country this size to publish religion there at all.

The table is sheet `5.31` of the ethnocultural annexe to the final 2024 results,
*"Populația după afilierea religioasă pe orașe (municipii) și sate (comune)"*. Sheet `5.29`
is the same thing at raion level and is read only as a cross-check.

**THE 2014 CENSUS IS THE OBVIOUS ALTERNATIVE AND IT IS THE WORSE ONE.** It is on disk here
as `Caracteristici_populatie_Comune_RPL_2014_rom_rus_eng.xls` because it was fetched first.
Three things settle it, and the reasoning is in `sources/md.md` §2:

  1. **2014 published religion at RAION level only** — 35 units against 901. Its
     commune-level sheet carries sex and age and no religion.
  2. **2014 under-enumerated badly.** BNS itself reports 2,804,801 persons actually
     enumerated against an estimated 2,998,235, so about one person in fifteen was never
     reached, and the shortfall is concentrated rather than spread.
  3. **6.4% of 2014 declared no religion answer**, against 0.75% in 2024.

What 2024 costs is real and small: **2014 named Judaism (584 people) and the Lutheran
church (2,291) as their own categories and 2024 does not**, folding both into
`Alte religii`. That is 0.1% of the country, and it is why `md.md` records the 2014 figures
even though nothing draws from them.

**THE UNITS ARE NOT ALL AT THE SAME TIER AND THE SHEET DOES NOT SAY WHICH ARE WHICH.**
Rows for the 35 raions/municipalities are interleaved with the UATs, and inside municipiul
Chişinău there is a third row — `or. Chişinău, din care pe sectoare`, 567,038 people —
which is itself the sum of the five sector rows below it. Read the sheet as delivered and
Chişinău is counted three times. The rule used here is structural, not positional:

  * a raion/municipiu is a code ending `00000` (35 of them), dropped;
  * `0101000`, the Chişinău city row, is dropped by code, keeping the five sectors;
  * everything else is a leaf, 901 of them.

`main()` asserts that the 901 leaves sum to the published national figure in every one of
the sixteen columns, and that each raion's leaves sum to its own published row. They do,
exactly. Nothing here is rounded and nothing is suppressed.

**THE UNIVERSE IS `populaţia cu reşedinţă obişnuită`** — usual residence, footnote 2 of the
sheet — and NOT the enumerated population, which is the quantity the 2014 tables carry.

**WHO IS NOT IN IT.** Footnote 1, verbatim in `EXCLUDED_NOTE` below: the census did not
cover the administrative-territorial units on the left bank of the Nistru, the municipality
of Bender, and six named places on the right bank that are administered from Tiraspol. That
is the whole of Transnistria plus a scatter of villages, and it is `gap=` in countries.py.

Usage:
    python sources/md.py --fetch    download the annexe (1.3MB) if missing
    python sources/md.py            normalise from data/raw/md/
"""

import csv
import os
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
RAW = os.path.join(ROOT, "data", "raw", "md")
OUT = os.path.join(ROOT, "data", "normalized", "md.csv")

SOURCE_ID = "md_rpl_2024"
YEAR = 2024
BASIS = "self_id"

XLSX_NAME = "Anexa_Caracteristici_Etnoculturale_RPL2024.xlsx"
XLSX_URL = ("https://statistica.gov.md/files/files/ComPresa/Recensamant/2024/Ro/"
            "Anexa_Caracteristici_Etnoculturale_RPL2024.xlsx")
MIN_BYTES = 900_000

SHEET_UAT = "5.31"          # oraşe (municipii) şi sate (comune)
SHEET_RAION = "5.29"        # regiuni de dezvoltare şi raioane/municipii

COLUMNS = ["geo_id", "geo_level", "geo_name", "source_category", "count",
           "basis", "year", "source_id", "note"]

# BNS, RPL 2024 final results, population with usual residence.
NATIONAL = 2_409_207
EXPECTED_LEAVES = 901
EXPECTED_RAIONS = 35

TOTAL_LABEL = "Total"
NOT_DECLARED = "Nu au declarat religia"
DECLARED_HEADER = "Religia declarată"
CODE_HEADER = "Cod statistic CUATM"

# The one row that is a subtotal of rows that also appear: municipiul Chişinău's city, whose
# five sectors are listed under it.  Dropped by CODE rather than by name or position.
CHISINAU_CITY = "0101000"
CHISINAU_SECTORS = ("0110000", "0120000", "0130000", "0140000", "0150000")
CHISINAU_CITY_TOTAL = 567_038

ZERO = "-"                  # ”-” = magnitudine zero, stated in the sheet's own key

# Footnote 1 of sheets 5.29 and 5.31, kept verbatim because it is the gap statement and
# countries.py's `gap=` must not drift from it.
EXCLUDED_NOTE = (
    "Datele se referă doar la UAT efectiv recenzate și nu includ unitățile "
    "administrativ-teritoriale din stânga Nistrului, municipiului Bender (inclusiv satul "
    "Proteagailovca), comuna Chițcani (inclusiv satele Merenești și Zahorna), satele "
    "Cremenciug și Gîsca din raionul Căușeni, comuna Corjova (inclusiv satul Mahala) din "
    "raionul Dubăsari, precum și satul Roghi din cadrul comunei Molovata Nouă, raionul "
    "Dubăsari.")


def fetch():
    import requests

    os.makedirs(RAW, exist_ok=True)
    dest = os.path.join(RAW, XLSX_NAME)
    if os.path.exists(dest) and os.path.getsize(dest) >= MIN_BYTES:
        print("already have", dest)
        return
    print("downloading", XLSX_URL)
    r = requests.get(XLSX_URL, headers={"User-Agent": "Mozilla/5.0"}, timeout=300)
    r.raise_for_status()
    with open(dest, "wb") as fh:
        fh.write(r.content)
    size = os.path.getsize(dest)
    # §5a: a 200 is not a download.  An error page is a 200 with HTML in it.
    if size < MIN_BYTES:
        raise SystemExit(f"{dest} is {size:,} bytes, expected at least {MIN_BYTES:,}")
    with open(dest, "rb") as fh:
        if fh.read(2) != b"PK":
            raise SystemExit(f"{dest} is not a zip container -- the server sent something "
                             "that is not an xlsx")
    print(f"  {dest} ({size:,} bytes)")


def _clean(v):
    return " ".join(str(v).split()) if v is not None else ""


def _num(v):
    """A cell of the count block.  `-` is a true zero and is stated as such in the sheet."""
    if v is None or _clean(v) in ("", ZERO):
        return 0
    return int(round(float(v)))


def _read_sheet(wb, name):
    """-> (categories, rows).

    categories is the ordered list of column labels; rows is a list of
    (code, geo_name, total, {category: count}).  Columns are located by HEADER TEXT and not
    by index: the two sheets have different leading blank columns and 5.31 carries 17 unused
    trailing ones.
    """
    ws = wb[name]
    grid = [list(r) for r in ws.iter_rows(values_only=True)]

    hdr = None
    for i, row in enumerate(grid):
        if any(_clean(v) == CODE_HEADER for v in row):
            hdr = i
            break
    if hdr is None:
        raise SystemExit(f"sheet {name}: no '{CODE_HEADER}' header row -- the workbook's "
                         "layout changed")

    top, sub = grid[hdr], grid[hdr + 1]
    code_col = next(j for j, v in enumerate(top) if _clean(v) == CODE_HEADER)
    name_col = code_col + 1
    total_col = next(j for j, v in enumerate(top) if _clean(v) == TOTAL_LABEL)
    nd_col = next(j for j, v in enumerate(top) if _clean(v) == NOT_DECLARED)
    dec_col = next(j for j, v in enumerate(top) if _clean(v) == DECLARED_HEADER)

    cats = []
    for j in range(dec_col, nd_col):
        lab = _clean(sub[j]) if j < len(sub) else ""
        if lab:
            cats.append((lab, j))
    if not cats:
        raise SystemExit(f"sheet {name}: no category labels under '{DECLARED_HEADER}'")

    rows = []
    for row in grid[hdr + 2:]:
        code = _clean(row[code_col]) if code_col < len(row) else ""
        if not code.isdigit():
            continue
        counts = {lab: _num(row[j]) for lab, j in cats}
        counts[NOT_DECLARED] = _num(row[nd_col])
        rows.append((code.zfill(7), _clean(row[name_col]), _num(row[total_col]), counts))
    return [c for c, _ in cats] + [NOT_DECLARED], rows


def main():
    import openpyxl

    if "--fetch" in sys.argv:
        fetch()

    src = os.path.join(RAW, XLSX_NAME)
    if not os.path.exists(src):
        raise SystemExit(f"missing {src} -- run with --fetch first")

    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)
    cats, rows = _read_sheet(wb, SHEET_UAT)
    rcats, rrows = _read_sheet(wb, SHEET_RAION)
    if cats != rcats:
        raise SystemExit("sheets 5.29 and 5.31 carry different category lists:\n"
                         f"  5.31 {cats}\n  5.29 {rcats}")
    print(f"{len(cats)} categories: {', '.join(cats)}")

    codes = [c for c, _, _, _ in rows]
    if len(codes) != len(set(codes)):
        raise SystemExit("sheet 5.31 has duplicate CUATM codes")

    raion_rows = {c: (n, t, k) for c, n, t, k in rows if c.endswith("00000")}
    if len(raion_rows) != EXPECTED_RAIONS:
        raise SystemExit(f"{len(raion_rows)} raion/municipiu rows in 5.31, "
                         f"expected {EXPECTED_RAIONS}")

    # 5.29 is a separate sheet with its own copy of the raion figures.  If the two disagree
    # the workbook is internally inconsistent and neither can be trusted.
    r29 = {c: (n, t, k) for c, n, t, k in rrows if c.endswith("00000")}
    if set(r29) != set(raion_rows):
        raise SystemExit("5.29 and 5.31 name different raion sets")
    for c in sorted(r29):
        if r29[c][1] != raion_rows[c][1] or r29[c][2] != raion_rows[c][2]:
            raise SystemExit(f"raion {c} differs between 5.29 and 5.31")
    print(f"{len(raion_rows)} raions/municipalities, identical in sheets 5.29 and 5.31")

    drop = set(raion_rows) | {CHISINAU_CITY}
    leaves = [(c, n, t, k) for c, n, t, k in rows if c not in drop]
    if len(leaves) != EXPECTED_LEAVES:
        raise SystemExit(f"{len(leaves)} leaf UATs, expected {EXPECTED_LEAVES}")

    # ---- 1. the five Chişinău sectors ARE the city row, so dropping it loses nothing ----
    sect = sum(t for c, _, t, _ in leaves if c in CHISINAU_SECTORS)
    if sect != CHISINAU_CITY_TOTAL:
        raise SystemExit(f"the five Chişinău sectors sum to {sect:,}, but the "
                         f"`or. Chişinău` row says {CHISINAU_CITY_TOTAL:,} -- dropping "
                         "that row would lose people")
    print(f"Chişinău: 5 sectors sum to {sect:,}, exactly the dropped city row")

    # ---- 2. every column of the 901 leaves sums to the published national figure ----
    # The national row is the sheet's own, found by its name and its total rather than by
    # position: it has no CUATM code, so _read_sheet skipped it.
    grid = [list(r) for r in wb[SHEET_UAT].iter_rows(values_only=True)]
    hdr = next(i for i, r in enumerate(grid) if any(_clean(v) == CODE_HEADER for v in r))
    top, sub = grid[hdr], grid[hdr + 1]
    nd_col = next(j for j, v in enumerate(top) if _clean(v) == NOT_DECLARED)
    dec_col = next(j for j, v in enumerate(top) if _clean(v) == DECLARED_HEADER)
    total_col = next(j for j, v in enumerate(top) if _clean(v) == TOTAL_LABEL)
    name_col = next(j for j, v in enumerate(top) if _clean(v) == CODE_HEADER) + 1
    colof = {_clean(sub[j]): j for j in range(dec_col, nd_col) if _clean(sub[j])}
    colof[NOT_DECLARED] = nd_col

    natl_row = None
    for row in grid[hdr + 2:]:
        code = _clean(row[name_col - 1]) if name_col - 1 < len(row) else ""
        if not code and _clean(row[name_col]) == TOTAL_LABEL:
            natl_row = row
            break
    if natl_row is None:
        raise SystemExit(f"could not find the uncoded '{TOTAL_LABEL}' row in {SHEET_UAT}")

    print("\ncolumn sums over the 901 leaves against the sheet's own Total row:")
    bad = 0
    tot = sum(t for _, _, t, _ in leaves)
    if tot != _num(natl_row[total_col]) or tot != NATIONAL:
        raise SystemExit(f"leaves total {tot:,}, published {NATIONAL:,}")
    print(f"  {'Total':<40} {tot:>9,}  OK")
    for cat in cats:
        s = sum(k[cat] for _, _, _, k in leaves)
        p = _num(natl_row[colof[cat]])
        ok = s == p
        bad += not ok
        print(f"  {cat:<40} {s:>9,}  {'OK' if ok else 'DIFFERS from %s' % f'{p:,}'}")
    if bad:
        raise SystemExit(f"{bad} columns do not reconcile")

    # ---- 3. every raion's own leaves sum to its published row, column by column ----
    for code, (rname, rtot, rk) in sorted(raion_rows.items()):
        kids = [(c, n, t, k) for c, n, t, k in leaves if c[:2] == code[:2]]
        if not kids:
            raise SystemExit(f"raion {code} {rname} has no UATs under it")
        if sum(t for _, _, t, _ in kids) != rtot:
            raise SystemExit(f"{rname}: leaves sum to "
                             f"{sum(t for _, _, t, _ in kids):,}, published {rtot:,}")
        for cat in cats:
            if sum(k[cat] for _, _, _, k in kids) != rk[cat]:
                raise SystemExit(f"{rname}, {cat}: leaves do not sum to the raion row")
    print(f"\nall {len(raion_rows)} raions reconcile against their own UATs, "
          f"in all {len(cats)} columns")

    # ---- 4. write ----
    raion_of = {}
    for code, (rname, _, _) in raion_rows.items():
        for c, _, _, _ in leaves:
            if c[:2] == code[:2]:
                raion_of[c] = rname

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    n = 0
    with open(OUT, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(COLUMNS)
        for code, name, total, k in leaves:
            note = (f"raion={raion_of[code]}; universe=populatia cu resedinta obisnuita; "
                    f"RPL 2024 table 5.31")
            for cat in cats:
                if k[cat] <= 0:
                    continue
                w.writerow([code, "uat", name, cat, k[cat], BASIS, YEAR, SOURCE_ID, note])
                n += 1
        for code, (rname, rtot, rk) in sorted(raion_rows.items()):
            note = "RPL 2024 table 5.29; cross-check level, not drawn"
            for cat in cats:
                if rk[cat] <= 0:
                    continue
                w.writerow([code, "raion", rname, cat, rk[cat], BASIS, YEAR,
                            SOURCE_ID, note])
                n += 1
        note = EXCLUDED_NOTE
        for cat in cats:
            s = sum(k[cat] for _, _, _, k in leaves)
            if s > 0:
                w.writerow(["MD", "country", "Republica Moldova", cat, s, BASIS, YEAR,
                            SOURCE_ID, note])
                n += 1

    print(f"\nwrote {OUT} ({n:,} rows; {len(leaves):,} UATs, {len(raion_rows)} raions, "
          "1 country)")
    declared = tot - sum(k[NOT_DECLARED] for _, _, _, k in leaves)
    print(f"  {tot:,} people with usual residence; {declared:,} declared an affiliation "
          f"({declared / tot:.2%})")


if __name__ == "__main__":
    main()
