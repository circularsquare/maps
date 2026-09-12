# -*- coding: utf-8 -*-
"""Shared parts of the Korail 광역철도 layers.

`build_donghae.py` and `build_gyeongchun.py` both take Korail's published
per-station 승하차 for one commuter service, fit a doubly-constrained gravity
model to a published mean trip length, and cumulate the result into per-segment
loads. What differs is where the corridor comes from -- 동해선's is an intercity
chain `solve.py` already builds, ITX-청춘's is assembled out of the Seoul metro
layer -- so only the counts reader and the decay fit live here.

Both sit on `info.korail.com` board 425; see README.md.
"""

import build_busan as BB


def read_counts(path, line=None, alias=None):
    """{station: (승차, 하차)} for the year, from a Korail 광역철도 sheet.

    `line` filters on the 노선명 column where the edition has one -- 2025 and
    2026 do, 2022 and 2023 do not, and on those a filter would silently return
    nothing, so asking for one there raises instead.

    **Columns are found by their header, not by position, and that is not
    fussiness.** Three editions, three layouts: 2022 is
    역명구분 | 역명 | 구분 | twelve months | 합계, 2023 inserts a blank column
    before 구분 and shifts everything one right, and 2025-26 rename the sheet
    승하차인원 and put 노선명 in front of the lot. Reading 2022 with 2023's
    offsets matches no station at all and sums to a confident zero.
    """
    import openpyxl

    alias = alias or {}
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[next(n for n in wb.sheetnames if n.startswith("승하차"))]

    rows = ws.iter_rows(values_only=True)
    col = {}
    for row in rows:
        head = {str(v).strip(): i for i, v in enumerate(row) if v is not None}
        if "역명" in head and "구분" in head and "합계" in head:
            col = head
            break
    if not col:
        wb.close()
        raise SystemExit("%s: no 역명/구분/합계 header row -- the layout has "
                         "changed again, look at the sheet" % path)
    if line and "노선명" not in col:
        wb.close()
        raise SystemExit("%s: no 노선명 column, so it cannot be filtered to "
                         "%s -- use a 2025 or later edition" % (path, line))
    i_nm, i_kind, i_sum = col["역명"], col["구분"], col["합계"]
    i_line = col.get("노선명")

    out = {}
    for row in rows:
        if len(row) <= max(i_nm, i_kind, i_sum):
            continue
        if line is not None:
            if len(row) <= i_line:
                continue
            got = row[i_line]
            if not got or str(got).strip() != line:
                continue
        nm, kind, v = row[i_nm], row[i_kind], row[i_sum]
        if not nm or kind not in ("승차", "하차"):
            continue
        if not isinstance(v, (int, float)):
            continue
        nm = alias.get(str(nm).strip(), str(nm).strip())
        on, off = out.get(nm, (0.0, 0.0))
        out[nm] = (on + v, off) if kind == "승차" else (on, off + v)
    wb.close()
    if not out:
        raise SystemExit("%s: header found but no 승차/하차 rows read%s"
                         % (path, (" for " + line) if line else ""))
    return out


def fit_od_beta(cost, board, alight, beta, iters=10000, tol=1e-5):
    """`build_busan.fit_od` with the distance multiplier free to be positive.

    That file writes the weight as `exp(-cost / decay)`, which can only make
    nearer pairs likelier, and rejects a non-positive decay. For a stopping
    service that is right. For a **limited-stop** one it is the wrong family
    outright: ITX-청춘 shares its rails with the 경춘선 전동차, so the short
    hops go to the stopping train and what is left rides a long way. Its
    published mean trip is 73 km on a 98 km route, and no decay reproduces that
    -- the most an independence fit gives is 54.

    Written as `exp(beta * cost)`, which is the same maximum-entropy model with
    the Lagrange multiplier on distance left unsigned: negative is a decay,
    positive a preference for distance. `beta = -1 / decay` recovers the other
    form exactly. The shift by `cost.max()` only keeps the exponential in range
    and cancels in the balancing.
    """
    import numpy as np

    if np.any(board < 0) or np.any(alight < 0):
        raise ValueError("negative demand")
    if min(board.sum(), alight.sum()) <= 0:
        raise ValueError("empty demand")
    target = alight * (board.sum() / alight.sum())
    od = np.exp(beta * (cost - cost.max()))
    np.fill_diagonal(od, 0.0)
    for _ in range(iters):
        rs = od.sum(axis=1)
        od *= np.divide(board, rs, out=np.zeros_like(board), where=rs > 0)[:, None]
        cs = od.sum(axis=0)
        od *= np.divide(target, cs, out=np.zeros_like(target), where=cs > 0)[None, :]
        err = max(np.max(np.abs(od.sum(axis=1) - board)),
                  np.max(np.abs(od.sum(axis=0) - target)))
        if err < tol:
            return od
    raise ValueError("OD balancing did not converge")


def fit_beta(cost, board, alight, want_km, span=20.0):
    """The distance multiplier whose fitted OD has the published mean trip.

    Monotone increasing in beta. Returns (beta, od). Use this where the service
    is limited-stop; `fit_decay` is the equivalent for a stopping one.

    **The bounds are scaled to the costs, not fixed.** beta multiplies a
    distance, so what counts as a wide search depends on how long the line is:
    a flat +/-2 per km is a factor of exp(196) on a 98 km route, which
    overflows, and a flat +/-0.01 would not move a two-km metro at all. `span`
    is the exponent range to sweep, and the bounds come out as +/-span divided
    by the longest pair. Narrowed further if the ends still will not balance.
    """
    import numpy as np

    far = float(np.max(cost)) or 1.0
    lo, hi = -span / far, span / far

    def mean_km(b):
        od = fit_od_beta(cost, board, alight, b)
        tot = od.sum()
        return (float((od * cost).sum() / tot) if tot else 0.0), od

    while True:
        try:
            lo_km, _ = mean_km(lo)
            hi_km, _ = mean_km(hi)
            break
        except ValueError:
            lo, hi = lo / 2.0, hi / 2.0
            if hi < 1e-9:
                raise SystemExit("no multiplier gives a balanced OD; check "
                                 "the boardings and alightings are non-zero")
    if not (lo_km <= want_km <= hi_km):
        raise SystemExit(
            "a mean trip of %.2f km is outside what these stations can produce "
            "even with the multiplier unsigned (%.2f to %.2f km) -- the "
            "marginals or the span are wrong, not the target"
            % (want_km, lo_km, hi_km))
    od = None
    for _ in range(80):
        mid = 0.5 * (lo + hi)
        got, od = mean_km(mid)
        if abs(got - want_km) < 1e-4:
            break
        if got < want_km:
            lo = mid
        else:
            hi = mid
    return 0.5 * (lo + hi), od


def fit_decay(cost, board, alight, want_km, lo=0.5, hi=400.0):
    """The decay whose fitted OD has the published mean trip length.

    Monotone in the decay -- a longer scale puts more weight on distant pairs --
    so a bisection is enough. Returns (decay, od).
    """
    def mean_km(decay):
        od, _, _ = BB.fit_od(cost, board, alight, decay)
        tot = od.sum()
        return (float((od * cost).sum() / tot) if tot else 0.0), od

    # A decay much shorter than the gap between neighbouring stations makes
    # exp(-cost/decay) underflow to an all-zero matrix, and the IPF then cannot
    # balance and raises. Where the stations are far apart -- ITX-청춘 stops
    # every few tens of km -- that rules out the nominal floor, so walk it up
    # until the fit is defined rather than starting from a number that is not.
    while lo < hi:
        try:
            lo_km, _ = mean_km(lo)
            break
        except ValueError:
            lo *= 2.0
    else:
        raise SystemExit("no decay between the bounds gives a balanced OD; "
                         "check the boardings and alightings are non-zero")
    hi_km, _ = mean_km(hi)
    if not (lo_km <= want_km <= hi_km):
        raise SystemExit(
            "a mean trip of %.2f km is outside what these stations can "
            "produce (%.2f to %.2f km) -- the marginals or the span are wrong, "
            "not the target" % (want_km, lo_km, hi_km))
    od = None
    for _ in range(60):
        mid = 0.5 * (lo + hi)
        got, od = mean_km(mid)
        if abs(got - want_km) < 1e-4:
            break
        if got < want_km:
            lo = mid
        else:
            hi = mid
    return 0.5 * (lo + hi), od
