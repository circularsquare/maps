"""Sint Maarten, Population and Housing Census 2011 (STAT, Department of Statistics): persons in
private households by the language most spoken in the household -> data/normalized/sx.csv.

    python sources/sx_census.py [--fetch]

THE TABLE. Census 2011 Table F-07, *Private households by most spoken language in the
household*, households and persons, for the whole country (the Dutch side of the island). STAT
publishes it as the "Housing 2011" workbook on its census tables page
(stats.sintmaartengov.org/tables.php?division=social&topic=cen ->
download.php?type=census&nummer=3, sheet table_f_07) and prints the same table in the census
report "Census 2011" (reports.php?cat=CEN, download.php?type=rep&section=CEN&nummer=9, p.216 of
the PDF). One answer per household, given to every member; 25 named answers plus "no response",
33,162 persons in 12,854 private households (the 2011 population is 33,609; the other 447 lived
in institutions and were not asked).

Why 2011 and not 2022: Census 2022 (November 2022) published language only as households' shares
in nine groups ("Sint Maarten Population 2022", reports.php?cat=CEN, download.php?type=rep&
section=CEN&nummer=18, p.58), not persons and no counts. sources/sx.md §1.

THE CHECKS: the 25 language rows plus no response sum to the printed persons total 33,162 and the
households to 12,854 within 3 (they come to 33,160 and 12,853, the same in the report's print); each row's printed persons share agrees with count / 33,162 within
0.06; the 2011 household shares quoted by the 2022 report (English 67.2, Spanish 13.3, Creole
10.3, Dutch 3.7, Papiamentu 1.5, Hindi 1.3, Chinese 0.4, French 0.3) match this table's; and
Table D-13 (persons attending day-school by household language, the Education workbook, nummer=1)
is a subset of F-07 row by row: no language has more school attendees than persons.
"""
import csv
import sys
import urllib.request
from pathlib import Path

import openpyxl

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = Path(__file__).resolve().parent.parent
RAW = HERE / "data" / "raw" / "sx"
OUT = HERE / "data" / "normalized" / "sx.csv"
BASE = "https://stats.sintmaartengov.org/download.php?type=census&nummer="
FILES = {"housing": (3, RAW / "census2011_housing.xlsx"),
         "education": (1, RAW / "census2011_education.xlsx")}
UA = ("Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) "
      "Chrome/124.0 Safari/537.36")
SOURCE_ID = "stat_sx_census2011_tableF07"
PERSONS, HOUSEHOLDS = 33_162, 12_854
# the 2011 column of the 2022 report's "Language most spoken in households" chart (p.58)
SHARES_2022_REPORT = {"english": 67.2, "spanish": 13.3, "french creole": 10.3, "dutch": 3.7,
                      "papiamentu": 1.5, "hindi": 1.3, "chinese": 0.4, "french": 0.3}


def fetch():
    RAW.mkdir(parents=True, exist_ok=True)
    for num, path in FILES.values():
        req = urllib.request.Request(BASE + str(num), headers={"User-Agent": UA})
        data = urllib.request.urlopen(req, timeout=60).read()
        if data[:2] != b"PK":
            raise SystemExit(f"nummer={num}: not an xlsx ({data[:40]!r})")
        path.write_bytes(data)
        print(f"  {path.name}: {len(data):,} bytes")


def num(v):
    if isinstance(v, (int, float)):
        return v
    return float(str(v).replace("\xa0", "").replace(",", "").strip())


def read_f07():
    wb = openpyxl.load_workbook(FILES["housing"][1], read_only=True)
    rows = list(wb["table_f_07"].iter_rows(values_only=True))
    if rows[2][0] != "Language" or rows[2][3] != "Persons Absolute":
        raise SystemExit(f"table_f_07 header moved: {rows[2]}")
    out, total = [], None
    for r in rows[3:]:
        if r[0] is None:
            continue
        lab = str(r[0]).strip()
        rec = (lab, int(num(r[1])), num(r[2]), int(num(r[3])), num(r[4]))
        if lab == "Total":
            total = rec
        else:
            out.append(rec)
    return out, total


def read_d13():
    wb = openpyxl.load_workbook(FILES["education"][1], read_only=True)
    out = {}
    for r in wb["table_d_13"].iter_rows(values_only=True, min_row=5):
        if r[0] and isinstance(r[-1], (int, float)) and str(r[0]).strip().lower() not in ("total", "grand total"):
            out[str(r[0]).strip().lower()] = int(r[-1])
    return out


def main():
    if "--fetch" in sys.argv:
        fetch()
    rows, total = read_f07()
    if len(rows) != 26:
        raise SystemExit(f"F-07: {len(rows)} rows, expected 25 languages + no response")
    if total[1] != HOUSEHOLDS or total[3] != PERSONS:
        raise SystemExit(f"F-07 total row {total}")
    hh, pp = sum(r[1] for r in rows), sum(r[3] for r in rows)
    print(f"  F-07: households {hh:,} (printed {HOUSEHOLDS:,}), persons {pp:,} (printed {PERSONS:,})")
    # The rows sum to 12,853 and 33,160, one household and two persons short of the printed
    # totals, in the workbook and the report alike; allow 3.
    if abs(hh - HOUSEHOLDS) > 3 or abs(pp - PERSONS) > 3:
        raise SystemExit("F-07 rows do not sum to the printed totals")
    worst = max(abs(r[3] / PERSONS * 100 - r[4]) for r in rows)
    print(f"  persons shares: largest gap to the printed % {worst:.3f}")
    if worst > 0.06:
        raise SystemExit("a persons share disagrees with its count")
    byl = {r[0]: r for r in rows}
    for lab, s in SHARES_2022_REPORT.items():
        if abs(byl[lab][2] - s) > 0.05:
            raise SystemExit(f"{lab}: F-07 household share {byl[lab][2]} vs 2022 report {s}")
    print("  household shares match the 2011 column of the 2022 report (8 languages)")
    d13 = read_d13()
    d13.pop("not reported", None)
    over = {k: (v, byl[k][3]) for k, v in d13.items() if k in byl and v > byl[k][3]}
    missing = sorted(set(d13) - set(byl))
    print(f"  D-13 (day-school attendees): {len(d13)} languages, {sum(d13.values()):,} people; "
          f"over F-07: {over or 'none'}; not in F-07: {missing or 'none'}")
    if over or missing:
        raise SystemExit("D-13 is not a subset of F-07")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    with open(OUT, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(["geo_id", "geo_level", "geo_name", "source_category", "count", "tier",
                    "source_id", "year", "note"])
        for lab, h, hs, p, ps in rows:
            w.writerow(["SX", "national", "Sint Maarten", lab, p, "measured", SOURCE_ID, 2011,
                        f"persons in {h:,} private households"])
    print(f"  wrote {OUT.relative_to(HERE)}: {len(rows)} rows, {pp:,} persons")


if __name__ == "__main__":
    main()
