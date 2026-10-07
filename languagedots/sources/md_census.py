"""Moldova: BNS, Recensământul Populaţiei şi al Locuinţelor 2024, mother tongue by UAT.

    python sources/md_census.py --fetch    download the ethnocultural annexe (1.3 MB) if missing
    python sources/md_census.py            -> data/normalized/md.csv

THE TABLE. Sheet `5.15` of `Anexa_Caracteristici_Etnoculturale_RPL2024.xlsx`, "Populaţia după
limba maternă pe oraşe (municipii) şi sate (comune)": mother tongue for the 901 UATs (towns,
communes and Chişinău's five sectors), the same workbook and the same units religiondots draws
religion on (its sources/md.py reads sheet 5.31 the same way). Eight answers: Moldovenească,
Română, Ucraineană, Rusă, Găgăuză, Bulgară, Romani (Ţigănească), Altă limbă; plus "Nu au
declarat limba maternă" (not drawn). The `Moldovenească sau Română` column is the sheet's own
sum of the first two (footnote 3) and is read only as a check.

ROWS. Interleaved in the sheet: 35 raions/municipalities (code ending 00000), the Chişinău city
row `0101000` (the sum of the five sector rows under it), and 901 leaves. The rule is religiondots'
and is structural: drop codes ending 00000 and 0101000, keep the rest; codes lost their leading
zero in some cells and are zero-filled to 7.

CHECKS, all asserted:
  * the 901 leaves sum to the sheet's uncoded `Total` row in every column, and to 2,409,207;
  * every raion row equals the sum of its own leaves, in every column;
  * sheet 5.13 (raions, a separate sheet) carries the same raion figures as 5.15;
  * sheet 5.10 (national, 2024 beside 2014) carries the same national figures and the same
    eight answers (it names nothing inside `Altă limbă`; only sheet 5.37, for ages 3+, does);
  * the five Chişinău sectors sum to the dropped city row;
  * Moldovenească + Română = the sheet's `Moldovenească sau Română` in every leaf.

UNIVERSE: population with usual residence. NOT COVERED (footnote 1): the left bank of the Nistru,
Bender, and six right-bank places administered from Tiraspol.
"""
import csv
import os
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
RAW = os.path.join(ROOT, "data", "raw", "md")
OUT = os.path.join(ROOT, "data", "normalized", "md.csv")

XLSX_NAME = "Anexa_Caracteristici_Etnoculturale_RPL2024.xlsx"
XLSX_URL = ("https://statistica.gov.md/files/files/ComPresa/Recensamant/2024/Ro/"
            "Anexa_Caracteristici_Etnoculturale_RPL2024.xlsx")
MIN_BYTES = 900_000
SOURCE_ID = "md_rpl_2024_t5.15"
YEAR = 2024
NATIONAL = 2_409_207
EXPECTED_LEAVES = 901
EXPECTED_RAIONS = 35
CHISINAU_CITY = "0101000"
CHISINAU_SECTORS = ("0110000", "0120000", "0130000", "0140000", "0150000")

# column index in sheets 5.15 and 5.13 -> label (header row 7/8, checked against the sheet below)
SUM_COL = 4                    # Moldovenească sau Română, the sheet's own sum of cols 5 and 6
COLS = {5: "Moldovenească", 6: "Română", 7: "Ucraineană", 8: "Rusă", 9: "Găgăuză",
        10: "Bulgară", 11: "Romani (Ţigănească)", 12: "Altă limbă",
        13: "Nu au declarat limba maternă"}
COLUMNS = ["geo_id", "geo_level", "geo_name", "source_category", "count", "tier", "year",
           "source_id", "note"]


def fetch():
    import requests
    os.makedirs(RAW, exist_ok=True)
    dest = os.path.join(RAW, XLSX_NAME)
    if os.path.exists(dest) and os.path.getsize(dest) >= MIN_BYTES:
        print("already have", dest)
        return
    print("downloading", XLSX_URL)
    r = requests.get(XLSX_URL, headers={"User-Agent": "Mozilla/5.0"}, timeout=300)
    r.raise_for_status()
    if len(r.content) < MIN_BYTES or r.content[:2] != b"PK":
        raise SystemExit(f"got {len(r.content):,} bytes that are not an xlsx")
    with open(dest, "wb") as fh:
        fh.write(r.content)
    print(f"  {dest} ({len(r.content):,} bytes)")


def clean(v):
    return " ".join(str(v).split()) if v is not None else ""


def num(v):
    if v is None or clean(v) in ("", "-"):
        return 0
    return int(round(float(v)))


def norm_label(s):
    # the sheets mix cedilla and comma-below ş/ţ and pad with spaces
    return clean(s).replace("ș", "ş").replace("ț", "ţ").replace("Ț", "Ţ").replace("Ș", "Ş")


def read(ws):
    """-> (total_row_counts, {code: (name, total, {label: n}, sum_col)})"""
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    # header check: the labels this file assumes are where it assumes them
    h7, h8 = grid[7], grid[8]
    want = {7: "Ucraineană", 8: "Rusă", 9: "Găgăuză", 10: "Bulgară", 11: "Romani (Ţigănească)",
            12: "Altă limbă"}
    for j, lab in want.items():
        if norm_label(h7[j]) != norm_label(lab):
            raise SystemExit(f"{ws.title}: column {j} is {h7[j]!r}, expected {lab!r}")
    if norm_label(h8[5]) != "Moldovenească" or norm_label(h8[6]) != "Română":
        raise SystemExit(f"{ws.title}: Moldovenească/Română not in columns 5/6: {h8[5:7]}")
    if "Moldovenească sau Română" not in norm_label(h7[SUM_COL]):
        raise SystemExit(f"{ws.title}: column {SUM_COL} is {h7[SUM_COL]!r}")
    # 5.15 puts this header on row 6, 5.13 on row 7
    if not any("Nu au declarat" in clean(grid[i][13]) for i in (6, 7)):
        raise SystemExit(f"{ws.title}: 'Nu au declarat' is not column 13")

    total, rows = None, {}
    for r in grid[10:]:
        code, name = clean(r[1]), clean(r[2])
        if not code and name == "Total":
            total = {COLS[j]: num(r[j]) for j in COLS} | {"_total": num(r[3])}
            continue
        if not code.isdigit():
            continue
        code = code.zfill(7)
        if code in rows:
            raise SystemExit(f"{ws.title}: duplicate code {code}")
        rows[code] = (name, num(r[3]), {COLS[j]: num(r[j]) for j in COLS}, num(r[SUM_COL]))
    if total is None:
        raise SystemExit(f"{ws.title}: no uncoded Total row")
    return total, rows


def main():
    import openpyxl
    if "--fetch" in sys.argv:
        fetch()
    src = os.path.join(RAW, XLSX_NAME)
    if not os.path.exists(src):
        raise SystemExit(f"missing {src}; run with --fetch")
    wb = openpyxl.load_workbook(src, read_only=True, data_only=True)

    total, rows = read(wb["5.15"])
    raions = {c: v for c, v in rows.items() if c.endswith("00000")}
    if len(raions) != EXPECTED_RAIONS:
        raise SystemExit(f"{len(raions)} raion rows, expected {EXPECTED_RAIONS}")
    leaves = {c: v for c, v in rows.items() if c not in raions and c != CHISINAU_CITY}
    if len(leaves) != EXPECTED_LEAVES:
        raise SystemExit(f"{len(leaves)} leaves, expected {EXPECTED_LEAVES}")

    # every leaf: its columns sum to its total; Moldovenească + Română = the sheet's own sum
    for c, (n, t, k, s) in leaves.items():
        if sum(k.values()) != t:
            raise SystemExit(f"{c} {n}: columns sum to {sum(k.values())}, total {t}")
        if k["Moldovenească"] + k["Română"] != s:
            raise SystemExit(f"{c} {n}: Moldovenească + Română != the sheet's sum column")
    print(f"{len(leaves)} leaves: each row's columns sum to its total, and Moldovenească + "
          "Română equals the sheet's own sum column in every one")

    # national
    tot = sum(t for _, t, _, _ in leaves.values())
    if tot != NATIONAL or tot != total["_total"]:
        raise SystemExit(f"leaves total {tot:,}; sheet {total['_total']:,}; expected {NATIONAL:,}")
    print(f"\ncolumn sums over the 901 leaves against the sheet's Total row ({tot:,}):")
    for lab in COLS.values():
        s = sum(k[lab] for _, _, k, _ in leaves.values())
        print(f"  {lab:<32} {s:>10,}  {'OK' if s == total[lab] else 'DIFFERS ' + str(total[lab])}")
        if s != total[lab]:
            raise SystemExit("national column does not reconcile")

    # raions against their own leaves, and against sheet 5.13
    for rc, (rn, rt, rk, _) in raions.items():
        kids = [v for c, v in leaves.items() if c[:2] == rc[:2]]
        if not kids or sum(t for _, t, _, _ in kids) != rt:
            raise SystemExit(f"raion {rc} {rn}: leaves do not sum to its total")
        for lab in COLS.values():
            if sum(k[lab] for _, _, k, _ in kids) != rk[lab]:
                raise SystemExit(f"raion {rc} {rn}, {lab}: leaves do not sum to the raion row")
    t13, r13 = read(wb["5.13"])
    r13 = {c: v for c, v in r13.items() if c.endswith("00000")}
    if set(r13) != set(raions):
        raise SystemExit("5.13 and 5.15 name different raions")
    for c in raions:
        if r13[c][1] != raions[c][1] or r13[c][2] != raions[c][2]:
            raise SystemExit(f"raion {c} differs between 5.13 and 5.15")
    print(f"\nall {len(raions)} raions equal the sum of their leaves in every column, and "
          "sheet 5.13 carries the same figures")

    sect = sum(leaves[c][1] for c in CHISINAU_SECTORS)
    if sect != rows[CHISINAU_CITY][1]:
        raise SystemExit("the Chişinău sectors do not sum to the city row")
    print(f"Chişinău: the five sectors sum to {sect:,}, the dropped city row")

    # sheet 5.10: the national table, a separate copy; same answers, same figures
    grid = [list(r) for r in wb["5.10"].iter_rows(values_only=True)]
    natl = {}
    for r in grid[10:]:
        lab = norm_label(r[1])
        if lab and isinstance(r[2], (int, float)):
            # the sheet repeats the list for Total, then Urban, then Rural; keep the first
            if lab == "Total" and natl:
                break
            natl[lab] = int(r[2])
    big = {"Total", "Moldovenească sau Română", "Moldovenească", "Română", "Ucraineană", "Rusă",
           "Găgăuză", "Bulgară", "Romani (Ţigănească)", "Nu au declarat limba maternă"}
    for lab in ("Moldovenească", "Română", "Ucraineană", "Rusă", "Găgăuză", "Bulgară",
                "Nu au declarat limba maternă"):
        if natl.get(lab) != total[lab]:
            raise SystemExit(f"5.10 {lab} = {natl.get(lab)}, 5.15 says {total[lab]}")
    others = {k: v for k, v in natl.items() if k not in big and not k.startswith("Romani")}
    print("\nsheet 5.10 (national) agrees; its rows beyond 5.15's named seven:")
    for k, v in sorted(others.items(), key=lambda kv: -kv[1]):
        print(f"  {k:<32} {v:>7,}")
    if sum(others.values()) != total["Altă limbă"]:
        raise SystemExit(f"5.10's other languages sum to {sum(others.values()):,}, "
                         f"5.15's Altă limbă is {total['Altă limbă']:,}")
    print(f"  {'sum':<32} {sum(others.values()):>7,}  = 5.15's Altă limbă")

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    n = 0
    with open(OUT, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(COLUMNS)
        for c, (name, _, k, _) in leaves.items():
            raion = raions[c[:2] + "00000"][0]
            for lab, v in k.items():
                if v > 0:
                    w.writerow([c, "uat", name, lab, v, "measured", YEAR, SOURCE_ID,
                                f"raion={raion}; RPL 2024 table 5.15"])
                    n += 1
        for c, (name, _, k, _) in raions.items():
            for lab, v in k.items():
                if v > 0:
                    w.writerow([c, "raion", name, lab, v, "measured", YEAR, SOURCE_ID,
                                "RPL 2024 table 5.13/5.15; cross-check level, not drawn"])
                    n += 1
        for lab in COLS.values():
            if total[lab] > 0:
                w.writerow(["MD", "country", "Republica Moldova", lab, total[lab], "measured",
                            YEAR, SOURCE_ID, "usual residence; excludes the left bank and Bender"])
                n += 1
    print(f"\nwrote {OUT} ({n:,} rows)")


if __name__ == "__main__":
    main()
