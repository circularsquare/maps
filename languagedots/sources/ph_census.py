"""Philippines 2020 census (CPH), ethnicity by province and HUC, read as language
-> data/normalized/ph.csv and data/normalized/ph_retention.csv.

    python sources/ph_census.py [--fetch]

WHY ETHNICITY. The 2020 CPH asked language too ("What is the language/dialect generally spoken
at home by members of this household?"), but PSA published it nationally only (one row per
language, counting HOUSEHOLDS), plus a few regional special releases (Cordillera, Oriental
Mindoro) with each province's top five. The 2000 CPH asked the same household question and
printed it by province (Report No. 2, Table 27), but in about a hundred provincial PDFs behind
psa.gov.ph's Cloudflare wall, six of them on the Wayback Machine. sources/ph.md has the search.
So the counts are the 2020 ethnicity table ("ethnicity by descent/blood relation/
consanguinity", 290 categories, every household member, full count), on the same 117 province /
HUC units religiondots draws religion on. Every row is `derived` (AGENT_BRIEF §2, ethnicity).

RETENTION (AGENT_BRIEF §2). The national household-language table is the check. For each group
the census names in both tables (or a set of them, where one table splits what the other
lumps), the households expected if every member's household spoke the group's language are
sum over units of persons x (households / household population) in that unit; the share the
language table actually reports is the group's retention, capped at 1. Groups below 1 lose the
rest of their people to the unit's lingua franca, removed first from units where the group is
not the largest ethnicity (the diaspora: Bikolanos in Manila speak Tagalog at home, Bikolanos
in Albay speak Bikol), then from the rest in proportion. The lingua franca of a unit is the
largest of a fixed list (Tagalog, Cebuano, Hiligaynon, Ilocano, Bikol, Waray, Tausug, Maranao,
Maguindanao, Chavacano) other than the group itself. Moved rows are `modelled`. The language
table counts households, and a household of mixed descent counts once, so this is approximate;
ph_retention.csv carries every figure.

PLACE-DEPENDENT LABEL. `Other Local Ethnicity` is 88% of Aklan and 23% of Zambales: the
ethnicity list has no Aklanon and no Sambal (non-IP regional peoples), and the language table
has 133,121 households of "Bukidnon/Binukid-Akeanon/Aklanon", about Aklan's household count. So
in those two units it is written as `Other Local Ethnicity (Aklan)` / `(Zambales)` and mapped to
Aklanon and Sambal; everywhere else it stays the Philippine remainder. Likewise `Dumagat` in
Mindanao (the Lumad name for lowland Visayan settlers) is written `Dumagat (Mindanao)`, and
`Bukidnon` in the Visayas (Panay and Negros highlanders, not Mindanao's Binukid) `Bukidnon
(Visayas)`.

CHECKS (all must pass, else nothing is written):
  1. 290 ethnicity columns (one name twice, `Buhid Mangyan`; the second is written `[2]`);
     every row's columns sum to its household population
  2. the 135 rows join religiondots' ph.csv row for row: same label, same household population,
     to the person; the 117 fine units (province, HUC, Pateros, Interim Province) sum to the
     national row per column
  3. the press release's top ten and total (Table 1, 3 July 2023) match the national row
  4. the OpenSTAT household table joins the same 135 rows with the same household population
  5. the national language table's categories sum to its total, 26,388,654 households
  6. after the shift, every unit still sums to its household population less `Not Reported`
"""
import argparse
import json
import re
import sys
import unicodedata
import urllib.request
from pathlib import Path

import openpyxl
import pandas as pd

HERE = Path(__file__).resolve().parent.parent
RAW = HERE / "data" / "raw" / "ph"
OUT = HERE / "data" / "normalized" / "ph.csv"
OUT_RET = HERE / "data" / "normalized" / "ph_retention.csv"
RD_NORM = HERE.parent / "religiondots" / "data" / "normalized" / "ph.csv"   # read-only

WB = "http://web.archive.org/web/{ts}id_/https://psa.gov.ph/system/files/phcd/{f}"
FILES = {
    "ph2020_ethnicity.xlsx": ("20230812114311", "Ethnicity_Statistical%20Table.xlsx"),
    "ph2020_ethnicity_pr.pdf": ("20230812114629", "PR%20on%20Ethnicity.pdf"),
    "ph2020_ethnicity_tn.pdf": ("20230812114348", "PR%20on%20Ethnicity_Technical%20Notes.pdf"),
    "ph2020_language.xlsx": ("20230812114327", "PR_Statistical%20Tables_Language_030323_PMMJ_CRD.xlsx"),
    "ph2020_language_pr.pdf": ("20230812114250",
                               "02%20-Press%20Release%20on%20Language_Dialect%20Generally%20Spoken"
                               "%20at%20Home_030323_PMMJ_CRD.pdf"),
    "ph2020_language_tn.pdf": ("20230813061906",
                               "Technical%20Notes%20for%20PR%20on%20Language-Dialect%20Generally%20"
                               "Spoken%20at%20Home%20(2020CPH).pdf"),
}
OPENSTAT = "https://openstat.psa.gov.ph/PXWeb/api/v1/en/DB/1A/PO_2020/0011A6DPHH0.px"
HH_JSON = RAW / "ph2020_households_openstat.json"
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
                    "(KHTML, like Gecko) Chrome/126.0 Safari/537.36"}

N_ETH = 290
HP_NATIONAL = 108_667_043
HH_NATIONAL = 26_388_654
PRESS_TOP10 = {"Tagalog": 28_273_666, "Bisaya/ Binisaya": 15_522_998, "Ilocano": 8_746_169,
               "Cebuano": 8_683_525, "Ilonggo": 8_608_191, "Bikol/Bicol": 7_079_814,
               "Waray": 4_106_539, "Kapampangan": 3_209_738, "Maguindanao": 2_021_099,
               "Pangasinan": 2_012_496, "Not Reported": 18_590}
FINE = ("province", "city", "municipality")
NOT_REPORTED = "Not Reported"
OTHER_LOCAL = "Other Local Ethnicity"
OTHER_LOCAL_SPLIT = {"0600400000": "Aklan", "0307100000": "Zambales"}
# Two more labels whose meaning depends on the island group (PSGC region = first two digits).
# "Dumagat" in Mindanao is the Lumad word for the lowland Visayan settlers, not the Negrito
# Dumagat of Luzon (43% of the census's Dumagat are in Northern Mindanao); "Bukidnon" in the
# Visayas is the highland peoples of Panay and Negros, not the Binukid of Mindanao.
MINDANAO = ("09", "10", "11", "12", "16", "19")
VISAYAS = ("06", "07", "08")
REGION_SPLIT = {"Dumagat": (MINDANAO, "Dumagat (Mindanao)"),
                "Bukidnon": (VISAYAS, "Bukidnon (Visayas)")}

# The lingua francas a group's non-speakers are moved to: the largest present in the unit,
# other than the group's own. Each is a set of ethnicity labels; the first is the label the
# moved people are written under.
LINGUA_FRANCA = [
    ("Tagalog", ["Tagalog", "Caviteño", "Batangan"]),
    ("Cebuano", ["Cebuano", "Bisaya/ Binisaya", "Boholano"]),
    ("Ilonggo", ["Ilonggo"]),
    ("Ilocano", ["Ilocano"]),
    ("Bikol/Bicol", ["Bikol/Bicol"]),
    ("Waray", ["Waray"]),
    ("Tausog/ Tausug", ["Tausog/ Tausug"]),
    ("Maranao", ["Maranao"]),
    ("Maguindanao", ["Maguindanao"]),
    ("Zamboangeño", ["Zamboangeño"]),
]

# Ethnicity set -> language-table set, for the retention check. A set where one table splits
# what the other lumps (the language table has one "Kalinga", the ethnicity table 40 Kalinga
# subgroups). Ethnicity labels are matched exactly; a prefix ending in "*" takes every label
# that starts with it. Groups in no set are kept whole: nothing measures them.
RETENTION_SETS = [
    (["Tagalog", "Caviteño", "Batangan"], ["Tagalog", "Caviteño", "Batangan"]),
    (["Bisaya/ Binisaya", "Cebuano", "Boholano", "Dumagat (Mindanao)"],
     ["Bisaya/Binisaya", "Cebuano", "Boholano"]),
    (["Ilocano", "Bago"], ["Ilocano", "Bago"]),
    (["Ilonggo"], ["Hiligaynon/Ilonggo"]),
    (["Bikol/Bicol"], ["Bikol/Bicol"]),
    (["Waray"], ["Waray"]),
    (["Kapampangan"], ["Kapampangan"]),
    (["Pangasinan"], ["Pangasinan/Panggalato"]),
    (["Maguindanao"], ["Maguindanao"]),
    (["Maranao"], ["Maranao"]),
    (["Tausog/ Tausug"], ["Tausug/Bahasa Sug"]),
    (["Capizeño"], ["Capizeño"]),
    (["Karay-a"], ["Karay-A/Kinaray-A"]),
    (["Masbateño/ Masbatenon"], ["Masbateño/Masbatenon"]),
    (["Surigaonon"], ["Surigaonon"]),
    (["Romblomanon"], ["Romblomanon/Ini"]),
    (["Cuyonen/ Cuyunon"], ["Cuyonon/Cuyonen"]),
    (["Bantoanon"], ["Bantoanon/Asi"]),
    (["Agutaynen"], ["Agutaynen"]),
    (["Zamboangeño"], ["Zamboangueño-Chavacano"]),
    (["Caviteño-Chavacano"], ["Caviteño-Chavacano"]),
    (["Cotabateño-Chavacano", "Cotabateño"], ["Cotabateño-Chavacano", "Cotabateño"]),
    (["Davao-Chavacano"], ["Davao-Chavacano"]),
    (["Davaweño"], ["Davaweño"]),
    (["Subanen/ Subanon", "Kolibugan"], ["Subanen/Subanon/Subanun", "Kalibugan/Kolibugan"]),
    (["Kankanaey", "Kankanaey- Hak'ki", "Applai", "Applai-Kachakran/ Kadaclan"],
     ["Kankanaey", "Kankanaey-Hak’Ki", "Applai", "Applai-Kachakran/Kadaclan"]),
    (["Ibaloy"], ["Ibaloi/Ibaloy"]),
    (["Kalanguya", "Kalanguya-Ikalahan", "Kalanguya-Yattuka"],
     ["Kalanguya", "Kalanguya-Ikalahan", "Kalanguya-Yattuka"]),
    (["Ifugao", "Tuwali", "Tuwali-Kele-i", "Ayangan", "Ayangan-Henanga"],
     ["Ifugao", "Tuwali", "Tuwali-Kele-I", "Ayangan", "Ayangan-Henanga"]),
    (["Kalinga*", "Calinga"], ["Kalinga", "Calinga"]),
    (["Itneg*", "Tingguian"], ["Itneg*", "Tingguian"]),
    (["Bontok", "Bontok-Majukayong", "Baliwon*"], ["Bontok", "Bontok-Majukayong", "Baliwon"]),
    (["Balangao", "Balangao-Lias"], ["Balangao", "Balangao-Lias"]),
    (["Isnag", "Isneg", "Isneg/ Isnag", "Yapayao"], ["Isnag", "Isneg", "Isneg/Isnag", "Yapayao"]),
    (["Ibanag"], ["Ibanag"]),
    (["Itawes"], ["Itawis"]),
    (["Gaddang"], ["Gaddang"]),
    (["Yogad"], ["Yogad"]),
    (["Malaueg"], ["Malaueg"]),
    (["Isinai"], ["Isinai"]),
    (["Ivatan", "Ibatan"], ["Ivatan", "Ibatan"]),
    (["Bugkalot/ Ilongot/ Egongot"], ["Bugkalot/Ilongot"]),
    (["Karao"], ["Karao"]),
    (["Iwak"], ["Iwak/Iowak/Owak/I-Wak"]),
    (["Aeta*", "Ayta", "Abelling/Aberling"], ["Aeta*", "Ayta", "Abelling/Aberling"]),
    (["Agta*", "Alta", "Kabihug*", "Dumagat", "Dumagat-*", "Dumagat/*"],
     ["Agta*", "Alta*", "Arta", "Kabihug*", "Dumagat*"]),
    (["Ati"], ["Ati/Inete/Inati (Negros)"]),
    (["Ata", "Ata/Negrito"], ["Ata", "Atta", "Ata/Inata/Negrito"]),
    (["Mamanwa"], ["Mamanwa"]),
    (["Batak"], ["Batak"]),
    (["Iraya Mangyan", "Alangan Mangyan", "Tadyawan Mangyan", "Buhid Mangyan",
      "Buhid Mangyan [2]", "Bangon Mangyan", "Hanunuo Mangyan", "Tau-buid Mangyan",
      "Ratagnon Mangyan", "Gubatnon Mangyan", "Mangyan"],
     ["Mangyan*"]),
    (["Tagbanua", "Tagbanua-Calamian", "Tagbanua-Kalamianen", "Tagbanua-Tandulanen"],
     ["Tagbanua*", "Kalamyanen"]),
    (["Palawan-o", "Palawan-O- Ken-ey", "Palawan-O- Tao't-Bato", "Palawani"],
     ["Palawan-O*", "Palawani"]),
    (["Molbog"], ["Molbog"]),
    (["Cagayanen"], ["Cagayanen"]),
    (["Manobo", "Manobo-*", "Aromanen-Manobo*", "Obu-Manuvu", "Ubo Monuvu*", "Matigsalog",
      "Tigwahanon", "Dibabawon", "Umayamnon", "Talaingod", "Tinananen", "Lambanguian"],
     ["Manobo*", "Aromanen-Manobo*", "Obu-Manuvu/Ubo-Manobo", "Ubo Monuvu*", "Manubo-Ubo*",
      "Tigwahanon", "Dibabawon", "Umayamnon", "Talaingod", "Tinananen", "Lambangian"]),
    (["Higaonon/ Higa-onon", "Higaonon-Tagoloanon"], ["Higaonon/Higa-Onon", "Higaonon-Tagoloanon"]),
    (["Bukidnon", "Bukidnon-Tagoloanon", "Talaandig"], ["Bukidnon/Binukid", "Talaandig"]),
    (["Bukidnon-Akeanon", OTHER_LOCAL + " (Aklan)"], ["Bukidnon/Binukid-Akeanon/Aklanon"]),
    (["Panay Bukidnon", "Bukidnon (Visayas)", "Bukidnon- Iraynon", "Bukidnon- Ituman", "Bukidnon- Pan-Anayon",
      "Bukidnon-Halowodnon", "Bukidnon-Magahat", "Magahats", "Pan-Ayanon"],
     ["Panay-Bukidnon", "Pan-Ayanon", "Bukidnon/Binukid- Halawodnon", "Magahats"]),
    (["Mandaya", "Mansaka", "Kagan/ Kalagan", "Mangguangan"],
     ["Mandaya", "Mansaka", "Kalagan", "Mangguangan"]),
    (["Tagakaulo"], ["Tagakaulo"]),
    (["Bagobo", "Bagobo Klata", "Bagobo Tagabawa", "Tagabawa", "Guiangan", "Diangan"],
     ["Bagobo*", "Tagabawa", "Guiangan", "Diangan"]),
    (["Blaan"], ["B’Laan/Blaan"]),
    (["T'boli/Tboli"], ["T‘Boli/Tboli"]),
    (["T'duray/ Teduray"], ["T’Duray/Teduray"]),
    (["Iranun/ Iraynun"], ["Iranon/Iranun/Iraynon"]),
    (["Yakan"], ["Yakan"]),
    (["Sama/Samal", "Sama Bangingi", "Sama Badjao", "Badjao", "Bajau", "Sama Dilaut/ Sama Laut"],
     ["Sama/Samal", "Sama Bangingi", "Sama Badjao", "Badjao", "Bajao/Bajau", "Sama Laut"]),
    (["Jama Mapun"], ["Jama Mapun"]),
    (["Sangir/Sangil"], ["Sangir/Sangil"]),
    (["Kamiguin"], ["Kamiguin"]),
    (["Banwaon"], ["Banwaon"]),
    (["Eskaya"], ["Eskaya"]),
    (["Chinese", "Taiwanese"], ["Chinese"]),
    (["Japanese"], ["Japanese"]),
    (["South Korean", "North Korean"], ["Korean"]),
]


def get(req):
    import time
    for attempt in range(5):
        try:
            return urllib.request.urlopen(req, timeout=300).read()
        except urllib.error.HTTPError as e:   # the Wayback Machine answers 429 to bursts
            if e.code != 429 or attempt == 4:
                raise
            time.sleep(30 * (attempt + 1))


def fetch():
    RAW.mkdir(parents=True, exist_ok=True)
    for name, (ts, f) in FILES.items():
        if (RAW / name).exists():
            continue
        url = WB.format(ts=ts, f=f)
        data = get(urllib.request.Request(url, headers=UA))
        want = b"PK" if name.endswith(".xlsx") else b"%PDF"
        if not data.startswith(want):
            raise SystemExit(f"{name}: not a {want!r} file ({len(data):,} bytes)")
        (RAW / name).write_bytes(data)
        print(f"  wrote {name} ({len(data):,} bytes)")
    q = json.dumps({"query": [], "response": {"format": "json-stat"}}).encode()
    req = urllib.request.Request(OPENSTAT, data=q, headers=dict(UA, **{"Content-Type": "application/json"}))
    data = urllib.request.urlopen(req, timeout=120).read()
    json.loads(data)
    HH_JSON.write_bytes(data)
    print(f"  wrote {HH_JSON.name} ({len(data):,} bytes)")


def clean(s):
    s = unicodedata.normalize("NFC", str(s))
    return re.sub(r"\s+", " ", s).strip()


def load_ethnicity():
    wb = openpyxl.load_workbook(RAW / "ph2020_ethnicity.xlsx", read_only=True, data_only=True)
    rows = list(wb["Table"].iter_rows(values_only=True))
    wb.close()
    cats, seen = [], set()
    for h in rows[3][2:]:
        if h is None:
            break
        c = clean(h)
        if c in seen:
            c += " [2]"
        seen.add(c)
        cats.append(c)
    units = []
    for r in rows[5:]:
        if not r or r[0] is None or r[1] is None or not isinstance(r[1], (int, float)):
            continue
        label = re.sub(r"\s+\d$", "", clean(r[0]))
        counts = [int(v or 0) for v in r[2:2 + len(cats)]]
        units.append({"label": label, "hp": int(r[1]), "counts": counts})
    return cats, units


def load_language():
    wb = openpyxl.load_workbook(RAW / "ph2020_language.xlsx", data_only=True)
    out = {}
    for r in wb.active.iter_rows(values_only=True):
        if isinstance(r[1], (int, float)) and r[0] is not None:
            out[clean(r[0])] = int(r[1])
    return out


def load_households():
    js = json.loads(HH_JSON.read_bytes())
    ds = js.get("dataset", js)
    dims = ds["dimension"]
    ids = ds["dimension"]["id"] if "id" in ds["dimension"] else ds["id"]
    size = ds["dimension"]["size"] if "size" in ds["dimension"] else ds["size"]
    geo_dim = [d for d in ids if "Geographic" in d][0]
    var_dim = [d for d in ids if d != geo_dim][0]
    geo = dims[geo_dim]["category"]
    var = dims[var_dim]["category"]
    gi = sorted(geo["index"].items(), key=lambda kv: kv[1])
    vi = sorted(var["index"].items(), key=lambda kv: kv[1])
    vlab = [var["label"][k] for k, _ in vi]
    vals = ds["value"]
    if isinstance(vals, dict):
        vals = [vals.get(str(i)) for i in range(size[0] * size[1])]
    order = ids.index(geo_dim)
    rows = []
    for g, gpos in gi:
        rec = {"label": geo["label"][g]}
        for v, vpos in vi:
            i = gpos * size[1] + vpos if order == 0 else vpos * size[1] + gpos
            rec[var["label"][v]] = vals[i]
        rows.append(rec)
    return rows, vlab


def expand(patterns, universe):
    out = []
    for p in patterns:
        if p.endswith("*"):
            hit = [u for u in universe if u.startswith(p[:-1])]
        else:
            hit = [p] if p in universe else []
        if not hit:
            raise SystemExit(f"retention set: {p!r} matches nothing")
        out += [h for h in hit if h not in out]
    return out


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--fetch", action="store_true")
    args = ap.parse_args()
    if args.fetch or not (RAW / "ph2020_ethnicity.xlsx").exists():
        fetch()

    ok = True

    def report(good, msg):
        nonlocal ok
        ok &= bool(good)
        print(f"  {'OK ' if good else 'BAD'} {msg}")

    print("Philippines, 2020 CPH, household population by ethnicity (PSA)\n")
    cats, units = load_ethnicity()
    bad = [u["label"] for u in units if sum(u["counts"]) != u["hp"]]
    report(len(cats) == N_ETH and not bad,
           f"{len(cats)} ethnicity columns; {len(units)} rows; rows not summing to household "
           f"population: {bad[:5]}")

    rd = pd.read_csv(RD_NORM, dtype={"geo_id": str}, keep_default_na=False, na_values=[""])
    rd = rd[rd["source_category"] == "Household Population"].reset_index(drop=True)
    same = (len(rd) == len(units) and all(
        clean(a) == b["label"] and int(h) == b["hp"]
        for a, h, b in zip(rd["geo_name"], rd["count"], units)))
    report(same, f"{len(units)} rows join religiondots' ph.csv row for row on label and "
                 f"household population ({len(rd)} there)")
    if not same:
        for a, h, b in zip(rd["geo_name"], rd["count"], units):
            if clean(a) != b["label"] or int(h) != b["hp"]:
                print(f"      {a!r} {h} | {b['label']!r} {b['hp']}")
                break
        raise SystemExit("join failed; nothing written")
    for u, (_, r) in zip(units, rd.iterrows()):
        u["geo_id"], u["level"] = r["geo_id"], r["geo_level"]

    nat = units[0]
    fine = [u for u in units if u["level"] in FINE]
    col_off = [sum(u["counts"][j] for u in fine) - nat["counts"][j] for j in range(len(cats))]
    report(len(fine) == 117 and sum(u["hp"] for u in fine) == HP_NATIONAL == nat["hp"]
           and not any(col_off),
           f"{len(fine)} fine units sum to {sum(u['hp'] for u in fine):,} and to the national "
           f"row in every column")
    idx = {c: j for j, c in enumerate(cats)}
    pr = {k: (nat["counts"][idx[k]], v) for k, v in PRESS_TOP10.items()}
    report(all(a == b for a, b in pr.values()),
           "press release Table 1 top ten and Not Reported match the national row"
           + "".join(f"; {k} {a:,} vs {b:,}" for k, (a, b) in pr.items() if a != b))

    hh_rows, vlab = load_households()
    hh_col = [v for v in vlab if "Number of Households" in v or "Households" == v.strip()]
    hp_col = [v for v in vlab if "Household Population" in v]
    # OpenSTAT lists HUCs in its own order, so the join is on household population, which is
    # unique across the 135 rows, with the names checked as a second key
    by_hp = {int(r[hp_col[0]]): r for r in hh_rows}

    def fold(s):
        s = re.sub(r"\(.*?\)|\*|\.|city of|municipality of", " ", clean(s).lower())
        return re.sub(r"[^a-z]", "", s)

    name_bad = [(u["label"], by_hp[u["hp"]]["label"]) for u in units
                if u["level"] in FINE and u["hp"] in by_hp
                and fold(u["label"]) != fold(by_hp[u["hp"]]["label"])
                # OpenSTAT's name for the BARMM interim province (its eight SGU clusters)
                and (u["label"], fold(by_hp[u["hp"]]["label"])) != ("Interim Province",
                                                                     "eightareaclusters")]
    names_ok = not name_bad
    if name_bad:
        print("      name mismatches:", name_bad[:6])
    report(len(hh_rows) == len(units) == len(by_hp) and hh_col and hp_col
           and all(u["hp"] in by_hp for u in units) and names_ok,
           f"OpenSTAT household table: {len(hh_rows)} rows, one per unit on household "
           f"population, names agree")
    for u in units:
        u["hh"] = int(by_hp[u["hp"]][hh_col[0]])
    report(nat["hh"] >= HH_NATIONAL, f"households {nat['hh']:,} (the language table's "
                                     f"{HH_NATIONAL:,} excludes the homeless)")

    lang = load_language()
    total = lang.pop("Total")
    report(total == HH_NATIONAL and sum(lang.values()) == total,
           f"language table: {len(lang)} categories sum to {sum(lang.values()):,} households "
           f"(total {total:,})")
    if not ok:
        raise SystemExit("\nreconciliation FAILED; nothing written")

    # ---- the long table, with the place-dependent label split ----
    long = []
    for u in fine:
        for c, v in zip(cats, u["counts"]):
            if v <= 0 or c == NOT_REPORTED:
                continue
            if c == OTHER_LOCAL and u["geo_id"] in OTHER_LOCAL_SPLIT:
                c = f"{OTHER_LOCAL} ({OTHER_LOCAL_SPLIT[u['geo_id']]})"
            if c in REGION_SPLIT and u["geo_id"][:2] in REGION_SPLIT[c][0]:
                c = REGION_SPLIT[c][1]
            long.append({"geo_id": u["geo_id"], "geo_level": u["level"], "geo_name": u["label"],
                         "ethnicity": c, "count": float(v)})
    df = pd.DataFrame(long)
    units_df = pd.DataFrame([{"geo_id": u["geo_id"], "hp": u["hp"], "hh": u["hh"]} for u in fine])
    k = dict(zip(units_df["geo_id"], units_df["hh"] / units_df["hp"]))
    df["k"] = df["geo_id"].map(k)
    eth_universe = sorted(df["ethnicity"].unique())
    lang_universe = sorted(lang)

    # unit -> largest ethnicity, and lingua franca totals
    largest = df.loc[df.groupby("geo_id")["count"].idxmax()].set_index("geo_id")["ethnicity"]
    lf_tot = {}
    for name, members in LINGUA_FRANCA:
        lf_tot[name] = df[df["ethnicity"].isin(members)].groupby("geo_id")["count"].sum()
    lf_df = pd.DataFrame(lf_tot).fillna(0.0)
    lf_of = {name: set(m) for name, m in LINGUA_FRANCA}

    # ---- retention ----
    print("\n  RETENTION (national household-language table vs ethnicity)")
    ret_rows, used = [], set()
    moves = []
    for eth_p, lang_p in RETENTION_SETS:
        e_set = expand(eth_p, eth_universe)
        l_set = expand(lang_p, lang_universe)
        clash = used & set(e_set)
        if clash:
            raise SystemExit(f"ethnicity in two retention sets: {sorted(clash)}")
        used |= set(e_set)
        sub = df[df["ethnicity"].isin(e_set)]
        persons = sub["count"].sum()
        expected = (sub["count"] * sub["k"]).sum()
        hh = sum(lang[x] for x in l_set)
        r = min(1.0, hh / expected) if expected else 1.0
        deficit = (1 - r) * persons
        own_lf = [n for n, m in LINGUA_FRANCA if set(e_set) & set(m)]
        ret_rows.append({"ethnicity": " + ".join(e_set), "language_table": " + ".join(l_set),
                         "persons": int(persons), "households_expected": round(expected),
                         "households_language": hh, "ratio": round(hh / expected, 3) if expected else None,
                         "retention": round(r, 3), "moved": round(deficit)})
        if deficit < 1:
            continue
        # where: units where this set is not the largest ethnicity first
        per_unit = sub.groupby("geo_id")["count"].sum()
        home = per_unit.index[per_unit.index.map(lambda g: largest.get(g) in e_set)]
        away = per_unit.drop(home)
        take = pd.Series(0.0, index=per_unit.index)
        a = away.sum()
        if a >= deficit:
            take[away.index] = away * (deficit / a)
        else:
            take[away.index] = away
            rest = deficit - a
            take[home] = per_unit[home] * (rest / per_unit[home].sum())
        for g, t in take.items():
            if t <= 0:
                continue
            cand = lf_df.loc[g].drop(labels=own_lf, errors="ignore")
            target = cand.idxmax()
            rows_g = sub[sub["geo_id"] == g]
            for _, row in rows_g.iterrows():
                m = t * row["count"] / per_unit[g]
                moves.append((g, row["ethnicity"], target, m))
    ret = pd.DataFrame(ret_rows).sort_values("persons", ascending=False)
    for r in ret.itertuples():
        print(f"    {r.ethnicity[:44]:44} {r.persons:>11,} hh expected {r.households_expected:>10,}"
              f"  language {r.households_language:>10,}  ratio {r.ratio:5.2f}  moved {r.moved:>9,}")
    kept_whole = sorted(set(eth_universe) - used)
    print(f"\n  {len(kept_whole)} ethnicities with no language-table counterpart, kept whole: "
          f"{int(df[df['ethnicity'].isin(kept_whole)]['count'].sum()):,} people")

    # apply
    df["origin"] = df["ethnicity"]
    df["tier"] = "derived"
    mv = pd.DataFrame(moves, columns=["geo_id", "ethnicity", "target", "count"])
    sub_m = mv.groupby(["geo_id", "ethnicity"])["count"].sum()
    df = df.set_index(["geo_id", "ethnicity"])
    df.loc[sub_m.index, "count"] -= sub_m
    df = df.reset_index()
    names = df.drop_duplicates("geo_id").set_index("geo_id")[["geo_level", "geo_name"]]
    add = mv.rename(columns={"ethnicity": "origin", "target": "ethnicity"})
    add["tier"] = "modelled"
    add = add.join(names, on="geo_id")
    df = pd.concat([df.drop(columns=["k"]), add], ignore_index=True)
    df = df.groupby(["geo_id", "geo_level", "geo_name", "ethnicity", "origin", "tier"],
                    as_index=False)["count"].sum()
    if (df["count"] < -1e-6).any():
        raise SystemExit("negative count after the shift")
    # integer counts per unit by largest remainder, so units still sum exactly
    out = []
    for g, part in df.groupby("geo_id"):
        target = int(round(part["count"].sum()))
        fl = part["count"].apply(int)
        short = target - int(fl.sum())
        order = (part["count"] - fl).sort_values(ascending=False).index[:short]
        fl[order] += 1
        p = part.copy()
        p["count"] = fl
        out.append(p)
    df = pd.concat(out)
    df = df[df["count"] > 0]

    want = {u["geo_id"]: u["hp"] - u["counts"][idx[NOT_REPORTED]] for u in fine}
    got = df.groupby("geo_id")["count"].sum()
    off = {g: int(got.get(g, 0) - w) for g, w in want.items() if got.get(g, 0) != w}
    report(not off, f"after the shift every unit sums to its household population less Not "
                    f"Reported ({sum(want.values()):,} people); off: {list(off.items())[:5]}")
    moved = int(df.loc[df["tier"] == "modelled", "count"].sum())
    print(f"  -- {moved:,} people ({moved / sum(want.values()):.1%}) moved to a lingua franca; "
          f"by target: " + ", ".join(f"{k} {v:,}" for k, v in
                                    df[df['tier'] == 'modelled'].groupby('ethnicity')['count'].sum()
                                    .sort_values(ascending=False).astype(int).items()))
    if not ok:
        raise SystemExit("\nFAILED; nothing written")

    df = df.rename(columns={"ethnicity": "source_category"})
    df = df[["geo_id", "geo_level", "geo_name", "source_category", "origin", "tier", "count"]]
    df = df.sort_values(["geo_id", "count"], ascending=[True, False])
    OUT.parent.mkdir(parents=True, exist_ok=True)
    df.to_csv(OUT, index=False, encoding="utf-8")
    ret.to_csv(OUT_RET, index=False, encoding="utf-8")
    print(f"\nwrote {OUT}: {len(df):,} rows, {df['geo_id'].nunique()} units, "
          f"{df['count'].sum():,} people")
    print(f"wrote {OUT_RET}")


if __name__ == "__main__":
    sys.stdout.reconfigure(encoding="utf-8")
    main()
