"""Croatia: DZS, Popis stanovnistva 2021, mother tongue by town/municipality.

    python sources/hr_census.py --fetch    download the workbook (18 MB) if missing
    python sources/hr_census.py            normalise from data/raw/hr/

-> data/normalized/hr.csv (levels `country`, `county`, `municipality`, `city_district`;
   alternatives, never summed; `municipality` and `city_district` together are the drawn cover)

One workbook, the census's "population by towns/municipalities" release
(https://podaci.dzs.hr/media/td3jvrbu/popis_2021-stanovnistvo_po_gradovima_opcinama.xlsx),
sheet `4.`: STANOVNISTVO PREMA MATERINSKOM JEZIKU PO GRADOVIMA/OPCINAMA, POPIS 2021. Croatian,
24 named languages, "Other languages" and "Unknown", each as a count and a percentage.

THE LAYOUT IS RELIGIONDOTS' RELIGION SHEET (religiondots/sources/hr.py reads sheet `2.` of the
same file), so this follows it exactly and produces the SAME geo_ids: the workbook has no codes,
so a unit is "ZUPANIJA|NAME", and religiondots' data/geo/hr/hr_lookup.csv (read only) routes
each one to its LAU code. Grad Zagreb appears only as its 17 gradske cetvrti (city districts),
so the cover is 555 municipalities plus 17 districts = 572 rows.

CHECKS: national total 3,871,833; the counties and the 572-unit cover each sum to the national
row in every column; the 26 categories partition every unit (DZS neither rounds nor
suppresses); every unit's total equals religiondots' total for the same geo_id from the
religion sheet of the same census, both ways (a join check: same units, same keys); and the
sheet's percentage columns agree with count / total to DZS's two decimals.
"""

import csv
import os
import sys

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
RAW = os.path.join(ROOT, "data", "raw", "hr")
OUT = os.path.join(ROOT, "data", "normalized", "hr.csv")
RD_NORM = os.path.join(os.path.dirname(ROOT), "religiondots", "data", "normalized", "hr.csv")

SOURCE_ID = "hr_popis_2021_mt"
YEAR = 2021

XLSX_NAME = "gradovi_opcine.xlsx"
XLSX_URL = ("https://podaci.dzs.hr/media/td3jvrbu/"
            "popis_2021-stanovnistvo_po_gradovima_opcinama.xlsx")
MIN_BYTES = 10_000_000

COLUMNS = ["geo_id", "geo_level", "geo_name", "source_category", "count", "tier", "year",
           "source_id", "note"]

NATIONAL = 3_871_833          # DZS, total population, census 2021
SHEET = "4."
HEADER_ROW = 8
FIRST_DATA_ROW = 9
COUNTY_COL, KIND_COL, NAME_COL, TOTAL_COL = 0, 1, 4, 5
NATIONAL_NAME = "Republika Hrvatska"
TOTAL_LABEL = "Total"

# Count column -> English half of DZS's bilingual header, which is the label written to hr.csv.
# The column after each is the same figure as a percentage. Asserted against the sheet.
CATEGORY_COLS = {
    7: "Croatian", 9: "Croato-Serbian", 11: "Albanian", 13: "Bosnian", 15: "Bulgarian",
    17: "Montenegrin", 19: "Czech", 21: "Hungarian", 23: "Macedonian", 25: "German",
    27: "Polish", 29: "Romani", 31: "Romanian", 33: "Russian", 35: "Ruthenian", 37: "Slovak",
    39: "Slovenian", 41: "Serbian", 43: "Serbo-Croatian", 45: "Italian", 47: "Turkish",
    49: "Ukrainian", 51: "Vlach", 53: "Hebrew", 55: "Other languages", 57: "Unknown",
}

KIND_LEVEL = {"Grad": "municipality", "Općina": "municipality",
              "Gradska četvrt": "city_district"}
TRUE_ZERO = ("-", "–", "—")


def fetch():
    import requests
    import urllib3
    urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)

    os.makedirs(RAW, exist_ok=True)
    dest = os.path.join(RAW, XLSX_NAME)
    if os.path.exists(dest) and os.path.getsize(dest) >= MIN_BYTES:
        print("already have", dest)
        return
    print("GET", XLSX_URL)
    # verify=False as religiondots does: podaci.dzs.hr's chain does not validate on this box
    r = requests.get(XLSX_URL, headers={"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"},
                     timeout=900, verify=False)
    r.raise_for_status()
    if not r.content.startswith(b"PK") or len(r.content) < MIN_BYTES:
        raise SystemExit(f"not the workbook: {len(r.content):,} bytes, starts {r.content[:16]!r}")
    tmp = dest + ".part"
    with open(tmp, "wb") as fh:
        fh.write(r.content)
    os.replace(tmp, dest)
    print(f"  {os.path.getsize(dest):,} bytes")


def _txt(v):
    return "" if v is None else " ".join(str(v).split())


def _num(cell, where):
    if cell is None:
        return None
    if isinstance(cell, (int, float)):
        if cell != int(cell):
            raise SystemExit(f"{where}: {cell!r} is not a whole number")
        return int(cell)
    s = str(cell).strip()
    if s in TRUE_ZERO:
        return 0
    if s == "":
        return None
    raise SystemExit(f"unexpected value {cell!r} at {where}; DZS changed the sheet")


def _pct(cell):
    if isinstance(cell, (int, float)):
        return float(cell)
    return 0.0 if _txt(cell) in TRUE_ZERO else None


def read():
    import openpyxl

    path = os.path.join(RAW, XLSX_NAME)
    if not os.path.exists(path) or os.path.getsize(path) < MIN_BYTES:
        raise SystemExit(f"missing or truncated {path}; run with --fetch first")
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[SHEET]
    title = _txt(list(ws.iter_rows(min_row=2, max_row=2, values_only=True))[0][0])
    if "MATERINSKOM JEZIKU" not in title:
        raise SystemExit(f"sheet {SHEET} is {title!r}, not the mother-tongue table")
    header = list(ws.iter_rows(min_row=HEADER_ROW, max_row=HEADER_ROW, values_only=True))[0]
    for c, label in CATEGORY_COLS.items():
        en = str(header[c]).split("\n")[-1].strip()
        pct_en = str(header[c + 1]).split("\n")[-1].strip()
        if en != label or pct_en.rstrip() != f"{label}, %":
            raise SystemExit(f"column {c}: header {header[c]!r} / {header[c + 1]!r}, expected "
                             f"{label!r}; the column map is stale")

    records, current_county, pct_bad = [], None, []
    for ri, r in enumerate(ws.iter_rows(min_row=FIRST_DATA_ROW, values_only=True), FIRST_DATA_ROW):
        if len(r) <= TOTAL_COL:
            continue
        county, kind, name = _txt(r[COUNTY_COL]), _txt(r[KIND_COL]), _txt(r[NAME_COL])
        total = _num(r[TOTAL_COL], f"row {ri}")
        if total is None:
            continue
        if county:
            current_county = county
        if county == NATIONAL_NAME:
            gid, level, gname = "HR", "country", NATIONAL_NAME
        elif not kind:
            gid, level, gname = current_county, "county", current_county
        else:
            level = KIND_LEVEL.get(kind)
            if level is None:
                raise SystemExit(f"row {ri}: unknown unit kind {kind!r}")
            if not name:
                raise SystemExit(f"row {ri}: {kind} with no name")
            gid, gname = f"{current_county}|{name}", name
        values = {TOTAL_LABEL: total}
        for c, label in CATEGORY_COLS.items():
            n = _num(r[c], f"row {ri} col {c}")
            values[label] = n or 0
            p = _pct(r[c + 1])
            if p is not None and total and abs(100.0 * values[label] / total - p) > 0.006:
                pct_bad.append((gid, label, values[label], total, p))
        records.append(dict(gid=gid, level=level, name=gname, county=current_county,
                            values=values))

    rows = []
    for rec in records:
        for label, n in rec["values"].items():
            note = f"DZS Popis 2021 sheet {SHEET}; level={rec['level']}"
            if rec["level"] in ("municipality", "city_district"):
                note += f"; county={rec['county']}"
            if label == TOTAL_LABEL:
                note += "; universe total, not a language"
            rows.append(dict(geo_id=rec["gid"], geo_level=rec["level"], geo_name=rec["name"],
                             source_category=label, count=n, tier="measured", year=YEAR,
                             source_id=SOURCE_ID, note=note))
    return rows, records, pct_bad


def check(records, pct_bad):
    ok = True
    lv = {}
    for rec in records:
        lv.setdefault(rec["level"], []).append(rec)
    expected = {"country": 1, "county": 21, "municipality": 555, "city_district": 17}
    for k, want in expected.items():
        got = len(lv.get(k, ()))
        good = got == want
        ok &= good
        print(f"  {'OK ' if good else 'BAD'} {k:<14} {got:>4} rows (expected {want})")

    nat = lv["country"][0]["values"]
    good = nat[TOTAL_LABEL] == NATIONAL
    ok &= good
    print(f"\n  {'OK ' if good else 'BAD'} national total {nat[TOTAL_LABEL]:,} (published {NATIONAL:,})")

    cover = lv["municipality"] + lv["city_district"]
    for name, group in (("21 county", lv["county"]), ("572-unit cover", cover)):
        bad = [(k, sum(r["values"][k] for r in group), v) for k, v in nat.items()
               if sum(r["values"][k] for r in group) != v]
        ok &= not bad
        print(f"  {'OK ' if not bad else 'BAD'} the {name} rows sum to the national row in all "
              f"{len(nat)} columns")
        for k, s, v in bad[:6]:
            print(f"        {k}: {s:,} vs {v:,}")

    bad = [r["gid"] for r in records
           if sum(v for k, v in r["values"].items() if k != TOTAL_LABEL) != r["values"][TOTAL_LABEL]]
    ok &= not bad
    print(f"  {'OK ' if not bad else 'BAD'} the {len(CATEGORY_COLS)} categories partition every "
          f"one of {len(records)} rows")
    for g in bad[:6]:
        print(f"        {g}")

    ids = [r["gid"] for r in cover]
    good = len(set(ids)) == len(ids)
    ok &= good
    print(f"  {'OK ' if good else 'BAD'} {len(set(ids))} distinct unit ids in the cover")

    good = not pct_bad
    ok &= good
    print(f"  {'OK ' if good else 'BAD'} every percentage column agrees with count / total "
          f"({len(pct_bad)} disagree)")
    for x in pct_bad[:6]:
        print(f"        {x}")

    if os.path.exists(RD_NORM):
        rd = {}
        with open(RD_NORM, encoding="utf-8") as fh:
            for row in csv.DictReader(fh):
                if row["geo_level"] in ("municipality", "city_district") \
                        and row["source_category"] == "Ukupno":
                    rd[row["geo_id"]] = int(row["count"])
        mine = {r["gid"]: r["values"][TOTAL_LABEL] for r in cover}
        miss = sorted(set(mine) ^ set(rd))
        diff = [(k, mine[k], rd[k]) for k in mine if k in rd and mine[k] != rd[k]]
        good = not miss and not diff
        ok &= good
        print(f"  {'OK ' if good else 'BAD'} all {len(mine)} unit ids match religiondots' hr.csv "
              f"(religion sheet) both ways, with identical totals in every unit")
        for k in miss[:6]:
            print(f"        only on one side: {k}")
        for k, a, b in diff[:6]:
            print(f"        {k}: {a:,} here, {b:,} in the religion sheet")
    else:
        ok = False
        print(f"  BAD {RD_NORM} missing, so the join is unchecked")

    print("\n  categories, national:")
    for k, v in sorted(nat.items(), key=lambda kv: -kv[1]):
        if k != TOTAL_LABEL:
            print(f"    {v:>10,}  {100.0 * v / NATIONAL:6.2f}%  {k}")
    if not ok:
        raise SystemExit("reconciliation FAILED")


def main():
    if "--fetch" in sys.argv:
        fetch()
    rows, records, pct_bad = read()
    check(records, pct_bad)
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=COLUMNS)
        w.writeheader()
        w.writerows(rows)
    print("\nwrote", OUT, f"({len(rows):,} rows)")


if __name__ == "__main__":
    main()
