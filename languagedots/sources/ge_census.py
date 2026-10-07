"""Georgia: Geostat, 2024 Population and Agricultural Census, native language by self-governed unit.

    python sources/ge_census.py --fetch    download the two 2024 workbooks and the 2014 PxWeb cube
    python sources/ge_census.py            normalise from data/raw/ge/

-> data/normalized/ge.csv (levels `country`, `region`, `unit`; alternatives, never summed; only
   `unit` is drawn: Tbilisi and the 63 other self-governed cities and municipalities)

THE DRAWN TABLE. Geostat's results page for the 2024 census ("Demographic and Social
Characteristics", geostat.ge/en/modules/categories/910) links a workbook per topic, published
22 June 2026:

  5. Population by regions, self-governed units, native language and Georgian language knowledge
     level.xlsx (media/80624)
     One sheet. Rows: Georgia, then each region followed by its self-governed units (Tbilisi is a
     region and a unit at once). Columns: the native-language block (Total; Georgian, Abkhaz,
     Ossetian, Azerbaijanian, Russian, Armenian, Other, Not stated), then four blocks splitting
     the people whose native language is NOT Georgian by Georgian knowledge (fluent, partial,
     none, not stated), each with the same languages less Georgian.

  4. Population by regions, self-governed units, sex and nationality.xlsx (media/80623)
     The same rows, by ethnicity. The second-table check: every row's total must be the same
     person count, and per unit the Armenian and Azerbaijani speakers must track the Armenians
     and Azerbaijanis (printed, with a loose band asserted).

The 2024 census portal (census2024.geostat.ge) is a React app whose results page only links back
to these category pages; the PxWeb at pc-axis.geostat.ge holds the 2014 census and nothing for
2024 yet (2026-10-04). 2014's native-language table is there at region level only
(`20_Population_by_region,_by_native_languages_and_fluently_speak_Georgian....px`), with the same
eight answers; it is fetched as a witness, not drawn.

The site's TLS chain does not verify from here (curl exit 60), so requests run with verify=False;
the files are public statistics, and each is checked for its own shape after download.

CHECKS: the national row against the Main Results PDF's printed figures; per row, the eight
answers sum to Total; per region, its units sum to it per answer; regions sum to Georgia per
answer; per row and non-Georgian language, the four Georgian-knowledge blocks sum to the native-
language column; 64 units; table 4's totals equal table 5's on every row; and the 2014 regional
shares beside the 2024 ones.
"""

import csv
import json
import os
import sys
import urllib.parse

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
RAW = os.path.join(ROOT, "data", "raw", "ge")
OUT = os.path.join(ROOT, "data", "normalized", "ge.csv")

SOURCE_ID = "ge_census_2024_nl"
YEAR = 2024
UA = {"User-Agent": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"}

FILES = {
    "t5_native_language_2024.xlsx": (
        "https://www.geostat.ge/media/80624/5.-Population-by-regions%2C-self-governed-units%2C-"
        "native-language-and-Georgian-language-knowledge-level.xlsx", 20_000),
    "t4_nationality_2024.xlsx": (
        "https://www.geostat.ge/media/80623/4.-Population-by-regions%2C-self-governed-units%2C-"
        "sex-and-nationality.xlsx", 20_000),
}
PX = ("https://pc-axis.geostat.ge/PXWeb/api/v1/en/" + urllib.parse.quote(
    "Database/Population Census 2014/Demographic And Social Characteristics/"
    "20_Population_by_region,_by_native_languages_and_fluently_speak_Georgian....px"))
PX_FILE = os.path.join(RAW, "px2014_native_language_region.json")

LANGS = ["Georgian", "Abkhaz", "Ossetian", "Azerbaijanian", "Russian", "Armenian", "Other",
         "Not stated"]
NON_GEO = ["Abkhaz", "Ossetian", "Azerbaijanian", "Russian", "Armenian", "Other"]
# Main Results of the 2024 Population and Agricultural Census (Geostat, 22.06.2026), p.3 and p.15
PRINTED = {"Total": 3_929_581, "Georgian": 3_343_987, "Azerbaijanian": 265_534,
           "Armenian": 139_438, "Russian": 54_460, "Ossetian": 3_839, "Abkhaz": 370}
EXPECTED_UNITS = 64
EXPECTED_REGIONS = 11      # Tbilisi counted as a region too

# Region rows as the table names them. Tbilisi is region and unit both.
REGIONS = ["C. Tbilisi", "Adjara A.R.", "Guria", "Imereti", "Kakheti", "Mtskheta-Mtianeti",
           "Racha-Lechkhumi and Kvemo Svaneti", "Samegrelo-Zemo Svaneti", "Samtskhe-Javakheti",
           "Kvemo Kartli", "Shida Kartli"]


def fetch():
    import requests
    import urllib3
    urllib3.disable_warnings()
    os.makedirs(RAW, exist_ok=True)
    for name, (url, floor) in FILES.items():
        path = os.path.join(RAW, name)
        if os.path.exists(path) and os.path.getsize(path) > floor:
            print("already have", path)
            continue
        print("GET", url)
        r = requests.get(url, headers=UA, timeout=300, verify=False)
        r.raise_for_status()
        if r.content[:2] != b"PK":
            raise SystemExit(f"{name}: not an xlsx, first bytes {r.content[:60]!r}")
        with open(path, "wb") as fh:
            fh.write(r.content)
        print(f"  {len(r.content):,} bytes")
    if os.path.exists(PX_FILE) and os.path.getsize(PX_FILE) > 1_000:
        print("already have", PX_FILE)
        return
    # PxWeb v1 here: explicit selection of every value, json-stat v1 (religiondots/sources/ge.py)
    print("GET ", PX)
    meta = requests.get(PX, headers=UA, timeout=120, verify=False)
    meta.raise_for_status()
    variables = meta.json()["variables"]
    query = {"query": [{"code": v["code"], "selection": {"filter": "item", "values": v["values"]}}
                       for v in variables], "response": {"format": "json-stat"}}
    print("POST", PX)
    r = requests.post(PX, headers=UA, json=query, timeout=300, verify=False)
    r.raise_for_status()
    doc = r.json()
    if "dataset" not in doc:
        raise SystemExit(f"2014 cube: not json-stat v1, keys {list(doc)}")
    with open(PX_FILE, "w", encoding="utf-8") as fh:
        json.dump(doc, fh, ensure_ascii=False)
    print(f"  {os.path.getsize(PX_FILE):,} bytes")


def _sheet(name):
    import openpyxl
    path = os.path.join(RAW, name)
    if not os.path.exists(path):
        raise SystemExit(f"missing {path} -- run with --fetch first")
    return openpyxl.load_workbook(path, data_only=True, read_only=True).worksheets[0]


def _rows(ws, first=8):
    """(name, values) for every data row, stopping at the blank row before the source note."""
    out = []
    for r in ws.iter_rows(min_row=first, values_only=True):
        if r[0] is None:
            break
        out.append((str(r[0]).strip(), list(r[1:])))
    return out


def read_t5():
    ws = _sheet("t5_native_language_2024.xlsx")
    head = [list(r) for r in ws.iter_rows(min_row=5, max_row=7, values_only=True)]
    # the layout this parser assumes, asserted cell by cell
    want_langs = ["Georgian", "Abkhaz", "Ossetian", "Azerbaijanian", "Russian", "Armenian",
                  "Other", "Not stated"]
    if head[1][1] != "Total" or head[2][2:10] != want_langs:
        raise SystemExit(f"t5 header has changed: {head[1][:3]} / {head[2][:10]}")
    blocks = {"fluent": 10, "partial": 17, "none": 24, "not stated": 31}
    for b, c in blocks.items():
        want = NON_GEO + (["Not stated"] if b == "not stated" else [])
        if head[1][c] != "Total" or head[2][c + 1:c + 1 + len(want)] != want:
            raise SystemExit(f"t5 block {b!r} at column {c} has changed: {head[2][c:c + 8]}")
    rows = []
    for name, v in _rows(ws):
        rec = {"name": name, "Total": int(v[0])}
        rec.update({lang: int(v[1 + i]) for i, lang in enumerate(LANGS)})
        rec["blocks"] = {}
        for b, c in blocks.items():
            n = len(NON_GEO) + (1 if b == "not stated" else 0)
            rec["blocks"][b] = {"Total": int(v[c - 1]),
                                **{lang: int(v[c + i]) for i, lang in enumerate(
                                    (NON_GEO + ["Not stated"])[:n])}}
        rows.append(rec)
    return rows


def levels(rows):
    """Tag each row country / region / unit and give each unit its region."""
    if rows[0]["name"] != "Georgia":
        raise SystemExit(f"first row is {rows[0]['name']!r}, not Georgia")
    out, region = [], None
    for r in rows[1:]:
        if r["name"] in REGIONS:
            region = r["name"]
            out.append(dict(r, level="region", region=region))
            if region == "C. Tbilisi":
                out.append(dict(r, level="unit", region=region))   # a region and a unit at once
        else:
            if region is None:
                raise SystemExit(f"unit {r['name']!r} before any region")
            out.append(dict(r, level="unit", region=region))
    return [dict(rows[0], level="country", region="")] + out


def check(rows, t4, px):
    ok = True

    def say(good, msg):
        nonlocal ok
        ok &= bool(good)
        print(f"  {'OK ' if good else 'BAD'} {msg}")

    nat = rows[0]
    units = [r for r in rows if r["level"] == "unit"]
    regions = [r for r in rows if r["level"] == "region"]
    say(len(units) == EXPECTED_UNITS, f"{len(units)} self-governed units (expected {EXPECTED_UNITS})")
    say(len(regions) == EXPECTED_REGIONS, f"{len(regions)} regions (expected {EXPECTED_REGIONS})")
    say(len({r['name'] for r in units}) == len(units), "unit names are unique")

    for k, v in PRINTED.items():
        say(nat[k] == v, f"national {k} {nat[k]:,} against the Main Results PDF's {v:,}")

    bad = [r["name"] for r in rows if sum(r[l] for l in LANGS) != r["Total"]]
    say(not bad, f"the 8 answers sum to Total on every row ({len(rows)} rows) {bad[:4]}")

    for reg in regions:
        if reg["name"] == "C. Tbilisi":
            continue
        kids = [u for u in units if u["region"] == reg["name"]]
        diff = {l: reg[l] - sum(u[l] for u in kids) for l in ["Total"] + LANGS}
        if any(diff.values()):
            say(False, f"{reg['name']}: units do not sum to the region {diff}")
    say(True, "every region's units sum to it, per answer (a BAD line above if not)")
    diff = {l: nat[l] - sum(r[l] for r in regions) for l in ["Total"] + LANGS}
    say(not any(diff.values()), f"the 11 regions sum to Georgia per answer {diff}")
    diff = {l: nat[l] - sum(u[l] for u in units) for l in ["Total"] + LANGS}
    say(not any(diff.values()), f"the 64 units sum to Georgia per answer {diff}")

    # the Georgian-knowledge blocks partition each non-Georgian native language
    nbad = 0
    for r in rows:
        for lang in NON_GEO + ["Not stated"]:
            s = sum(b.get(lang, 0) for b in r["blocks"].values())
            if s != r[lang]:
                nbad += 1
        if sum(b["Total"] for b in r["blocks"].values()) != r["Total"] - r["Georgian"]:
            nbad += 1
    say(nbad == 0, f"the four Georgian-knowledge blocks sum to each non-Georgian native "
                   f"language on every row ({nbad} mismatches)")

    # ---- second table: nationality ----
    names5 = [r["name"] for r in rows if not (r["level"] == "unit" and r["region"] == r["name"])]
    names4 = [n for n, _ in t4]
    say(names4 == names5, "table 4 (nationality) has the same rows in the same order")
    tot4 = {n: int(v[0]) for n, v in t4}
    bad = [r["name"] for r in rows if tot4.get(r["name"]) != r["Total"]]
    say(not bad, f"table 4's total equals table 5's on every row {bad[:4]}")
    # columns: Total, Georgians, Abkhazians, Ossetians, Azerbaijanis, Russians, Armenians, ...
    eth = {n: {"Armenian": int(v[6]), "Azerbaijanian": int(v[4]), "Georgian": int(v[1])}
           for n, v in t4}
    print("\n  speakers / members of the group, per unit with 2,000+ members "
          "(the column join's witness):")
    for lang in ("Armenian", "Azerbaijanian"):
        rs = sorted((r[lang] / eth[r["name"]][lang], r["name"]) for r in units
                    if eth[r["name"]][lang] >= 2000)
        print(f"    {lang:<14} {len(rs)} units, {rs[0][0]:.2f} ({rs[0][1]}) to "
              f"{rs[-1][0]:.2f} ({rs[-1][1]}), median {rs[len(rs) // 2][0]:.2f}")
        # the city Armenians of Batumi (0.52) and Tbilisi often name Georgian or Russian; a
        # shifted column would put these near 0 or far above 1
        say(all(0.5 <= x <= 1.15 for x, _ in rs), f"{lang}: every ratio inside 0.5-1.15")

    # ---- 2014 at region level, a witness only ----
    if px:
        print("\n  2014 (PxWeb, region) against 2024 (this table), share of the region:")
        print(f"    {'region':<34}" + "".join(f"{l[:8]:>10}" for l in
                                               ("Georgian", "Azerbaijanian", "Armenian", "Russian")))
        for reg in regions:
            p = px.get(reg["name"])
            if not p:
                print(f"    {reg['name']:<34} (no 2014 row)")
                continue
            s14 = [p[l] / p["Total"] for l in ("Georgian", "Azerbaijanian", "Armenian", "Russian")]
            s24 = [reg[l] / reg["Total"] for l in ("Georgian", "Azerbaijanian", "Armenian",
                                                    "Russian")]
            print(f"    {reg['name']:<34}" + "".join(f"  {a:>4.0%}>{b:<4.0%}" for a, b in
                                                     zip(s14, s24)))
    if not ok:
        raise SystemExit("reconciliation FAILED")


PX_REGION = {"Autonomous Republic of Adjara": "Adjara A.R."}
PX_LANG = {"Abkhazian": "Abkhaz"}


def read_px():
    if not os.path.exists(PX_FILE):
        return None
    ds = json.load(open(PX_FILE, encoding="utf-8"))["dataset"]
    dim, val = ds["dimension"], ds["value"]
    ids, size = dim["id"], dim["size"]

    def axis(name):
        cat = dim[name]["category"]
        return [cat["label"][c] for c, _ in sorted(cat["index"].items(), key=lambda kv: kv[1])]

    reg, flu, lang = (axis(i) for i in ids)
    if flu[0] != "Total":
        raise SystemExit(f"2014 cube: Georgian-knowledge axis starts {flu[0]!r}")
    out = {}
    for i, rn in enumerate(reg):
        rn = PX_REGION.get(rn.strip(), rn.strip())
        rec = {}
        for k, ln in enumerate(lang):
            v = val[i * size[1] * size[2] + k]
            rec[PX_LANG.get(ln, ln)] = v if isinstance(v, (int, float)) else 0
        out[rn] = rec
    return out


def main():
    if "--fetch" in sys.argv:
        fetch()
    rows = levels(read_t5())
    t4 = _rows(_sheet("t4_nationality_2024.xlsx"))
    px = read_px()
    check(rows, t4, px)

    out = []
    for r in rows:
        gid = "Georgia" if r["level"] == "country" else r["name"]
        if r["level"] == "region":
            gid = "region:" + r["name"]
        for cat in ["Total"] + LANGS:
            note = f"level={r['level']}"
            if r["region"]:
                note += f"; region={r['region']}"
            if cat == "Total":
                note += "; universe total, not a language"
            if cat == "Not stated":
                note += "; not drawn (gap)"
            out.append({"geo_id": gid, "geo_level": r["level"], "geo_name": r["name"],
                        "source_category": cat, "count": r[cat], "tier": "measured",
                        "year": YEAR, "source_id": SOURCE_ID, "note": note})
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    with open(OUT, "w", encoding="utf-8", newline="") as fh:
        w = csv.DictWriter(fh, fieldnames=list(out[0]))
        w.writeheader()
        w.writerows(out)
    nat = rows[0]
    print(f"\n  national, {nat['Total']:,}:")
    for l in LANGS:
        print(f"    {nat[l]:>10,}  {100 * nat[l] / nat['Total']:6.2f}%  {l}")
    print(f"\nwrote {OUT} ({len(out):,} rows)")


if __name__ == "__main__":
    main()
