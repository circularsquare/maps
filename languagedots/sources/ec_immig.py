"""Ecuador, Censo 2022 (INEC): people born abroad per canton, by country of birth ->
data/normalized/ec_immig.csv (unit = canton, as data/geo/ec/ec_lookup.csv; iso; count). Read by
countries/ec.py, which turns them into languages with sources/latam_immig.py (record:
sources/ec.md, "Immigrant languages").

    python sources/ec_immig.py [--fetch]

THE SOURCE: INEC's tabulado "2022_CPV_Migracion.xlsx". www.censoecuador.gob.ec answers 403 from
here (sources/ec.md, Access), so --fetch takes the Wayback Machine's 2025-01-15 capture (it
arrives gzip-encoded). Two tables:
- 1.1: population by place of birth per province, canton and area; "En otro país" per canton
  is the COUNT drawn (425,045 nationally).
- 7: foreign-born by country of birth (190 columns) per province: the MIX, applied to every
  canton of the province. There is no canton-by-country table. The unnamed columns ("Otras
  Naciones De ...", "Zonas No Especificadas") are spread over the province's named countries
  in proportion.

CHECKS: 221 cantons, the same set as ec.csv's; cantons sum to their province's "En otro país",
and provinces to the national; table 7's province totals equal 1.1's province "En otro país";
table 7's countries sum to its own total per province; every country column maps to ISO or is a
listed unnamed one; ec_immig.csv sums to 425,045.
"""
import sys
import unicodedata
from pathlib import Path

import pandas as pd

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")

ROOT = Path(__file__).resolve().parent.parent
RAW = ROOT / "data" / "raw" / "ec"
NORM = ROOT / "data" / "normalized"
GEO = ROOT / "data" / "geo"
BOOK = RAW / "2022_CPV_Migracion.xlsx"
WAYBACK = ("http://web.archive.org/web/20250115011848id_/https://www.censoecuador.gob.ec/"
           "wp-content/uploads/2024/02/2022_CPV_Migracion.xlsx")
FOREIGN = 425_045

UNNAMED = {"Otras Naciones Americanas", "Otras Naciones De Asia", "Otras Naciones De Europa",
           "Otras Naciones De Oceanía", "Otras Naciones De África", "Zonas No Especificadas"}
ISO = {
    "Afganistan": "AF", "Albania": "AL", "Alemania": "DE", "Andorra": "AD", "Angola": "AO",
    "Anguila": "AI", "Antigua y Barbuda": "AG",
    "Antillas Neerlandesas (Antillas Holandesas)": "CW", "Arabia Saudita": "SA",
    "Argelia": "DZ", "Argentina": "AR", "Armenia": "AM", "Aruba": "AW", "Australia": "AU",
    "Austria": "AT", "Azerbaiyan": "AZ", "Bahamas": "BS", "Bahrein": "BH", "Bangladesh": "BD",
    "Barbados": "BB", "Belarus (Republica De Beilorrusia)": "BY", "Belice": "BZ", "Benin": "BJ",
    "Bermuda": "BM", "Bhutan": "BT", "Bolivia": "BO", "Bosnia Y Herzegovina": "BA",
    "Botswana (Botsuana)": "BW", "Brasil": "BR", "Bulgaria": "BG",
    "Burkina Faso (Antigua Republica Alto Bolta)": "BF", "Burundi": "BI", "Bélgica": "BE",
    "Cabo Verde": "CV", "Camerún": "CM", "Canadá": "CA", "Chile": "CL", "China": "CN",
    "Chipre": "CY", "Colombia": "CO", "Comoras": "KM", "Congo": "CG", "Costa Rica": "CR",
    "Cote D  Ivoire (Costa De Marfil)": "CI", "Croacia": "HR", "Cuba": "CU", "Dinamarca": "DK",
    "Dominica": "DM", "Egipto": "EG", "El Salvador": "SV", "Emiratos Arabes Unidos": "AE",
    "Eritrea": "ER", "Eslovaquia": "SK", "Eslovenia": "SI", "España": "ES",
    "Estados Unidos": "US", "Estonia": "EE", "Etiopia": "ET",
    "Ex Republica Yugoslavia (De Macedonia)": "MK", "Faja De Gaza (Franja De Gaza)": "PS",
    "Federación De Rusia (Antigua Union Sovietica)": "RU", "Fiji": "FJ", "Filipinas": "PH",
    "Finlandia": "FI", "Francia": "FR", "Gabón": "GA", "Georgia": "GE", "Ghana": "GH",
    "Granada": "GD", "Grecia": "GR", "Guadalupe": "GP", "Guatemala": "GT",
    "Guayana Francesa": "GF", "Guinea": "GN", "Guinea Bissau": "GW", "Guinea Ecuatorial": "GQ",
    "Guyana": "GY", "Haití": "HT", "Honduras": "HN",
    "Hong Kong (Región Administrativa Especial De China)": "HK", "Hungría": "HU",
    "India": "IN", "Indonesia": "ID", "Iraq": "IQ", "Irlanda": "IE",
    "Irán (Republica Islámica Del)": "IR", "Islandia": "IS", "Islas Caiman": "KY",
    "Islas Turcas Y Caicos": "TC", "Islas Vírgenes Británicas": "VG",
    "Islas Vírgenes De Los Estados Unidos": "VI", "Israel": "IL", "Italia": "IT",
    "Jamahiriya Arabe Libia (Libia)": "LY", "Jamaica": "JM", "Japón": "JP", "Jordania": "JO",
    "Kazajstan": "KZ", "Kenya": "KE", "Kirguistan": "KG", "Kuwait": "KW",
    "Letonia (Latvijas)": "LV", "Liberia": "LR", "Lituania": "LT", "Luxemburgo": "LU",
    "Líbano": "LB", "Macao": "MO", "Madagascar": "MG", "Malasia (Malasia Peninsular)": "MY",
    "Malawi": "MW", "Mali": "ML", "Malta": "MT", "Marruecos": "MA", "Mauritania": "MR",
    "Micronesia (Estados Federados De)": "FM", "Mongolia": "MN", "Mozambique": "MZ",
    "Myanmar (Birmania)": "MM", "México": "MX", "Nepal": "NP", "Nicaragua": "NI",
    "Nigeria": "NG", "Noruega": "NO", "Nueva Caledonia": "NC", "Nueva Zelandia": "NZ",
    "Oman": "OM", "Pakistan": "PK", "Panamá": "PA", "Paraguay": "PY",
    "Países Bajos ( Isla Curazao; Holanda; Zelanda) ": "NL", "Perú": "PE", "Polonia": "PL",
    "Portugal": "PT", "Provincia China De Taiwan": "TW", "Puerto Rico": "PR", "Qatar": "QA",
    "Reino Unido": "GB", "Republica Arabe Siria": "SY", "Republica Checa": "CZ",
    "Republica De Moldova": "MD", "Republica Unida De Tanzania": "TZ",
    "República Centroafricana": "CF", "República De Corea (Corea Del Sur)": "KR",
    "República Democrática Del Congo (Zaire)": "CD", "República Dominicana": "DO",
    "República Popular Democrática De Corea (Corea Del Norte)": "KP",
    "Reunion (Isla De La)": "RE", "Rumania": "RO", "Rwanda (Ruanda)": "RW",
    "Sahara Occidental": "EH", "Samoa (Samoa Occidental)": "WS", "Samoa Americana": "AS",
    "San Vicente Y Las Granadinas": "VC", "Santa Elena (Isla)": "SH",
    "Santo Tome y Príncipe": "ST", "Senegal": "SN", "Serbia Y Montenegro (Yugoslavia)": "RS",
    "Sierra Leona": "SL", "Singapur": "SG", "Sri Lanka (Ceilan)": "LK", "Sudan": "SD",
    "Sudáfrica": "ZA", "Suecia": "SE", "Suiza": "CH", "Suriname": "SR", "Tailandia": "TH",
    "Tayikistan (Tajikistan)": "TJ", "Timor Oriental (Timor Del Este)": "TL", "Togo": "TG",
    "Trinidad Y Tobago": "TT", "Tunez": "TN", "Turkmenistan": "TM", "Turquia": "TR",
    "Ucrania": "UA", "Uganda": "UG", "Uruguay": "UY", "Uzbekistan": "UZ", "Vanuatu": "VU",
    "Venezuela": "VE", "Vietnam": "VN", "Yemen": "YE", "Zambia": "ZM",
    "Zimbabwe (Zimbabue)": "ZW",
}


def fold(s):
    s = unicodedata.normalize("NFKD", str(s)).encode("ascii", "ignore").decode()
    return " ".join(s.lower().split())


def fetch():
    import gzip
    import urllib.request
    req = urllib.request.Request(WAYBACK, headers={"User-Agent": "Mozilla/5.0"})
    b = urllib.request.urlopen(req, timeout=300).read()
    if b[:2] == b"\x1f\x8b":
        b = gzip.decompress(b)
    BOOK.write_bytes(b)
    print("fetched", BOOK.name, len(b))


def read():
    import openpyxl
    wb = openpyxl.load_workbook(BOOK, read_only=True)
    r11 = [r for r in wb["1.1"].iter_rows(values_only=True)]
    r7 = [r for r in wb["7"].iter_rows(values_only=True)]
    # 1.1: rows (_, province, canton label, area label, total, ..., foreign at column 13)
    cant, prov, nat = {}, {}, None
    for r in r11:
        if not r or not isinstance(r[1], str) or r[3] is None:
            continue
        p, c, a = r[1].strip(), str(r[2]).strip(), str(r[3]).strip()
        if p == "Total Nacional" and a == "Nacional":
            nat = r[13]
        elif c.startswith("Total ") and a == c:
            prov[p] = r[13]
        elif a == f"Total {c}":
            cant[(p, c)] = r[13]
    head = next(r for r in r7 if r and "Afganistan" in r)
    cols = [(i, h) for i, h in enumerate(head) if isinstance(h, str)]
    mix, tot7 = {}, {}
    for r in r7:
        if not r or not isinstance(r[1], str) or not isinstance(r[2], str):
            continue
        if r[2].strip() == f"Total {r[1].strip()}":
            tot7[r[1].strip()] = r[3]
            mix[r[1].strip()] = {h: (r[i] or 0) for i, h in cols}
    return cant, prov, nat, mix, tot7, [h for _, h in cols]


def main():
    if "--fetch" in sys.argv or not BOOK.exists():
        fetch()
    ok = True

    def check(c, m):
        nonlocal ok
        print(("ok    " if c else "FAIL  ") + m)
        ok &= bool(c)

    cant, prov, nat, mix, tot7, names = read()
    unknown = [h for h in names if h not in ISO and h not in UNNAMED]
    check(not unknown, f"{len(names)} country columns map to ISO or are unnamed ({unknown[:4]})")
    check(len(cant) == 221 and len(prov) == 24, f"{len(cant)} cantons, {len(prov)} provinces")
    check(nat == FOREIGN and sum(prov.values()) == nat, f"provinces sum to {nat:,}")
    bad = [p for p in prov if sum(v for (q, _), v in cant.items() if q == p) != prov[p]]
    check(not bad, f"cantons sum to their province ({bad[:3]})")
    check(set(tot7) == set(prov) and all(tot7[p] == prov[p] for p in prov),
          "table 7's province totals equal table 1.1's")
    bad = [p for p in mix if sum(mix[p].values()) != tot7[p]]
    check(not bad, f"table 7's countries sum to its province total ({bad[:3]})")

    lut = pd.read_csv(GEO / "ec" / "ec_lookup.csv", dtype=str)
    key = {(fold(p), fold(c)): u for p, c, u in zip(lut["province"], lut["canton"], lut["unit"])}
    units = {k: key.get((fold(k[0]), fold(k[1]))) for k in cant}
    miss = [k for k, u in units.items() if u is None]
    check(not miss and len(set(units.values())) == 221,
          f"every canton joins ec_lookup.csv one to one ({miss[:3]})")
    if not ok:
        raise SystemExit("checks failed")
    rows = []
    for (p, c), n in cant.items():
        named = {ISO[h]: v for h, v in mix[p].items() if h in ISO and v}
        t = sum(named.values())
        rows += [(units[(p, c)], iso, n * v / t) for iso, v in named.items()]
    df = pd.DataFrame(rows, columns=["unit", "iso", "count"])
    df = df.groupby(["unit", "iso"], as_index=False)["count"].sum()
    check(abs(df["count"].sum() - FOREIGN) < 1, f"ec_immig sums to {df['count'].sum():,.0f}")
    df.to_csv(NORM / "ec_immig.csv", index=False)
    top = df.groupby("iso")["count"].sum().sort_values(ascending=False).head(10)
    unnamed = sum(v for p in mix for h, v in mix[p].items() if h in UNNAMED)
    print(f"wrote ec_immig.csv ({len(df):,} rows; {unnamed:,} unnamed spread); top: "
          + ", ".join(f"{k} {v:,.0f}" for k, v in top.items()))


if __name__ == "__main__":
    main()
